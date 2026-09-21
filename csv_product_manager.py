"""CSV price-list uploader and editor for product records."""

import csv
import os
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

import openpyxl

import category_order
import config


class CSVProductManagerWindow(tk.Toplevel):
    def __init__(self, parent, default_file=None):
        super().__init__(parent)
        self.title("Panel Price List Manager")
        self.geometry("1100x680")
        self.minsize(900, 550)
        self.transient(parent)
        self.grab_set()
        self.lift()
        self.focus_force()
        self.file_path = None
        self.headers = []
        self.rows = []
        self.field_vars = {}
        self.selected_index = None
        self.file_name_var = tk.StringVar(value="No CSV file selected")
        self._build()
        self.build_editor_for_empty_state()
        default_file = default_file or os.path.join(config.SCRIPT_DIR, "price_list_clean.csv")
        if os.path.exists(default_file):
            self.open_file(default_file)

    def _build(self):
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)

        top = ttk.Frame(root)
        top.pack(fill="x", pady=(12, 8))
        ttk.Button(top, text="Upload CSV File", command=self.choose_file).pack(side="left")
        ttk.Button(top, text="⬆ Upload Excel", command=self.choose_excel_file).pack(side="left", padx=(8, 0))
        ttk.Button(top, text="⬇ Download Excel", command=self.export_excel_file).pack(side="left", padx=(8, 0))
        ttk.Label(top, textvariable=self.file_name_var).pack(side="left", padx=12)
        ttk.Button(top, text="Save Changes", command=self.update_csv_file).pack(side="right", padx=(8, 0))
        ttk.Button(top, text="Reload", command=self.reload_file).pack(side="right")

        table_box = ttk.LabelFrame(root, text=" Panel Products Price List ", padding=8)
        table_box.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(table_box, show="headings", selectmode="browse")
        scroll_y = ttk.Scrollbar(table_box, orient="vertical", command=self.tree.yview)
        scroll_x = ttk.Scrollbar(table_box, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scroll_y.pack(side="right", fill="y")
        scroll_x.pack(side="bottom", fill="x")
        self.tree.bind("<<TreeviewSelect>>", self.load_selected_product)
        self.tree.bind("<Double-1>", self.edit_selected_product)

        self.editor = ttk.LabelFrame(root, text=" Product Details ", padding=10)
        self.editor.pack(fill="x", pady=(10, 0))
        self.editor_inner = ttk.Frame(self.editor)
        self.editor_inner.pack(fill="x")

        actions = ttk.Frame(root, padding=(0, 8, 0, 0))
        actions.pack(fill="x")
        ttk.Button(actions, text="+ New Product", command=self.new_product).pack(side="left")
        ttk.Button(actions, text="Edit Selected", command=self.edit_selected_product).pack(side="left", padx=8)
        ttk.Button(actions, text="Save Changes", command=self.save_product).pack(side="left", padx=8)
        ttk.Button(actions, text="Update CSV File", command=self.update_csv_file).pack(side="left", padx=8)
        ttk.Button(actions, text="+ Add Column", command=self.add_column).pack(side="left", padx=8)
        ttk.Button(actions, text="Rename Column", command=self.rename_column).pack(side="left")
        ttk.Button(actions, text="Remove Column", command=self.remove_column).pack(side="left", padx=8)
        ttk.Button(actions, text="Edit Header Row", command=self.edit_all_column_names).pack(side="left")

    def build_editor_for_empty_state(self):
        for widget in self.editor_inner.winfo_children():
            widget.destroy()
        self.field_vars = {}
        ttk.Label(self.editor_inner, text="Upload a CSV file to edit products or create a new product after opening a file.", wraplength=700).pack(anchor="w")

    def choose_file(self):
        path = filedialog.askopenfilename(parent=self, title="Choose Price List CSV File", filetypes=(("CSV files", "*.csv"),))
        if path:
            self.open_file(path)

    def choose_excel_file(self):
        """Select an Excel workbook whose sheet will update the price list."""
        path = filedialog.askopenfilename(
            parent=self,
            title="Upload Excel Price List",
            filetypes=(
                ("Excel files", "*.xlsx *.xlsm"),
                ("All files", "*.*"),
            ),
        )
        if path:
            self.import_excel_to_price_list(path)

    def import_excel_to_price_list(self, excel_path):
        """Read an Excel sheet and save it as the active price list CSV."""
        try:
            workbook = openpyxl.load_workbook(excel_path, data_only=True)
        except Exception as error:
            messagebox.showerror(
                "Cannot Open Excel",
                f"Could not open this Excel file:\n{error}",
                parent=self,
            )
            return

        sheet_name = self.pick_excel_sheet(workbook, os.path.basename(excel_path))
        if not sheet_name:
            return

        worksheet = workbook[sheet_name]
        rows = [
            ["" if value is None else value for value in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(str(value).strip() for value in row)]
        if not rows:
            messagebox.showwarning(
                "Empty Sheet",
                "The selected worksheet does not contain any data.",
                parent=self,
            )
            return
        if not any(str(value).strip() for value in rows[0]):
            messagebox.showwarning(
                "Missing Headers",
                "The first row of the worksheet must contain column names.",
                parent=self,
            )
            return

        # Arrange the uploaded sheet category-wise before saving.
        rows = category_order.sort_rows_by_category(rows)

        target_path = self.file_path or config.PRODUCT_PRICE_CALCULATOR_CSV
        if not messagebox.askyesno(
            "Update Price List",
            "Replace the current price list with the uploaded Excel sheet?\n\n"
            f"Target file:\n{target_path}",
            parent=self,
        ):
            return

        try:
            os.makedirs(os.path.dirname(target_path), exist_ok=True)
            with open(target_path, "w", newline="", encoding="utf-8-sig") as file:
                csv.writer(file).writerows(rows)
        except PermissionError:
            messagebox.showerror(
                "File Is Open",
                "Close the price list CSV in another program, then upload again.",
                parent=self,
            )
            return
        except OSError as error:
            messagebox.showerror(
                "Excel Upload Failed",
                f"Could not update the price list:\n{error}",
                parent=self,
            )
            return

        self.open_file(target_path)
        messagebox.showinfo(
            "Excel Uploaded",
            f"The sheet '{sheet_name}' was uploaded and the price list was updated:\n{target_path}",
            parent=self,
        )

    def pick_excel_sheet(self, workbook, file_name):
        """Ask which worksheet to upload. Returns the sheet name or None."""
        sheet_names = workbook.sheetnames
        if not sheet_names:
            messagebox.showwarning(
                "No Worksheets",
                f"The workbook '{file_name}' does not contain any worksheets.",
                parent=self,
            )
            return None
        if len(sheet_names) == 1:
            return sheet_names[0]

        dialog = tk.Toplevel(self)
        dialog.title("Select Worksheet")
        dialog.geometry("440x180")
        dialog.resizable(False, False)
        dialog.transient(self)
        dialog.grab_set()
        dialog.lift()
        dialog.focus_force()

        active_title = workbook.active.title if workbook.active else sheet_names[0]
        selected = tk.StringVar(value=active_title)

        content = ttk.Frame(dialog, padding=18)
        content.pack(fill="both", expand=True)
        ttk.Label(
            content,
            text=f"Choose the worksheet to upload from '{file_name}':",
            wraplength=400,
        ).pack(anchor="w", pady=(0, 12))
        ttk.Combobox(
            content,
            textvariable=selected,
            values=sheet_names,
            state="readonly",
            width=42,
        ).pack(fill="x")

        result = {"sheet": None}

        def confirm():
            result["sheet"] = selected.get()
            dialog.destroy()

        buttons = ttk.Frame(content)
        buttons.pack(anchor="e", pady=(18, 0))
        ttk.Button(buttons, text="Upload", command=confirm).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="Cancel", command=dialog.destroy).pack(side="left")

        dialog.wait_window()
        return result["sheet"]

    def export_excel_file(self):
        """Save the currently open price list to an Excel workbook."""
        if not self.headers:
            messagebox.showwarning("Upload Needed", "Open a price list first.", parent=self)
            return

        destination = filedialog.asksaveasfilename(
            parent=self,
            title="Download Price List as Excel",
            defaultextension=".xlsx",
            initialfile="price_list.xlsx",
            filetypes=(("Excel workbook", "*.xlsx"),),
        )
        if not destination:
            return

        try:
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.title = "Price List"
            worksheet.append(self.headers)
            for row in self.rows:
                worksheet.append([row.get(header, "") for header in self.headers])
            workbook.save(destination)
            messagebox.showinfo(
                "Excel Downloaded",
                f"The price list was saved to:\n{destination}",
                parent=self,
            )
        except Exception as error:
            messagebox.showerror(
                "Excel Export Failed",
                f"Could not download the price list:\n{error}",
                parent=self,
            )

    def open_file(self, path):
        try:
            with open(path, newline="", encoding="utf-8-sig") as file:
                rows = list(csv.reader(file))
            if not rows:
                raise ValueError("The CSV file is empty.")

            header_row = rows[0]
            named_columns = [
                (index, header.strip())
                for index, header in enumerate(header_row)
                if header.strip()
            ]
            self.headers = [header for _, header in named_columns]
            self.rows = [
                {
                    header: data[index] if index < len(data) else ""
                    for index, header in named_columns
                }
                for data in rows[1:]
            ]
            if not self.headers:
                raise ValueError("The CSV file does not have a header row.")
            # Keep the list arranged category-wise every time it is opened.
            self.rows = category_order.sort_rows_by_category(self.rows)
            self.file_path = path
            self.file_name_var.set(os.path.basename(path))
            self.build_editor()
            self.refresh_table()
            self.new_product()
        except Exception as error:
            messagebox.showerror("Cannot Open CSV", f"Could not open this CSV file:\n{error}", parent=self)

    def build_editor(self):
        for widget in self.editor_inner.winfo_children():
            widget.destroy()
        self.field_vars = {header: tk.StringVar() for header in self.headers}
        for index, header in enumerate(self.headers):
            row, group = divmod(index, 2)
            column = group * 2
            ttk.Label(self.editor_inner, text=f"{header}:").grid(row=row, column=column, sticky="w", padx=(0, 6), pady=4)
            ttk.Entry(self.editor_inner, textvariable=self.field_vars[header], width=34).grid(row=row, column=column + 1, sticky="ew", padx=(0, 18), pady=4)
        self.editor_inner.columnconfigure(1, weight=1)
        self.editor_inner.columnconfigure(3, weight=1)
        ttk.Button(
            self.editor_inner,
            text="Save Changes",
            command=self.save_product,
        ).grid(row=(len(self.headers) + 1) // 2, column=0, columnspan=4, sticky="e", pady=(10, 0))

    def refresh_table(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.tree["columns"] = self.headers
        for header in self.headers:
            self.tree.heading(header, text=header)
            self.tree.column(header, width=145, minwidth=90, anchor="w")
        for index, row in enumerate(self.rows):
            self.tree.insert("", "end", iid=str(index), values=[row.get(header, "") for header in self.headers])

    def load_selected_product(self, _event=None):
        selected = self.tree.selection()
        if not selected:
            return
        self.selected_index = int(selected[0])
        row = self.rows[self.selected_index]
        for header, variable in self.field_vars.items():
            variable.set(row.get(header, ""))

    def edit_selected_product(self, _event=None):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning(
                "Select Product",
                "Select a panel product row first, then click Edit Selected.",
                parent=self,
            )
            return "break" if _event else None
        self.load_selected_product()
        self.editor.focus_set()
        return "break" if _event else None

    def new_product(self):
        self.selected_index = None
        for variable in self.field_vars.values():
            variable.set("")
        for item in self.tree.selection():
            self.tree.selection_remove(item)

    def save_product(self):
        if not self.file_path:
            messagebox.showwarning("Upload Needed", "Upload a CSV price list first.", parent=self)
            return
        if not self.field_vars:
            messagebox.showwarning("Upload Needed", "Open a valid CSV file before creating or editing products.", parent=self)
            return
        row = {header: variable.get().strip() for header, variable in self.field_vars.items()}
        if not any(row.values()):
            messagebox.showwarning("Product Details", "Enter at least one product detail before saving.", parent=self)
            return
        if self.selected_index is None:
            self.rows.append(row)
        else:
            self.rows[self.selected_index] = row
        if self.write_csv():
            self.refresh_table()
            self.new_product()
            messagebox.showinfo("Changes Saved", "Your changes were saved to the CSV file.", parent=self)

    def update_csv_file(self):
        """Write all current editor rows and columns to the active CSV file."""
        if not self.file_path or not self.headers:
            messagebox.showwarning("Upload Needed", "Open a CSV price list first.", parent=self)
            return
        if self.selected_index is not None and self.selected_index < len(self.rows):
            self.rows[self.selected_index] = {
                header: variable.get().strip()
                for header, variable in self.field_vars.items()
            }
        if self.write_csv():
            self.open_file(self.file_path)
            messagebox.showinfo(
                "CSV Updated",
                f"CSV file updated successfully:\n{self.file_path}",
                parent=self,
            )

    def add_column(self):
        if not self.file_path:
            messagebox.showwarning("Upload Needed", "Open a CSV price list first.", parent=self)
            return

        column_name = simpledialog.askstring(
            "Add Column",
            "Enter the new column name:",
            parent=self,
        )
        column_name = (column_name or "").strip()
        if not column_name:
            return
        if column_name.lower() in {header.lower() for header in self.headers}:
            messagebox.showwarning("Duplicate Column", "That column name already exists.", parent=self)
            return

        self.headers.append(column_name)
        for row in self.rows:
            row[column_name] = ""
        self.write_csv()
        self.build_editor()
        self.refresh_table()
        self.new_product()

    def rename_column(self):
        if not self.file_path or not self.headers:
            messagebox.showwarning("Upload Needed", "Open a CSV price list first.", parent=self)
            return

        old_name = simpledialog.askstring(
            "Rename Column",
            "Current column name:",
            parent=self,
        )
        old_name = (old_name or "").strip()
        if not old_name:
            return
        matching_headers = [header for header in self.headers if header.lower() == old_name.lower()]
        if not matching_headers:
            messagebox.showwarning("Column Not Found", f"No column named '{old_name}' exists.", parent=self)
            return

        new_name = simpledialog.askstring(
            "Rename Column",
            f"New name for '{matching_headers[0]}':",
            parent=self,
        )
        new_name = (new_name or "").strip()
        if not new_name:
            return
        if new_name.lower() in {header.lower() for header in self.headers if header != matching_headers[0]}:
            messagebox.showwarning("Duplicate Column", "That column name already exists.", parent=self)
            return

        old_name = matching_headers[0]
        header_index = self.headers.index(old_name)
        self.headers[header_index] = new_name
        for row in self.rows:
            row[new_name] = row.pop(old_name, "")
        self.write_csv()
        self.build_editor()
        self.refresh_table()
        self.new_product()

    def remove_column(self):
        if not self.file_path or not self.headers:
            messagebox.showwarning("Upload Needed", "Open a CSV price list first.", parent=self)
            return
        if len(self.headers) == 1:
            messagebox.showwarning("Cannot Remove Column", "A CSV file must keep at least one column.", parent=self)
            return

        column_name = simpledialog.askstring(
            "Remove Column",
            "Enter the column name to remove:",
            parent=self,
        )
        column_name = (column_name or "").strip()
        if not column_name:
            return
        matching_headers = [header for header in self.headers if header.lower() == column_name.lower()]
        if not matching_headers:
            messagebox.showwarning("Column Not Found", f"No column named '{column_name}' exists.", parent=self)
            return

        column_name = matching_headers[0]
        if not messagebox.askyesno(
            "Remove Column",
            f"Remove '{column_name}' and all values in this column?",
            parent=self,
        ):
            return

        self.headers.remove(column_name)
        for row in self.rows:
            row.pop(column_name, None)
        self.write_csv()
        self.build_editor()
        self.refresh_table()
        self.new_product()

    def edit_all_column_names(self):
        if not self.file_path or not self.headers:
            messagebox.showwarning("Upload Needed", "Open a CSV price list first.", parent=self)
            return

        editor = tk.Toplevel(self)
        editor.title("Edit Panel Price List Header Row")
        editor.geometry("620x650")
        editor.minsize(520, 420)
        editor.transient(self)
        editor.grab_set()

        content = ttk.Frame(editor, padding=14)
        content.pack(fill="both", expand=True)
        ttk.Label(
            content,
            text="Edit Panel Price List Header Row",
            font=("Helvetica", 13, "bold"),
        ).pack(anchor="w", pady=(0, 4))
        ttk.Label(
            content,
            text="Edit SKU, Category, make, sub-category fields, or any other first-row heading.",
            wraplength=560,
        ).pack(anchor="w", pady=(0, 10))

        canvas = tk.Canvas(content, highlightthickness=0)
        scrollbar = ttk.Scrollbar(content, orient="vertical", command=canvas.yview)
        fields = ttk.Frame(canvas)
        fields.bind("<Configure>", lambda _event: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=fields, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        name_vars = []
        for index, header in enumerate(self.headers, start=1):
            variable = tk.StringVar(value=header)
            name_vars.append(variable)
            ttk.Label(fields, text=f"Column {index}:").grid(row=index - 1, column=0, sticky="w", padx=(0, 10), pady=4)
            ttk.Entry(fields, textvariable=variable, width=50).grid(row=index - 1, column=1, sticky="ew", pady=4)
        fields.columnconfigure(1, weight=1)

        def save_column_names():
            new_headers = [variable.get().strip() for variable in name_vars]
            if any(not header for header in new_headers):
                messagebox.showwarning("Invalid Column Name", "Every column must have a name.", parent=editor)
                return
            lowered_headers = [header.lower() for header in new_headers]
            if len(set(lowered_headers)) != len(lowered_headers):
                messagebox.showwarning("Duplicate Column Name", "Column names must be unique.", parent=editor)
                return

            old_headers = self.headers[:]
            self.headers = new_headers
            self.rows = [
                {
                    new_header: row.get(old_header, "")
                    for old_header, new_header in zip(old_headers, new_headers)
                }
                for row in self.rows
            ]
            if self.write_csv():
                self.build_editor()
                self.refresh_table()
                self.new_product()
                editor.destroy()

        ttk.Button(content, text="Save Column Names", command=save_column_names).pack(anchor="e", pady=(12, 0))

    def write_csv(self):
        try:
            with open(self.file_path, "w", newline="", encoding="utf-8-sig") as file:
                writer = csv.DictWriter(file, fieldnames=self.headers, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(self.rows)
            return True
        except PermissionError:
            messagebox.showerror(
                "File Is Open",
                "Close the CSV file in another program, then update again.",
                parent=self,
            )
        except OSError as error:
            messagebox.showerror(
                "CSV Update Failed",
                f"Could not update the CSV file:\n{error}",
                parent=self,
            )
        return False

    def reload_file(self):
        if self.file_path:
            self.open_file(self.file_path)
