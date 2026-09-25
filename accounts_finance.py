"""Accounts and finance workspace with invoice-based inward entries."""

import csv
import os
import re
import shutil
import tkinter as tk
from datetime import date, datetime
from tkinter import filedialog, messagebox, ttk

import config
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape
from tkcalendar import DateEntry
from outward_window import OutwardWindow
from vendor_registration import VendorRegistrationWindow


DEPARTMENTS = ("Production", "R&D", "Panel Department", "Office", "Sales", "Marketing", "Other")
INWARD_HEADERS = ["date", "supplier", "invoice_no", "department", "item_description", "quantity", "rate", "amount", "payment_status", "payment_method", "cheque_date", "cheque_number", "cheque_photo", "notes"]
INWARD_CSV = os.path.join(config.CSV_DIR, "inward_entries.csv")
PAYMENTS_CSV = os.path.join(config.CSV_DIR, "invoice_payments.csv")
PAYMENT_HEADERS = ["payment_date", "supplier", "invoice_no", "amount", "payment_method"]
PRODUCT_CATALOG_CSV = os.path.join(config.SCRIPT_DIR, "price_list_clean.csv")
CHEQUE_PHOTOS_DIR = os.path.join(config.SCRIPT_DIR, "cheque_photos")
INWARD_OUTPUT_DIR = os.path.join(config.SCRIPT_DIR, "inward_invoices")
CUSTOMERS_DETAILED_CSV = os.path.join(config.CSV_DIR, "customers_detailed.csv")

# Sales / Income
SALES_INCOME_HEADERS = [
    "date", "customer_name", "invoice_number", "product_service",
    "invoice_amount", "amount_received", "pending_amount", "notes"
]
SALES_INCOME_CSV = os.path.join(config.CSV_DIR, "sales_income.csv")

# Expenses
EXPENSE_CATEGORIES = (
    "Salary", "Rent", "Electricity", "Travel", "Material Purchase",
    "Software", "Internet / Mobile", "Repairs", "Other",
)
EXPENSES_HEADERS = [
    "date", "category", "description", "amount", "payment_method", "paid_to", "notes"
]
EXPENSES_CSV = os.path.join(config.CSV_DIR, "expenses.csv")

# Bank & Cash
BANK_CASH_TRANSACTION_TYPES = (
    "Deposit", "Withdrawal", "Bank Charge", "Cash In", "Cash Out", "Transfer",
)
BANK_ACCOUNTS_HEADERS = ["account_name", "account_number", "bank_name", "opening_balance", "notes"]
BANK_ACCOUNTS_CSV     = os.path.join(config.CSV_DIR, "bank_accounts.csv")
BANK_TRANSACTIONS_HEADERS = [
    "date", "account_name", "transaction_type", "amount",
    "description", "reference", "notes"
]
BANK_TRANSACTIONS_CSV = os.path.join(config.CSV_DIR, "bank_transactions.csv")

# Accounts Receivable
AR_STATUS_OPTIONS = ("Unpaid", "Partial", "Paid", "Overdue", "Disputed")
AR_HEADERS = [
    "invoice_date", "due_date", "customer_name", "invoice_number",
    "product_service", "invoice_amount", "amount_received",
    "pending_amount", "status", "notes"
]
AR_CSV = os.path.join(config.CSV_DIR, "accounts_receivable.csv")

# Accounts Payable
AP_PAYMENT_TYPES = ("Supplier Payment", "Vendor Bill", "Loan / EMI", "Rent", "Utility", "Other Liability")
AP_STATUS_OPTIONS = ("Unpaid", "Partial", "Paid", "Overdue", "Disputed")
AP_HEADERS = [
    "bill_date", "due_date", "payee_name", "bill_number",
    "payment_type", "bill_amount", "amount_paid",
    "pending_amount", "status", "notes"
]
AP_CSV = os.path.join(config.CSV_DIR, "accounts_payable.csv")

# Inventory / Stock
INVENTORY_HEADERS = [
    "item_code", "item_name", "category", "unit",
    "opening_qty", "purchased_qty", "used_qty", "finished_qty",
    "current_stock", "unit_cost", "stock_value", "reorder_level", "notes"
]
INVENTORY_CATEGORIES = (
    "Raw Material", "Semi-Finished", "Finished Goods",
    "Consumable", "Spare Parts", "Packaging", "Other",
)
INVENTORY_CSV = os.path.join(config.CSV_DIR, "inventory_stock.csv")

# Employee Salary
SALARY_STATUS_OPTIONS = ("Paid", "Unpaid", "Partial", "Hold")
SALARY_HEADERS = [
    "month_year", "employee_name", "designation", "monthly_salary",
    "working_days", "present_days", "advance", "deductions",
    "net_payable", "amount_paid", "balance", "status", "notes"
]
SALARY_CSV = os.path.join(config.CSV_DIR, "employee_salary.csv")

# GST & Tax
GST_ENTRY_TYPES = (
    "Sales GST", "Purchase GST / ITC", "GST Payable",
    "GST Return Filed", "TDS Payable", "TDS Deducted",
    "Income Tax", "Advance Tax", "Other Tax",
)
GST_HEADERS = [
    "date", "entry_type", "reference_number", "party_name",
    "taxable_amount", "cgst", "sgst", "igst", "total_gst",
    "tds_amount", "period", "status", "notes"
]
GST_STATUS_OPTIONS = ("Pending", "Filed", "Paid", "Overdue")
GST_CSV = os.path.join(config.CSV_DIR, "gst_tax.csv")


def ensure_inward_file():
    """Create the register or add the department column to old data."""
    if not os.path.exists(INWARD_CSV):
        with open(INWARD_CSV, "w", newline="", encoding="utf-8") as file:
            csv.DictWriter(file, fieldnames=INWARD_HEADERS).writeheader()
        return
    with open(INWARD_CSV, newline="", encoding="utf-8") as file:
        reader = csv.DictReader(file)
        old_headers, rows = reader.fieldnames or [], list(reader)
    if any(header not in old_headers for header in INWARD_HEADERS):
        for row in rows:
            for header in INWARD_HEADERS:
                row.setdefault(header, "")
        with open(INWARD_CSV, "w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=INWARD_HEADERS)
            writer.writeheader()
            writer.writerows(rows)


def ensure_payments_file():
    if not os.path.exists(PAYMENTS_CSV):
        with open(PAYMENTS_CSV, "w", newline="", encoding="utf-8") as file:
            csv.DictWriter(file, fieldnames=PAYMENT_HEADERS).writeheader()


def ensure_sales_income_file():
    """Create sales_income.csv if missing, or migrate missing columns."""
    if not os.path.exists(SALES_INCOME_CSV):
        with open(SALES_INCOME_CSV, "w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=SALES_INCOME_HEADERS).writeheader()
        return
    with open(SALES_INCOME_CSV, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        old_headers, rows = reader.fieldnames or [], list(reader)
    if any(h not in old_headers for h in SALES_INCOME_HEADERS):
        for row in rows:
            for h in SALES_INCOME_HEADERS:
                row.setdefault(h, "")
        with open(SALES_INCOME_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=SALES_INCOME_HEADERS)
            writer.writeheader()
            writer.writerows(rows)


def ensure_expenses_file():
    """Create expenses.csv if missing, or migrate missing columns."""
    if not os.path.exists(EXPENSES_CSV):
        with open(EXPENSES_CSV, "w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=EXPENSES_HEADERS).writeheader()
        return
    with open(EXPENSES_CSV, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        old_headers, rows = reader.fieldnames or [], list(reader)
    if any(h not in old_headers for h in EXPENSES_HEADERS):
        for row in rows:
            for h in EXPENSES_HEADERS:
                row.setdefault(h, "")
        with open(EXPENSES_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=EXPENSES_HEADERS)
            writer.writeheader()
            writer.writerows(rows)


def _migrate_csv(csv_path, headers):
    """Generic helper — create CSV if missing or add any new columns."""
    if not os.path.exists(csv_path):
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=headers).writeheader()
        return
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        old_headers, rows = reader.fieldnames or [], list(reader)
    if any(h not in old_headers for h in headers):
        for row in rows:
            for h in headers:
                row.setdefault(h, "")
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)


def ensure_bank_files():
    """Create / migrate both bank CSVs."""
    _migrate_csv(BANK_ACCOUNTS_CSV,     BANK_ACCOUNTS_HEADERS)
    _migrate_csv(BANK_TRANSACTIONS_CSV, BANK_TRANSACTIONS_HEADERS)


def ensure_ar_file():
    """Create / migrate accounts_receivable.csv."""
    _migrate_csv(AR_CSV, AR_HEADERS)


def ensure_ap_file():
    """Create / migrate accounts_payable.csv."""
    _migrate_csv(AP_CSV, AP_HEADERS)


def ensure_inventory_file():
    """Create / migrate inventory_stock.csv."""
    _migrate_csv(INVENTORY_CSV, INVENTORY_HEADERS)


def ensure_salary_file():
    """Create / migrate employee_salary.csv."""
    _migrate_csv(SALARY_CSV, SALARY_HEADERS)


def ensure_gst_file():
    """Create / migrate gst_tax.csv."""
    _migrate_csv(GST_CSV, GST_HEADERS)


def load_customer_names():
    """Return a sorted unique list of company names from customers_detailed.csv."""
    names = []
    if os.path.exists(CUSTOMERS_DETAILED_CSV):
        with open(CUSTOMERS_DETAILED_CSV, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                name = row.get("company_name", "").strip()
                if name and name not in names:
                    names.append(name)
    return sorted(names)


def get_invoice_balances():
    """Return invoice totals, recorded payments, and current balances."""
    ensure_inward_file()
    ensure_payments_file()
    invoices = {}
    with open(INWARD_CSV, newline="", encoding="utf-8") as file:
        for row in csv.DictReader(file):
            key = (row.get("supplier", ""), row.get("invoice_no", ""))
            if not all(key):
                continue
            try:
                amount = float(row.get("amount", 0) or 0)
            except ValueError:
                amount = 0.0
            invoice = invoices.setdefault(key, {"supplier": key[0], "invoice_no": key[1], "department": row.get("department", ""), "total": 0.0, "paid": 0.0, "status": row.get("payment_status", "Pending")})
            invoice["total"] += amount
    with open(PAYMENTS_CSV, newline="", encoding="utf-8") as file:
        for row in csv.DictReader(file):
            invoice = invoices.get((row.get("supplier", ""), row.get("invoice_no", "")))
            if not invoice:
                continue
            try:
                invoice["paid"] += float(row.get("amount", 0) or 0)
            except ValueError:
                pass
    for invoice in invoices.values():
        if invoice["status"] == "Paid" and invoice["paid"] == 0:
            invoice["paid"] = invoice["total"]
        invoice["paid"] = min(invoice["paid"], invoice["total"])
        invoice["balance"] = max(invoice["total"] - invoice["paid"], 0.0)
    return list(invoices.values())


class InvoicePaymentDialog(tk.Toplevel):
    def __init__(self, parent, invoice, on_saved):
        super().__init__(parent)
        self.invoice, self.on_saved = invoice, on_saved
        self.title("Pay Pending Invoice")
        self.minsize(360, 300)
        self.resizable(True, True)
        self.transient(parent.winfo_toplevel())
        self.grab_set()
        self.amount_var = tk.StringVar(value=f"{invoice['balance']:.2f}")
        self.method_var = tk.StringVar(value="Account Pay")
        form = ttk.Frame(self, padding=16)
        form.pack(fill="both", expand=True)
        for label, value in (("Vendor", invoice["supplier"]), ("Invoice No.", invoice["invoice_no"]), ("Pending Balance", f"{invoice['balance']:,.2f}")):
            ttk.Label(form, text=f"{label}:").pack(anchor="w", pady=(0, 2))
            ttk.Label(form, text=value, font=("Helvetica", 10, "bold")).pack(anchor="w", pady=(0, 8))
        ttk.Label(form, text="Payment Amount:").pack(anchor="w")
        ttk.Entry(form, textvariable=self.amount_var, width=28).pack(anchor="w", pady=(2, 8))
        ttk.Label(form, text="Payment Method:").pack(anchor="w")
        ttk.Combobox(form, textvariable=self.method_var, values=("Hard Cash", "Account Pay", "Cheque"), state="readonly", width=25).pack(anchor="w", pady=(2, 12))
        actions = ttk.Frame(form)
        actions.pack(fill="x")
        ttk.Button(actions, text="Cancel", command=self.destroy).pack(side="right")
        ttk.Button(actions, text="Save Payment", command=self.save).pack(side="right", padx=(0, 8))

    def save(self):
        try:
            amount = float(self.amount_var.get())
            if amount <= 0 or amount > self.invoice["balance"] + 0.0001:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Payment Amount", "Enter an amount greater than zero and not more than the pending balance.", parent=self)
            return
        ensure_payments_file()
        with open(PAYMENTS_CSV, "a", newline="", encoding="utf-8") as file:
            csv.DictWriter(file, fieldnames=PAYMENT_HEADERS).writerow({"payment_date": date.today().isoformat(), "supplier": self.invoice["supplier"], "invoice_no": self.invoice["invoice_no"], "amount": f"{amount:.2f}", "payment_method": self.method_var.get()})
        self.on_saved()
        self.destroy()
        messagebox.showinfo("Payment Saved", "Payment saved and pending balance updated.", parent=self.master)


def load_product_catalog():
    """Load selectable existing products from the maintained CSV price list."""
    if not os.path.exists(PRODUCT_CATALOG_CSV):
        return {}
    catalog = {}
    with open(PRODUCT_CATALOG_CSV, newline="", encoding="utf-8-sig") as file:
        for row in csv.DictReader(file):
            item = (row.get("Item Name") or "").strip()
            if not item:
                continue
            capacity = (row.get("CAPACITY") or "").strip()
            unit = (row.get("Unit") or "").strip()
            supplier = (row.get("Supplier") or "").strip()
            label = f"{item} | {capacity} {unit} | {supplier}".strip()
            catalog[label] = row
    return catalog


def safe_file_part(value):
    """Make a vendor or cheque number safe for a Windows file name."""
    return re.sub(r'[<>:"/\\|?*]+', "_", value).strip(" ._") or "unknown"


def load_vendor_names():
    """Vendor selection is deliberately limited to the registered vendor master."""
    from vendor_registration import VENDOR_CSV, ensure_vendor_file
    ensure_vendor_file()
    with open(VENDOR_CSV, newline="", encoding="utf-8-sig") as file:
        return sorted({(row.get("vendor_legal_name") or "").strip() for row in csv.DictReader(file) if (row.get("vendor_legal_name") or "").strip()}, key=str.casefold)


class ChequeDetailsWindow(ttk.Frame):
    """Search the saved cheque register by vendor, invoice, or cheque number."""

    def __init__(self, parent, on_complete=None):
        super().__init__(parent)
        self.on_complete = on_complete or (lambda: None)
        self.search_var = tk.StringVar()
        self.start_date_var = tk.StringVar()
        self.end_date_var = tk.StringVar()
        self.filtered_records = []
        self._build()
        self.load_cheques()

    def _build(self):
        top = ttk.Frame(self, padding=12)
        top.pack(fill="x")
        ttk.Label(top, text="Search Vendor / Invoice / Cheque No.:").pack(side="left")
        search = ttk.Entry(top, textvariable=self.search_var, width=40)
        search.pack(side="left", padx=8)
        search.bind("<KeyRelease>", lambda _event: self.load_cheques())
        ttk.Label(top, text="Cheque Date From:").pack(side="left", padx=(12, 5))
        DateEntry(top, textvariable=self.start_date_var, width=12, date_pattern="yyyy-mm-dd").pack(side="left")
        ttk.Label(top, text="To:").pack(side="left", padx=(8, 5))
        DateEntry(top, textvariable=self.end_date_var, width=12, date_pattern="yyyy-mm-dd").pack(side="left")
        ttk.Button(top, text="Apply Dates", command=self.apply_date_filter).pack(side="left", padx=8)
        ttk.Button(top, text="Export Excel", command=self.export_excel).pack(side="right")
        ttk.Button(top, text="Open Selected Photo", command=self.open_selected_photo).pack(side="right")
        ttk.Button(top, text="Back", command=self.on_complete).pack(side="right", padx=(0, 8))
        box = ttk.Frame(self, padding=(12, 0, 12, 12))
        box.pack(fill="both", expand=True)
        columns = ("supplier", "invoice", "date", "department", "cheque_date", "cheque_number", "photo")
        headings = ("Vendor", "Invoice", "Entry Date", "Department", "Cheque Date", "Cheque Number", "Photo File")
        self.tree = ttk.Treeview(box, columns=columns, show="headings", selectmode="browse")
        for column, heading in zip(columns, headings):
            self.tree.heading(column, text=heading)
            self.tree.column(column, width=160 if column in ("supplier", "photo") else 110, anchor="w")
        scroll = ttk.Scrollbar(box, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

    def load_cheques(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.filtered_records = []
        query = self.search_var.get().strip().casefold()
        start_date = self.start_date_var.get().strip()
        end_date = self.end_date_var.get().strip()
        if not os.path.exists(INWARD_CSV):
            return
        with open(INWARD_CSV, newline="", encoding="utf-8") as file:
            for index, row in enumerate(csv.DictReader(file)):
                if row.get("payment_method") != "Cheque":
                    continue
                searchable = " ".join((row.get("supplier", ""), row.get("invoice_no", ""), row.get("cheque_number", ""))).casefold()
                if query and query not in searchable:
                    continue
                cheque_date = row.get("cheque_date", "")
                if start_date and cheque_date < start_date:
                    continue
                if end_date and cheque_date > end_date:
                    continue
                self.filtered_records.append(row)
                self.tree.insert("", "end", iid=str(index), values=(row.get("supplier", ""), row.get("invoice_no", ""), row.get("date", ""), row.get("department", ""), row.get("cheque_date", ""), row.get("cheque_number", ""), row.get("cheque_photo", "")))

    def apply_date_filter(self):
        start_date = self.start_date_var.get().strip()
        end_date = self.end_date_var.get().strip()
        try:
            if start_date:
                datetime.strptime(start_date, "%Y-%m-%d")
            if end_date:
                datetime.strptime(end_date, "%Y-%m-%d")
            if start_date and end_date and start_date > end_date:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Date Range", "Use YYYY-MM-DD dates, with the start date before the end date.", parent=self)
            return False
        self.load_cheques()
        return True

    def export_excel(self):
        if not self.apply_date_filter():
            return
        if not self.filtered_records:
            messagebox.showwarning("No Cheques", "There are no cheque records in the selected date range.", parent=self)
            return
        start_date = self.start_date_var.get().strip() or "all"
        end_date = self.end_date_var.get().strip() or "all"
        path = filedialog.asksaveasfilename(
            parent=self, title="Save Cheque Details Excel",
            defaultextension=".xlsx", initialfile=f"cheque_details_{start_date}_to_{end_date}.xlsx",
            filetypes=(("Excel workbook", "*.xlsx"),),
        )
        if not path:
            return
        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Cheque Details"
            headings = ("Vendor", "Invoice No.", "Entry Date", "Department", "Cheque Date", "Cheque Number", "Payment Status", "Cheque Photo Path")
            sheet.append(headings)
            for cell in sheet[1]:
                cell.font = cell.font.copy(bold=True)
            for row in self.filtered_records:
                sheet.append((row.get("supplier", ""), row.get("invoice_no", ""), row.get("date", ""), row.get("department", ""), row.get("cheque_date", ""), row.get("cheque_number", ""), row.get("payment_status", ""), row.get("cheque_photo", "")))
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
            workbook.save(path)
            messagebox.showinfo("Exported", f"Cheque details Excel file saved successfully.\n\n{path}", parent=self)
        except Exception as error:
            messagebox.showerror("Export Failed", f"Could not export the Excel file:\n{error}", parent=self)

    def open_selected_photo(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Select Cheque", "Select a cheque record first.", parent=self)
            return
        photo_path = self.tree.item(selected[0], "values")[6]
        if not photo_path or not os.path.exists(photo_path):
            messagebox.showwarning("Photo Not Found", "The cheque photo file is not available.", parent=self)
            return
        os.startfile(photo_path)


class InwardEntryDialog(ttk.Frame):
    """Records an invoice header and multiple products under the invoice."""

    def __init__(self, parent, on_saved, on_cancel=None, on_register_vendor=None):
        super().__init__(parent)
        self.on_saved, self.items = on_saved, []
        self.on_cancel = on_cancel or on_saved
        self.on_register_vendor = on_register_vendor
        self.header = {"date": tk.StringVar(value=date.today().isoformat()), "supplier": tk.StringVar(), "invoice_no": tk.StringVar(), "department": tk.StringVar(), "payment_status": tk.StringVar(value="Pending"), "payment_method": tk.StringVar(value="Hard Cash"), "cheque_date": tk.StringVar(), "cheque_number": tk.StringVar(), "cheque_photo": tk.StringVar()}
        self.vendor_names = load_vendor_names()
        self.item = {"item_description": tk.StringVar(), "quantity": tk.StringVar(value="1"), "rate": tk.StringVar(value="0")}
        self.catalog = load_product_catalog()
        self.catalog_labels = list(self.catalog)
        self.catalog_var = tk.StringVar()
        self.total_var = tk.StringVar(value="0.00")
        self._build()

    def _build(self):
        form = ttk.Frame(self, padding=16)
        form.pack(fill="both", expand=True)
        form.columnconfigure(0, weight=1)
        form.rowconfigure(3, weight=1)

        details = ttk.LabelFrame(form, text=" Invoice Details ", padding=10)
        details.grid(row=0, column=0, sticky="ew")
        for column in (1, 3):
            details.columnconfigure(column, weight=1)
        for index, (label, key) in enumerate((("Date (YYYY-MM-DD)", "date"), ("Supplier", "supplier"), ("Invoice / Challan No.", "invoice_no"), ("Department", "department"))):
            row, group = divmod(index, 2)
            column = group * 2
            ttk.Label(details, text=f"{label}:").grid(row=row, column=column, sticky="w", padx=(0, 8), pady=5)
            if key == "department":
                widget = ttk.Combobox(details, textvariable=self.header[key], values=DEPARTMENTS, state="readonly")
            elif key == "supplier":
                supplier_frame = ttk.Frame(details)
                supplier_frame.grid(row=row, column=column + 1, sticky="ew", padx=(0, 16), pady=5)
                supplier_frame.columnconfigure(0, weight=1)
                self.supplier_combo = ttk.Combobox(supplier_frame, textvariable=self.header[key], values=self.vendor_names, state="readonly")
                self.supplier_combo.grid(row=0, column=0, sticky="ew")
                ttk.Button(supplier_frame, text="+ Add Vendor", command=self.add_vendor).grid(row=0, column=1, padx=(6, 0))
                continue
            elif key == "date":
                widget = DateEntry(details, textvariable=self.header[key], date_pattern="yyyy-mm-dd")
            else:
                widget = ttk.Entry(details, textvariable=self.header[key])
            widget.grid(row=row, column=column + 1, sticky="ew", padx=(0, 16), pady=5)

        product = ttk.LabelFrame(form, text=" Add Product to This Invoice ", padding=10)
        product.grid(row=1, column=0, sticky="ew", pady=(12, 0))
        product.columnconfigure(1, weight=1)
        ttk.Label(product, text="Product / Material:").grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Entry(product, textvariable=self.item["item_description"]).grid(row=0, column=1, sticky="ew")
        ttk.Label(product, text="Qty:").grid(row=0, column=2, padx=(12, 5))
        ttk.Entry(product, textvariable=self.item["quantity"], width=9).grid(row=0, column=3)
        ttk.Label(product, text="Rate:").grid(row=0, column=4, padx=(12, 5))
        ttk.Entry(product, textvariable=self.item["rate"], width=12).grid(row=0, column=5)
        ttk.Button(product, text="Add Product", command=self.add_product).grid(row=0, column=6, padx=(12, 0))
        ttk.Label(product, text="Existing Product:").grid(row=1, column=0, sticky="w", pady=(10, 0))
        self.catalog_combo = ttk.Combobox(product, textvariable=self.catalog_var, values=self.catalog_labels[:5], width=70, height=5)
        self.catalog_combo.grid(row=1, column=1, columnspan=4, sticky="ew", pady=(10, 0))
        self.catalog_combo.bind("<<ComboboxSelected>>", self.use_existing_product)
        self.catalog_combo.bind("<KeyRelease>", self.filter_existing_products)
        ttk.Button(product, text="Use Selected Product", command=self.use_existing_product).grid(row=1, column=5, columnspan=2, padx=(12, 0), pady=(10, 0))
        self.matches_list = tk.Listbox(product, height=5, exportselection=False)
        self.matches_list.grid(row=2, column=1, columnspan=4, sticky="ew", pady=(4, 0))
        self.matches_list.grid_remove()
        self.matches_list.bind("<<ListboxSelect>>", self.select_product_match)

        ttk.Label(form, text="Products in this invoice:", font=("Helvetica", 10, "bold")).grid(row=2, column=0, sticky="w", pady=(12, 4))
        box = ttk.Frame(form)
        box.grid(row=3, column=0, sticky="nsew")
        columns = ("item", "quantity", "rate", "amount")
        self.items_tree = ttk.Treeview(box, columns=columns, show="headings", height=8)
        for col, heading, width in (("item", "Product / Material", 430), ("quantity", "Qty", 90), ("rate", "Rate", 110), ("amount", "Amount", 120)):
            self.items_tree.heading(col, text=heading)
            self.items_tree.column(col, width=width, anchor="w" if col == "item" else "e")
        scroll = ttk.Scrollbar(box, orient="vertical", command=self.items_tree.yview)
        self.items_tree.configure(yscrollcommand=scroll.set)
        self.items_tree.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

        actions = ttk.Frame(form)
        actions.grid(row=4, column=0, sticky="ew", pady=(7, 10))
        ttk.Button(actions, text="Remove Selected Product", command=self.remove_selected_product).pack(side="left")
        ttk.Label(actions, textvariable=self.total_var, font=("Helvetica", 11, "bold")).pack(side="right")
        ttk.Label(actions, text="Invoice Total:", font=("Helvetica", 10, "bold")).pack(side="right", padx=(0, 6))

        footer = ttk.LabelFrame(form, text=" Payment and Notes ", padding=10)
        footer.grid(row=5, column=0, sticky="ew")
        footer.columnconfigure(3, weight=1)
        ttk.Label(footer, text="Payment Status:").grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Combobox(footer, textvariable=self.header["payment_status"], values=("Pending", "Paid", "Partial"), state="readonly", width=14).grid(row=0, column=1, sticky="w")
        ttk.Label(footer, text="Payment Method:").grid(row=0, column=2, sticky="w", padx=(18, 8))
        payment_combo = ttk.Combobox(footer, textvariable=self.header["payment_method"], values=("Hard Cash", "Account Pay", "Cheque"), state="readonly", width=16)
        payment_combo.grid(row=0, column=3, sticky="w")
        payment_combo.bind("<<ComboboxSelected>>", self.toggle_cheque_fields)
        self.cheque_frame = ttk.Frame(footer)
        self.cheque_frame.grid(row=1, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        self.cheque_frame.columnconfigure(5, weight=1)
        ttk.Label(self.cheque_frame, text="Cheque Date:").grid(row=0, column=0, sticky="w", padx=(0, 6))
        DateEntry(self.cheque_frame, textvariable=self.header["cheque_date"], width=14, date_pattern="yyyy-mm-dd").grid(row=0, column=1, sticky="w")
        ttk.Label(self.cheque_frame, text="Cheque Number:").grid(row=0, column=2, sticky="w", padx=(16, 6))
        ttk.Entry(self.cheque_frame, textvariable=self.header["cheque_number"], width=18).grid(row=0, column=3, sticky="w")
        ttk.Label(self.cheque_frame, text="Cheque Photo:").grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(self.cheque_frame, textvariable=self.header["cheque_photo"], state="readonly").grid(row=1, column=1, columnspan=4, sticky="ew", pady=(8, 0))
        ttk.Button(self.cheque_frame, text="Attach Photo", command=self.choose_cheque_photo).grid(row=1, column=5, padx=(8, 0), pady=(8, 0))
        ttk.Button(self.cheque_frame, text="Capture from Camera", command=self.capture_cheque_photo).grid(row=1, column=6, padx=(8, 0), pady=(8, 0))
        ttk.Label(footer, text="Notes:").grid(row=2, column=0, sticky="nw", pady=(10, 0))
        self.notes = tk.Text(footer, height=3)
        self.notes.grid(row=2, column=1, columnspan=3, sticky="ew", pady=(10, 0))
        buttons = ttk.Frame(footer)
        buttons.grid(row=3, column=0, columnspan=4, sticky="e", pady=(10, 0))
        ttk.Button(buttons, text="Cancel", command=self.on_cancel).pack(side="right")
        ttk.Button(buttons, text="Save Invoice", command=self.save).pack(side="right", padx=(0, 8))
        self.toggle_cheque_fields()

    def add_vendor(self):
        if self.on_register_vendor:
            self.on_register_vendor()

    def vendor_added(self, vendor_name):
        self.vendor_names = load_vendor_names()
        self.supplier_combo["values"] = self.vendor_names
        self.header["supplier"].set(vendor_name)

    def add_product(self):
        product = self.item["item_description"].get().strip()
        try:
            quantity, rate = float(self.item["quantity"].get()), float(self.item["rate"].get())
            if quantity <= 0 or rate < 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Invalid Product", "Enter a product, a quantity greater than zero, and a valid rate.", parent=self)
            return
        if not product:
            messagebox.showwarning("Missing Product", "Enter a product or material description.", parent=self)
            return
        self.items.append({"item_description": product, "quantity": quantity, "rate": rate, "amount": quantity * rate})
        self.item["item_description"].set("")
        self.item["quantity"].set("1")
        self.item["rate"].set("0")
        self.refresh_items()

    def use_existing_product(self, _event=None):
        product = self.catalog.get(self.catalog_var.get())
        if not product:
            messagebox.showwarning("Select Product", "Select an existing product from the list first.", parent=self)
            return
        capacity = " ".join(part for part in (product.get("CAPACITY", "").strip(), product.get("Unit", "").strip()) if part)
        self.item["item_description"].set(f"{product.get('Item Name', '').strip()} {capacity}".strip())
        price = (product.get("DP") or product.get("List Price") or "").replace(",", "").strip()
        if price:
            self.item["rate"].set(price)

    def filter_existing_products(self, event=None):
        """Show up to five matching products while the user types."""
        if event and event.keysym in ("Up", "Down", "Return", "Escape", "Tab"):
            return
        search = self.catalog_var.get().strip().casefold()
        matches = [label for label in self.catalog_labels if search in label.casefold()][:5]
        self.catalog_combo["values"] = matches
        if matches and search:
            self.matches_list.delete(0, tk.END)
            for label in matches:
                self.matches_list.insert(tk.END, label)
            self.matches_list.grid()
        else:
            self.matches_list.grid_remove()

    def select_product_match(self, _event=None):
        selected = self.matches_list.curselection()
        if not selected:
            return
        self.catalog_var.set(self.matches_list.get(selected[0]))
        self.matches_list.grid_remove()
        self.use_existing_product()

    def toggle_cheque_fields(self, _event=None):
        if self.header["payment_method"].get() == "Cheque":
            self.cheque_frame.grid()
        else:
            self.cheque_frame.grid_remove()

    def choose_cheque_photo(self):
        path = filedialog.askopenfilename(
            parent=self, title="Attach Cheque Photo",
            filetypes=(("Image files", "*.png *.jpg *.jpeg *.bmp"), ("All files", "*.*")),
        )
        if path:
            self.header["cheque_photo"].set(path)

    def capture_cheque_photo(self):
        """Capture a cheque image from the default PC camera using OpenCV."""
        try:
            import cv2
        except ImportError:
            messagebox.showerror("Camera Unavailable", "Camera support is not installed. Please install OpenCV first.", parent=self)
            return
        camera = cv2.VideoCapture(0)
        if not camera.isOpened():
            camera.release()
            messagebox.showerror("Camera Unavailable", "Could not open the PC camera. Check that the camera is connected and Windows camera permission is enabled.", parent=self)
            return
        captured_path = ""
        try:
            while True:
                success, frame = camera.read()
                if not success:
                    messagebox.showerror("Camera Error", "Could not receive an image from the camera.", parent=self)
                    break
                cv2.putText(frame, "Press SPACE to capture cheque | ESC to cancel", (12, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (0, 255, 0), 2)
                cv2.imshow("Capture Cheque Photo", frame)
                key = cv2.waitKey(1) & 0xFF
                if key == 32:
                    capture_dir = os.path.join(CHEQUE_PHOTOS_DIR, "camera_captures")
                    os.makedirs(capture_dir, exist_ok=True)
                    captured_path = os.path.join(capture_dir, f"captured_cheque_{datetime.now():%Y%m%d_%H%M%S_%f}.png")
                    cv2.imwrite(captured_path, frame)
                    break
                if key == 27:
                    break
        finally:
            camera.release()
            cv2.destroyAllWindows()
        if captured_path:
            self.header["cheque_photo"].set(captured_path)
            messagebox.showinfo("Cheque Captured", "Cheque photo captured and attached to this invoice.", parent=self)

    def remove_selected_product(self):
        selected = self.items_tree.selection()
        if not selected:
            messagebox.showwarning("Select Product", "Select a product line to remove.", parent=self)
            return
        del self.items[int(selected[0])]
        self.refresh_items()

    def refresh_items(self):
        for row_id in self.items_tree.get_children():
            self.items_tree.delete(row_id)
        total = 0.0
        for index, item in enumerate(self.items):
            total += item["amount"]
            self.items_tree.insert("", "end", iid=str(index), values=(item["item_description"], f"{item['quantity']:g}", f"{item['rate']:.2f}", f"{item['amount']:.2f}"))
        self.total_var.set(f"{total:,.2f}")

    def save(self):
        if any(not self.header[key].get().strip() for key in ("supplier", "invoice_no", "department")):
            messagebox.showwarning("Missing Details", "Supplier, invoice/challan number, and department are required.", parent=self)
            return
        if not self.items:
            messagebox.showwarning("No Products", "Add at least one product to this invoice.", parent=self)
            return
        cheque_photo = ""
        if self.header["payment_method"].get() == "Cheque":
            cheque_date = self.header["cheque_date"].get().strip()
            cheque_number = self.header["cheque_number"].get().strip()
            photo_path = self.header["cheque_photo"].get().strip()
            if not cheque_date or not cheque_number or not photo_path:
                messagebox.showwarning("Cheque Details", "Cheque date, cheque number, and cheque photo are required.", parent=self)
                return
            if not os.path.exists(photo_path):
                messagebox.showwarning("Cheque Photo", "The selected cheque photo could not be found.", parent=self)
                return
            try:
                vendor_folder = os.path.join(CHEQUE_PHOTOS_DIR, safe_file_part(self.header["supplier"].get()))
                os.makedirs(vendor_folder, exist_ok=True)
                extension = os.path.splitext(photo_path)[1]
                cheque_photo = os.path.join(vendor_folder, f"{safe_file_part(self.header['supplier'].get())}_{safe_file_part(cheque_number)}_{cheque_date}_{datetime.now():%Y%m%d_%H%M%S}{extension}")
                shutil.copy2(photo_path, cheque_photo)
            except OSError as error:
                messagebox.showerror("Cheque Photo", f"Could not save the cheque photo:\n{error}", parent=self)
                return
        ensure_inward_file()
        notes = self.notes.get("1.0", "end-1c").strip()
        with open(INWARD_CSV, "a", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=INWARD_HEADERS)
            for item in self.items:
                writer.writerow({"date": self.header["date"].get().strip(), "supplier": self.header["supplier"].get().strip(), "invoice_no": self.header["invoice_no"].get().strip(), "department": self.header["department"].get().strip(), "item_description": item["item_description"], "quantity": f"{item['quantity']:g}", "rate": f"{item['rate']:.2f}", "amount": f"{item['amount']:.2f}", "payment_status": self.header["payment_status"].get(), "payment_method": self.header["payment_method"].get(), "cheque_date": self.header["cheque_date"].get().strip() if cheque_photo else "", "cheque_number": self.header["cheque_number"].get().strip() if cheque_photo else "", "cheque_photo": cheque_photo, "notes": notes})
        try:
            output_dir = self.generate_invoice_files(notes)
        except Exception as error:
            messagebox.showerror("Invoice Saved", f"The inward invoice was saved, but its PDF and Excel files could not be generated:\n{error}", parent=self)
            self.on_saved()
            return
        messagebox.showinfo("Saved", f"Invoice saved with {len(self.items)} product line(s).\n\nPDF and Excel files created in:\n{output_dir}", parent=self)
        self.on_saved()

    def generate_invoice_files(self, notes):
        """Create one Excel workbook and one printable PDF for this new invoice."""
        invoice_no = self.header["invoice_no"].get().strip()
        base_name = safe_file_part(f"Inward_Invoice_{invoice_no}_{datetime.now():%Y%m%d_%H%M%S}")
        output_dir = os.path.join(INWARD_OUTPUT_DIR, base_name)
        os.makedirs(output_dir, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inward Invoice"
        sheet.append(("Date", self.header["date"].get().strip()))
        sheet.append(("Vendor", self.header["supplier"].get().strip()))
        sheet.append(("Invoice / Challan No.", invoice_no))
        sheet.append(("Department", self.header["department"].get().strip()))
        sheet.append(())
        sheet.append(("No.", "Product / Material", "Qty", "Rate", "Amount"))
        for index, item in enumerate(self.items, 1):
            sheet.append((index, item["item_description"], item["quantity"], item["rate"], item["amount"]))
        sheet.append(("", "", "", "Total", sum(item["amount"] for item in self.items)))
        sheet.append(())
        sheet.append(("Notes", notes))
        for cell in sheet[6]:
            cell.font = cell.font.copy(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
        workbook.save(os.path.join(output_dir, f"{base_name}.xlsx"))
        styles = getSampleStyleSheet()
        pdf = SimpleDocTemplate(os.path.join(output_dir, f"{base_name}.pdf"), pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm, topMargin=14 * mm)
        story = [Paragraph("Inward Invoice", styles["Title"]), Spacer(1, 5 * mm), Paragraph(
            f"<b>Date:</b> {escape(self.header['date'].get())}<br/><b>Vendor:</b> {escape(self.header['supplier'].get())}<br/><b>Invoice / Challan No.:</b> {escape(invoice_no)}<br/><b>Department:</b> {escape(self.header['department'].get())}", styles["Normal"]), Spacer(1, 5 * mm)]
        table_data = [["No.", "Product / Material", "Qty", "Rate", "Amount"]]
        for index, item in enumerate(self.items, 1):
            table_data.append((str(index), item["item_description"], f"{item['quantity']:g}", f"{item['rate']:.2f}", f"{item['amount']:.2f}"))
        table_data.append(("", "", "", "Total", f"{sum(item['amount'] for item in self.items):.2f}"))
        table = Table(table_data, colWidths=(12 * mm, 88 * mm, 18 * mm, 28 * mm, 28 * mm))
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("ALIGN", (2, 1), (-1, -1), "RIGHT"), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (3, -1), (-1, -1), "Helvetica-Bold")]))
        story.extend((table, Spacer(1, 5 * mm), Paragraph(f"<b>Notes:</b> {escape(notes or '-')}", styles["Normal"])))
        pdf.build(story)
        return output_dir


class SalesIncomeView(ttk.Frame):
    """Sales / Income register — track customer invoices, payments, and pending amounts."""

    COLS = ("date", "customer_name", "invoice_number", "product_service",
            "invoice_amount", "amount_received", "pending_amount", "notes")
    HEADINGS = ("Date", "Customer", "Invoice No.", "Product / Service",
                "Invoice Amount", "Received", "Pending", "Notes")

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_sales_income_file()
        self._customer_names = load_customer_names()
        self._editing_index = None          # row index being edited (None = new)
        self._build()
        self.load_entries()

    # ------------------------------------------------------------------ UI --
    def _build(self):
        # ── Top header bar ────────────────────────────────────────────────
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="💰 Sales / Income", font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel", command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh", command=self.load_entries).pack(side="right", padx=(0, 6))

        # ── Summary cards ─────────────────────────────────────────────────
        summary = ttk.LabelFrame(self, text=" Summary ", padding=10)
        summary.pack(fill="x", pady=(0, 10))
        self._total_invoiced_var = tk.StringVar(value="₹0.00")
        self._total_received_var = tk.StringVar(value="₹0.00")
        self._total_pending_var  = tk.StringVar(value="₹0.00")
        for col, (lbl, var) in enumerate((
            ("Total Invoiced",  self._total_invoiced_var),
            ("Total Received",  self._total_received_var),
            ("Total Pending",   self._total_pending_var),
        )):
            card = ttk.Frame(summary, padding=(12, 4))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            ttk.Label(card, textvariable=var, font=("Helvetica", 14, "bold")).pack(anchor="w")
            summary.columnconfigure(col, weight=1)

        # ── Entry form ────────────────────────────────────────────────────
        form_box = ttk.LabelFrame(self, text=" Add / Edit Entry ", padding=10)
        form_box.pack(fill="x", pady=(0, 10))

        # Row 0: Date | Customer | Invoice No. | Product/Service
        ttk.Label(form_box, text="Date:").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self._ent_date = DateEntry(form_box, date_pattern="yyyy-mm-dd", width=14)
        self._ent_date.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Customer:").grid(row=0, column=2, sticky="w", padx=4, pady=4)
        self._customer_var = tk.StringVar()
        self._cmb_customer = ttk.Combobox(form_box, textvariable=self._customer_var,
                                          values=self._customer_names, width=28)
        self._cmb_customer.grid(row=0, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Invoice No.:").grid(row=0, column=4, sticky="w", padx=4, pady=4)
        self._ent_invoice = ttk.Entry(form_box, width=16)
        self._ent_invoice.grid(row=0, column=5, sticky="ew", padx=4, pady=4)

        # Row 1: Product/Service | Invoice Amount | Amount Received | Notes
        ttk.Label(form_box, text="Product / Service:").grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self._ent_product = ttk.Entry(form_box, width=28)
        self._ent_product.grid(row=1, column=1, columnspan=2, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Invoice Amount (₹):").grid(row=1, column=3, sticky="w", padx=4, pady=4)
        self._ent_invoice_amt = ttk.Entry(form_box, width=14)
        self._ent_invoice_amt.grid(row=1, column=4, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Amount Received (₹):").grid(row=1, column=5, sticky="w", padx=4, pady=4)
        self._ent_received = ttk.Entry(form_box, width=14)
        self._ent_received.grid(row=1, column=6, sticky="ew", padx=4, pady=4)

        # Row 2: Notes
        ttk.Label(form_box, text="Notes:").grid(row=2, column=0, sticky="w", padx=4, pady=4)
        self._ent_notes = ttk.Entry(form_box, width=60)
        self._ent_notes.grid(row=2, column=1, columnspan=5, sticky="ew", padx=4, pady=4)

        # Action buttons
        btn_row = ttk.Frame(form_box)
        btn_row.grid(row=2, column=6, sticky="e", padx=4, pady=4)
        ttk.Button(btn_row, text="💾 Save", command=self.save_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="🗑 Delete", command=self.delete_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="✖ Clear", command=self.clear_form).pack(side="left")

        for c in range(7):
            form_box.columnconfigure(c, weight=1)

        # ── Treeview ──────────────────────────────────────────────────────
        table_box = ttk.LabelFrame(self, text=" Sales / Income Register ", padding=8)
        table_box.pack(fill="both", expand=True)

        self._tree = ttk.Treeview(table_box, columns=self.COLS, show="headings", selectmode="browse")
        col_widths = {
            "date": 95, "customer_name": 190, "invoice_number": 110,
            "product_service": 210, "invoice_amount": 120,
            "amount_received": 110, "pending_amount": 100, "notes": 200,
        }
        right_align = {"invoice_amount", "amount_received", "pending_amount"}
        for col, heading in zip(self.COLS, self.HEADINGS):
            self._tree.heading(col, text=heading,
                               command=lambda c=col: self._sort_by(c))
            self._tree.column(col, width=col_widths.get(col, 110),
                              anchor="e" if col in right_align else "w")

        vsb = ttk.Scrollbar(table_box, orient="vertical", command=self._tree.yview)
        hsb = ttk.Scrollbar(table_box, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_box.rowconfigure(0, weight=1)
        table_box.columnconfigure(0, weight=1)

        self._tree.bind("<<TreeviewSelect>>", self._on_row_select)

        # Track sort direction
        self._sort_reverse = {c: False for c in self.COLS}

    # -------------------------------------------------------- data helpers --
    def _read_all(self):
        ensure_sales_income_file()
        with open(SALES_INCOME_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_all(self, rows):
        with open(SALES_INCOME_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=SALES_INCOME_HEADERS)
            writer.writeheader()
            writer.writerows(rows)

    @staticmethod
    def _calc_pending(invoice_amt: str, received: str) -> float:
        try:
            inv = float(str(invoice_amt).replace(",", "") or 0)
        except ValueError:
            inv = 0.0
        try:
            rec = float(str(received).replace(",", "") or 0)
        except ValueError:
            rec = 0.0
        return max(inv - rec, 0.0)

    # ------------------------------------------------------------ actions --
    def load_entries(self):
        for iid in self._tree.get_children():
            self._tree.delete(iid)

        rows = self._read_all()
        total_inv = total_rec = total_pend = 0.0
        for idx, row in enumerate(rows):
            try:
                inv_amt = float(str(row.get("invoice_amount", "0")).replace(",", "") or 0)
            except ValueError:
                inv_amt = 0.0
            try:
                rec_amt = float(str(row.get("amount_received", "0")).replace(",", "") or 0)
            except ValueError:
                rec_amt = 0.0
            pend = self._calc_pending(inv_amt, rec_amt)

            # Recompute & store pending in memory for display
            row["pending_amount"] = f"{pend:.2f}"
            total_inv  += inv_amt
            total_rec  += rec_amt
            total_pend += pend

            self._tree.insert("", "end", iid=str(idx), values=(
                row.get("date", ""),
                row.get("customer_name", ""),
                row.get("invoice_number", ""),
                row.get("product_service", ""),
                f"{inv_amt:,.2f}",
                f"{rec_amt:,.2f}",
                f"{pend:,.2f}",
                row.get("notes", ""),
            ))

        self._total_invoiced_var.set(f"₹{total_inv:,.2f}")
        self._total_received_var.set(f"₹{total_rec:,.2f}")
        self._total_pending_var.set(f"₹{total_pend:,.2f}")
        self._editing_index = None
        self.clear_form()

    def _on_row_select(self, _event=None):
        sel = self._tree.selection()
        if not sel:
            return
        idx = int(sel[0])
        rows = self._read_all()
        if idx >= len(rows):
            return
        row = rows[idx]
        self._editing_index = idx

        # Populate form
        try:
            self._ent_date.set_date(row.get("date", date.today().isoformat()))
        except Exception:
            pass
        self._customer_var.set(row.get("customer_name", ""))
        self._ent_invoice.delete(0, tk.END)
        self._ent_invoice.insert(0, row.get("invoice_number", ""))
        self._ent_product.delete(0, tk.END)
        self._ent_product.insert(0, row.get("product_service", ""))
        self._ent_invoice_amt.delete(0, tk.END)
        self._ent_invoice_amt.insert(0, row.get("invoice_amount", ""))
        self._ent_received.delete(0, tk.END)
        self._ent_received.insert(0, row.get("amount_received", ""))
        self._ent_notes.delete(0, tk.END)
        self._ent_notes.insert(0, row.get("notes", ""))

    def clear_form(self):
        self._editing_index = None
        try:
            self._ent_date.set_date(date.today())
        except Exception:
            pass
        self._customer_var.set("")
        for w in (self._ent_invoice, self._ent_product,
                  self._ent_invoice_amt, self._ent_received, self._ent_notes):
            w.delete(0, tk.END)
        self._tree.selection_remove(self._tree.selection())

    def save_entry(self):
        customer = self._customer_var.get().strip()
        invoice_no = self._ent_invoice.get().strip()
        product = self._ent_product.get().strip()
        inv_amt_str = self._ent_invoice_amt.get().strip()
        rec_amt_str = self._ent_received.get().strip()

        if not customer:
            messagebox.showwarning("Missing Field", "Customer name is required.", parent=self)
            return
        if not invoice_no:
            messagebox.showwarning("Missing Field", "Invoice number is required.", parent=self)
            return
        try:
            inv_amt = float(inv_amt_str.replace(",", "") or 0)
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Invoice amount must be a number.", parent=self)
            return
        try:
            rec_amt = float(rec_amt_str.replace(",", "") or 0)
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Amount received must be a number.", parent=self)
            return

        pend = max(inv_amt - rec_amt, 0.0)

        new_row = {
            "date":            self._ent_date.get(),
            "customer_name":   customer,
            "invoice_number":  invoice_no,
            "product_service": product,
            "invoice_amount":  f"{inv_amt:.2f}",
            "amount_received": f"{rec_amt:.2f}",
            "pending_amount":  f"{pend:.2f}",
            "notes":           self._ent_notes.get().strip(),
        }

        rows = self._read_all()
        if self._editing_index is not None and self._editing_index < len(rows):
            rows[self._editing_index] = new_row
            msg = "Entry updated successfully."
        else:
            rows.append(new_row)
            msg = "Entry saved successfully."

        self._write_all(rows)
        messagebox.showinfo("Saved", msg, parent=self)
        self.load_entries()

    def delete_entry(self):
        if self._editing_index is None:
            messagebox.showwarning("Select Row", "Select a row to delete.", parent=self)
            return
        rows = self._read_all()
        if self._editing_index >= len(rows):
            return
        row = rows[self._editing_index]
        confirm = messagebox.askyesno(
            "Confirm Delete",
            f"Delete entry for {row.get('customer_name','')} — {row.get('invoice_number','')}?",
            parent=self,
        )
        if confirm:
            rows.pop(self._editing_index)
            self._write_all(rows)
            self.load_entries()

    def _sort_by(self, col):
        rows = self._read_all()
        reverse = self._sort_reverse[col]
        numeric_cols = {"invoice_amount", "amount_received", "pending_amount"}

        def key(r):
            v = r.get(col, "")
            if col in numeric_cols:
                try:
                    return float(str(v).replace(",", "") or 0)
                except ValueError:
                    return 0.0
            return v.lower()

        rows.sort(key=key, reverse=reverse)
        self._sort_reverse[col] = not reverse
        self._write_all(rows)
        self.load_entries()

    def export_excel(self):
        rows = self._read_all()
        if not rows:
            messagebox.showinfo("No Data", "There are no entries to export.", parent=self)
            return

        path = filedialog.asksaveasfilename(
            parent=self,
            title="Save Sales / Income Report",
            defaultextension=".xlsx",
            initialfile=f"Sales_Income_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Sales Income"

            # Header row
            ws.append(list(self.HEADINGS))
            header_row = ws[1]
            from openpyxl.styles import Font, PatternFill, Alignment
            for cell in header_row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2E4057")
                cell.alignment = Alignment(horizontal="center")

            # Data rows
            total_inv = total_rec = total_pend = 0.0
            for row in rows:
                try:
                    inv = float(str(row.get("invoice_amount", "0")).replace(",", "") or 0)
                except ValueError:
                    inv = 0.0
                try:
                    rec = float(str(row.get("amount_received", "0")).replace(",", "") or 0)
                except ValueError:
                    rec = 0.0
                pend = max(inv - rec, 0.0)
                total_inv += inv; total_rec += rec; total_pend += pend

                ws.append([
                    row.get("date", ""),
                    row.get("customer_name", ""),
                    row.get("invoice_number", ""),
                    row.get("product_service", ""),
                    inv, rec, pend,
                    row.get("notes", ""),
                ])

            # Totals row
            ws.append(["", "", "", "TOTAL", total_inv, total_rec, total_pend, ""])
            total_row = ws[ws.max_row]
            for cell in total_row:
                cell.font = Font(bold=True)

            # Column widths
            for col, width in zip("ABCDEFGH", [12, 30, 16, 35, 16, 16, 14, 30]):
                ws.column_dimensions[col].width = width

            # Number format for amount columns
            for row_cells in ws.iter_rows(min_row=2, min_col=5, max_col=7):
                for cell in row_cells:
                    cell.number_format = '#,##0.00'

            wb.save(path)
            messagebox.showinfo("Exported", f"Excel file saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", f"Could not save Excel file:\n{e}", parent=self)


class FinancialReportsView(ttk.Frame):
    """📊 Financial Reports — auto-calculated P&L, Cash Flow, Balance Sheet, Outstanding."""

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        self._build()
        self.refresh_all()

    # ================================================================ BUILD ==
    def _build(self):
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="📊 Financial Reports",
                  font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="🔄 Refresh All",
                   command=self.refresh_all).pack(side="right")
        ttk.Button(header, text="⬇ Export All to Excel",
                   command=self.export_all).pack(side="right", padx=(0, 6))

        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True)

        # ── Tab frames
        self._tab_dash   = ttk.Frame(nb, padding=16)
        self._tab_pl     = ttk.Frame(nb, padding=16)
        self._tab_cf     = ttk.Frame(nb, padding=16)
        self._tab_bs     = ttk.Frame(nb, padding=16)
        self._tab_out    = ttk.Frame(nb, padding=16)

        nb.add(self._tab_dash, text="  🏠 Dashboard  ")
        nb.add(self._tab_pl,   text="  📊 Profit & Loss  ")
        nb.add(self._tab_cf,   text="  🏦 Cash Flow  ")
        nb.add(self._tab_bs,   text="  📋 Balance Sheet  ")
        nb.add(self._tab_out,  text="  ⏳ Outstanding  ")

        self._build_dashboard()
        self._build_pl()
        self._build_cf()
        self._build_bs()
        self._build_outstanding()

    # ─────────────────────────────────────── TAB 1 — DASHBOARD ──────────────
    def _build_dashboard(self):
        ttk.Label(self._tab_dash, text="Today's Financial Snapshot",
                  font=("Helvetica", 13, "bold")).pack(anchor="w", pady=(0, 12))

        grid = ttk.Frame(self._tab_dash)
        grid.pack(fill="x")

        self._dash_vars = {}
        metrics = [
            ("today_sales",          "💰 Today's Sales",          "navy"),
            ("monthly_sales",        "📈 Monthly Sales",          "navy"),
            ("bank_balance",         "🏦 Bank Balance",           "green"),
            ("cash_balance",         "💵 Cash in Hand",           "green"),
            ("customer_outstanding", "📥 Customer Outstanding",   "red"),
            ("vendor_outstanding",   "📤 Vendor Outstanding",     "red"),
            ("monthly_expense",      "💸 Monthly Expenses",       "darkorange"),
            ("monthly_profit",       "📊 Monthly Profit / Loss",  "purple"),
            ("gst_payable",          "🧾 GST Payable",            "brown"),
        ]
        for idx, (key, label, fg) in enumerate(metrics):
            var = tk.StringVar(value="₹0.00")
            self._dash_vars[key] = var
            row, col = divmod(idx, 3)
            card = ttk.LabelFrame(grid, text=f" {label} ", padding=(14, 6))
            card.grid(row=row, column=col, sticky="ew", padx=8, pady=6)
            lbl = ttk.Label(card, textvariable=var, font=("Helvetica", 18, "bold"))
            lbl.configure(foreground=fg)
            lbl.pack(anchor="w")
            grid.columnconfigure(col, weight=1)

    # ──────────────────────────────────── TAB 2 — PROFIT & LOSS ──────────────
    def _build_pl(self):
        ttk.Label(self._tab_pl, text="Profit & Loss Statement",
                  font=("Helvetica", 13, "bold")).pack(anchor="w", pady=(0, 12))

        # Period filter
        fbar = ttk.Frame(self._tab_pl)
        fbar.pack(fill="x", pady=(0, 10))
        ttk.Label(fbar, text="Period (YYYY-MM, or leave blank for all):").pack(side="left", padx=(0, 6))
        self._pl_period = tk.StringVar()
        ttk.Entry(fbar, textvariable=self._pl_period, width=10).pack(side="left")
        ttk.Button(fbar, text="Apply", command=self._refresh_pl).pack(side="left", padx=4)
        ttk.Button(fbar, text="All",   command=lambda: [self._pl_period.set(""), self._refresh_pl()]).pack(side="left")

        # Table
        self._pl_tree = ttk.Treeview(self._tab_pl,
                                     columns=("category", "description", "amount"),
                                     show="headings", height=22)
        self._pl_tree.heading("category",    text="Category")
        self._pl_tree.heading("description", text="Description")
        self._pl_tree.heading("amount",      text="Amount (₹)")
        self._pl_tree.column("category",    width=160, anchor="w")
        self._pl_tree.column("description", width=300, anchor="w")
        self._pl_tree.column("amount",      width=140, anchor="e")
        self._pl_tree.tag_configure("header",  font=("Helvetica", 10, "bold"), background="#D5E8D4")
        self._pl_tree.tag_configure("total",   font=("Helvetica", 10, "bold"), background="#DAE8FC")
        self._pl_tree.tag_configure("profit",  font=("Helvetica", 10, "bold"), foreground="#27AE60")
        self._pl_tree.tag_configure("loss",    font=("Helvetica", 10, "bold"), foreground="#E74C3C")
        self._pl_tree.tag_configure("expense", foreground="#C0392B")
        self._pl_tree.tag_configure("income",  foreground="#1A5276")
        vsb = ttk.Scrollbar(self._tab_pl, orient="vertical", command=self._pl_tree.yview)
        self._pl_tree.configure(yscrollcommand=vsb.set)
        self._pl_tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

    # ────────────────────────────────────── TAB 3 — CASH FLOW ────────────────
    def _build_cf(self):
        ttk.Label(self._tab_cf, text="Cash Flow Statement",
                  font=("Helvetica", 13, "bold")).pack(anchor="w", pady=(0, 12))

        fbar = ttk.Frame(self._tab_cf)
        fbar.pack(fill="x", pady=(0, 10))
        ttk.Label(fbar, text="Period (YYYY-MM, blank = all):").pack(side="left", padx=(0, 6))
        self._cf_period = tk.StringVar()
        ttk.Entry(fbar, textvariable=self._cf_period, width=10).pack(side="left")
        ttk.Button(fbar, text="Apply", command=self._refresh_cf).pack(side="left", padx=4)
        ttk.Button(fbar, text="All",   command=lambda: [self._cf_period.set(""), self._refresh_cf()]).pack(side="left")

        self._cf_tree = ttk.Treeview(self._tab_cf,
                                     columns=("section", "description", "amount"),
                                     show="headings", height=22)
        self._cf_tree.heading("section",     text="Section")
        self._cf_tree.heading("description", text="Description")
        self._cf_tree.heading("amount",      text="Amount (₹)")
        self._cf_tree.column("section",     width=200, anchor="w")
        self._cf_tree.column("description", width=280, anchor="w")
        self._cf_tree.column("amount",      width=140, anchor="e")
        self._cf_tree.tag_configure("header",  font=("Helvetica", 10, "bold"), background="#D5E8D4")
        self._cf_tree.tag_configure("total",   font=("Helvetica", 10, "bold"), background="#DAE8FC")
        self._cf_tree.tag_configure("in",      foreground="#27AE60")
        self._cf_tree.tag_configure("out",     foreground="#E74C3C")
        vsb2 = ttk.Scrollbar(self._tab_cf, orient="vertical", command=self._cf_tree.yview)
        self._cf_tree.configure(yscrollcommand=vsb2.set)
        self._cf_tree.pack(side="left", fill="both", expand=True)
        vsb2.pack(side="right", fill="y")

    # ──────────────────────────────────── TAB 4 — BALANCE SHEET ──────────────
    def _build_bs(self):
        ttk.Label(self._tab_bs, text="Balance Sheet",
                  font=("Helvetica", 13, "bold")).pack(anchor="w", pady=(0, 12))

        self._bs_tree = ttk.Treeview(self._tab_bs,
                                     columns=("section", "item", "amount"),
                                     show="headings", height=24)
        self._bs_tree.heading("section", text="Section")
        self._bs_tree.heading("item",    text="Item")
        self._bs_tree.heading("amount",  text="Amount (₹)")
        self._bs_tree.column("section", width=180, anchor="w")
        self._bs_tree.column("item",    width=300, anchor="w")
        self._bs_tree.column("amount",  width=140, anchor="e")
        self._bs_tree.tag_configure("header",  font=("Helvetica", 10, "bold"), background="#D5E8D4")
        self._bs_tree.tag_configure("total",   font=("Helvetica", 10, "bold"), background="#DAE8FC")
        self._bs_tree.tag_configure("asset",   foreground="#1A5276")
        self._bs_tree.tag_configure("liab",    foreground="#922B21")
        vsb3 = ttk.Scrollbar(self._tab_bs, orient="vertical", command=self._bs_tree.yview)
        self._bs_tree.configure(yscrollcommand=vsb3.set)
        self._bs_tree.pack(side="left", fill="both", expand=True)
        vsb3.pack(side="right", fill="y")

    # ──────────────────────────────────── TAB 5 — OUTSTANDING ────────────────
    def _build_outstanding(self):
        ttk.Label(self._tab_out, text="Outstanding Report",
                  font=("Helvetica", 13, "bold")).pack(anchor="w", pady=(0, 12))

        # Summary strip
        self._out_summary = ttk.Frame(self._tab_out)
        self._out_summary.pack(fill="x", pady=(0, 10))
        self._out_cust_var   = tk.StringVar(value="₹0.00")
        self._out_vendor_var = tk.StringVar(value="₹0.00")
        for col, (lbl, var, fg) in enumerate((
            ("Total Customer Outstanding", self._out_cust_var,   "red"),
            ("Total Vendor Outstanding",   self._out_vendor_var, "red"),
        )):
            card = ttk.Frame(self._out_summary, padding=(14, 6))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            ttk.Label(card, textvariable=var,
                      font=("Helvetica", 16, "bold"), foreground=fg).pack(anchor="w")
            self._out_summary.columnconfigure(col, weight=1)

        # Two treeviews stacked
        ttk.Label(self._tab_out, text="Customer Pending Payments",
                  font=("Helvetica", 11, "bold")).pack(anchor="w", pady=(8, 2))
        self._cust_tree = self._make_out_tree(self._tab_out,
            ("Customer", "Invoice No.", "Invoice Date", "Due Date",
             "Invoice Amt", "Received", "Pending", "Status"))

        ttk.Label(self._tab_out, text="Vendor Pending Payments",
                  font=("Helvetica", 11, "bold")).pack(anchor="w", pady=(10, 2))
        self._vend_tree = self._make_out_tree(self._tab_out,
            ("Vendor", "Bill No.", "Bill Date", "Due Date",
             "Bill Amt", "Paid", "Pending", "Status"))

    def _make_out_tree(self, parent, headings):
        cols = tuple(h.lower().replace(" ", "_").replace("/","") for h in headings)
        tree = ttk.Treeview(parent, columns=cols, show="headings",
                            height=7, selectmode="browse")
        for col, hdg in zip(cols, headings):
            tree.heading(col, text=hdg)
            tree.column(col, width=140 if "name" in col or "vendor" in col or "customer" in col else 100,
                        anchor="e" if "amt" in col or "received" in col or "paid" in col or "pending" in col else "w")
        tree.tag_configure("overdue", foreground="#E74C3C")
        tree.tag_configure("partial", foreground="#F39C12")
        vsb = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.pack(side="left", fill="x", expand=True, pady=(0, 4))
        vsb.pack(side="right", fill="y", pady=(0, 4))
        return tree

    # ======================================================= DATA HELPERS ==
    @staticmethod
    def _tof(v):
        try: return float(str(v).replace(",", "") or 0)
        except ValueError: return 0.0

    def _safe_read(self, csv_path, headers):
        if not os.path.exists(csv_path):
            return []
        try:
            with open(csv_path, newline="", encoding="utf-8") as f:
                rows = list(csv.DictReader(f))
            return rows
        except Exception:
            return []

    def _filter_by_period(self, rows, date_key, period):
        if not period:
            return rows
        return [r for r in rows if str(r.get(date_key, "")).startswith(period)]

    # ──────────────────────────── aggregate reads ────────────────────────────
    def _get_total_sales(self, period=""):
        rows = self._filter_by_period(
            self._safe_read(SALES_INCOME_CSV, SALES_INCOME_HEADERS), "date", period)
        return sum(self._tof(r.get("invoice_amount", 0)) for r in rows)

    def _get_total_received(self, period=""):
        rows = self._filter_by_period(
            self._safe_read(SALES_INCOME_CSV, SALES_INCOME_HEADERS), "date", period)
        return sum(self._tof(r.get("amount_received", 0)) for r in rows)

    def _get_total_expenses(self, period=""):
        rows = self._filter_by_period(
            self._safe_read(EXPENSES_CSV, EXPENSES_HEADERS), "date", period)
        return sum(self._tof(r.get("amount", 0)) for r in rows)

    def _get_expense_by_category(self, period=""):
        rows = self._filter_by_period(
            self._safe_read(EXPENSES_CSV, EXPENSES_HEADERS), "date", period)
        cats = {}
        for r in rows:
            cat = r.get("category", "Other")
            cats[cat] = cats.get(cat, 0.0) + self._tof(r.get("amount", 0))
        return cats

    def _get_bank_balance(self):
        accounts = self._safe_read(BANK_ACCOUNTS_CSV, BANK_ACCOUNTS_HEADERS)
        total = 0.0
        for acc in accounts:
            try:
                opening = self._tof(acc.get("opening_balance", 0))
                txns = self._safe_read(BANK_TRANSACTIONS_CSV, BANK_TRANSACTIONS_HEADERS)
                credit = {"Deposit", "Cash In"}
                debit  = {"Withdrawal", "Cash Out", "Bank Charge"}
                for t in txns:
                    if t.get("account_name") != acc.get("account_name"):
                        continue
                    amt = self._tof(t.get("amount", 0))
                    if t.get("transaction_type") in credit: opening += amt
                    elif t.get("transaction_type") in debit: opening -= amt
                total += opening
            except Exception:
                pass
        return total

    def _get_cash_in_hand(self):
        txns = self._safe_read(BANK_TRANSACTIONS_CSV, BANK_TRANSACTIONS_HEADERS)
        cash = 0.0
        for t in txns:
            tt = t.get("transaction_type", "")
            amt = self._tof(t.get("amount", 0))
            if tt == "Cash In":    cash += amt
            elif tt == "Cash Out": cash -= amt
        return cash

    def _get_customer_outstanding(self):
        rows = self._safe_read(AR_CSV, AR_HEADERS)
        return sum(self._tof(r.get("pending_amount", 0))
                   for r in rows if r.get("status", "") not in ("Paid",))

    def _get_vendor_outstanding(self):
        rows = self._safe_read(AP_CSV, AP_HEADERS)
        return sum(self._tof(r.get("pending_amount", 0))
                   for r in rows if r.get("status", "") not in ("Paid",))

    def _get_gst_payable(self):
        rows = self._safe_read(GST_CSV, GST_HEADERS)
        sales = sum(self._tof(r.get("total_gst", 0))
                    for r in rows if r.get("entry_type") == "Sales GST")
        itc   = sum(self._tof(r.get("total_gst", 0))
                    for r in rows if "Purchase" in r.get("entry_type", "") or
                    "ITC" in r.get("entry_type", ""))
        return max(sales - itc, 0.0)

    def _get_inventory_value(self):
        rows = self._safe_read(INVENTORY_CSV, INVENTORY_HEADERS)
        return sum(self._tof(r.get("stock_value", 0)) for r in rows)

    def _get_salary_payable(self):
        rows = self._safe_read(SALARY_CSV, SALARY_HEADERS)
        return sum(self._tof(r.get("balance", 0))
                   for r in rows if r.get("status", "") in ("Unpaid", "Partial"))

    # ============================================================= REFRESH ==
    def refresh_all(self):
        self._refresh_dashboard()
        self._refresh_pl()
        self._refresh_cf()
        self._refresh_bs()
        self._refresh_outstanding()

    def _refresh_dashboard(self):
        today   = date.today().isoformat()
        month   = date.today().strftime("%Y-%m")

        # Today's sales
        sales_rows = self._safe_read(SALES_INCOME_CSV, SALES_INCOME_HEADERS)
        today_sales = sum(self._tof(r.get("invoice_amount", 0))
                          for r in sales_rows if r.get("date", "") == today)
        monthly_sales = sum(self._tof(r.get("invoice_amount", 0))
                            for r in sales_rows if r.get("date", "").startswith(month))

        monthly_exp = self._get_total_expenses(month)
        monthly_profit = monthly_sales - monthly_exp

        self._dash_vars["today_sales"].set(f"₹{today_sales:,.2f}")
        self._dash_vars["monthly_sales"].set(f"₹{monthly_sales:,.2f}")
        self._dash_vars["bank_balance"].set(f"₹{self._get_bank_balance():,.2f}")
        self._dash_vars["cash_balance"].set(f"₹{self._get_cash_in_hand():,.2f}")
        self._dash_vars["customer_outstanding"].set(f"₹{self._get_customer_outstanding():,.2f}")
        self._dash_vars["vendor_outstanding"].set(f"₹{self._get_vendor_outstanding():,.2f}")
        self._dash_vars["monthly_expense"].set(f"₹{monthly_exp:,.2f}")
        self._dash_vars["monthly_profit"].set(f"₹{monthly_profit:,.2f}")
        self._dash_vars["gst_payable"].set(f"₹{self._get_gst_payable():,.2f}")

    def _refresh_pl(self):
        for iid in self._pl_tree.get_children():
            self._pl_tree.delete(iid)
        period = self._pl_period.get().strip()

        def row(cat, desc, amt, tag=""):
            self._pl_tree.insert("", "end", values=(cat, desc, f"{amt:,.2f}"), tags=(tag,))

        # ── INCOME ──
        self._pl_tree.insert("", "end", values=("INCOME", "", ""), tags=("header",))
        total_sales = self._get_total_sales(period)
        row("Sales", "Total Sales / Revenue", total_sales, "income")

        received = self._get_total_received(period)
        row("Sales", "Amount Received", received, "income")
        pending_sales = total_sales - received
        row("Sales", "Pending Receivable", pending_sales, "income")
        self._pl_tree.insert("", "end", values=("", "Total Income", f"{total_sales:,.2f}"), tags=("total",))

        # ── EXPENSES ──
        self._pl_tree.insert("", "end", values=("", "", ""), tags=())
        self._pl_tree.insert("", "end", values=("EXPENSES", "", ""), tags=("header",))
        exp_cats = self._get_expense_by_category(period)
        total_exp = 0.0
        for cat, amt in sorted(exp_cats.items(), key=lambda x: -x[1]):
            row("Expense", cat, amt, "expense")
            total_exp += amt
        salary_payable = self._get_salary_payable()
        if salary_payable:
            row("Expense", "Salary Payable", salary_payable, "expense")
            total_exp += salary_payable
        self._pl_tree.insert("", "end", values=("", "Total Expenses", f"{total_exp:,.2f}"), tags=("total",))

        # ── NET ──
        net = total_sales - total_exp
        self._pl_tree.insert("", "end", values=("", "", ""), tags=())
        tag = "profit" if net >= 0 else "loss"
        label = "NET PROFIT" if net >= 0 else "NET LOSS"
        self._pl_tree.insert("", "end",
                             values=("", label, f"{abs(net):,.2f}"), tags=(tag,))

    def _refresh_cf(self):
        for iid in self._cf_tree.get_children():
            self._cf_tree.delete(iid)
        period = self._cf_period.get().strip()

        def row(section, desc, amt, tag=""):
            self._cf_tree.insert("", "end", values=(section, desc, f"{amt:,.2f}"), tags=(tag,))

        # ── INFLOWS ──
        self._cf_tree.insert("", "end", values=("CASH INFLOWS", "", ""), tags=("header",))
        received = self._get_total_received(period)
        row("Operating", "Cash received from customers", received, "in")

        bank_deposits = 0.0
        for t in self._filter_by_period(
                self._safe_read(BANK_TRANSACTIONS_CSV, BANK_TRANSACTIONS_HEADERS),
                "date", period):
            if t.get("transaction_type") in ("Deposit", "Cash In"):
                bank_deposits += self._tof(t.get("amount", 0))
        row("Financing", "Bank deposits / Cash in", bank_deposits, "in")
        total_in = received + bank_deposits
        self._cf_tree.insert("", "end", values=("", "Total Cash In", f"{total_in:,.2f}"), tags=("total",))

        # ── OUTFLOWS ──
        self._cf_tree.insert("", "end", values=("", "", ""), tags=())
        self._cf_tree.insert("", "end", values=("CASH OUTFLOWS", "", ""), tags=("header",))
        total_exp = self._get_total_expenses(period)
        row("Operating", "Business Expenses", total_exp, "out")

        ap_paid = sum(self._tof(r.get("amount_paid", 0))
                      for r in self._filter_by_period(
                          self._safe_read(AP_CSV, AP_HEADERS), "bill_date", period))
        row("Operating", "Vendor / Supplier Payments", ap_paid, "out")

        sal_paid = sum(self._tof(r.get("amount_paid", 0))
                       for r in self._filter_by_period(
                           self._safe_read(SALARY_CSV, SALARY_HEADERS), "month_year", period))
        row("Payroll", "Salary Paid", sal_paid, "out")

        bank_out = 0.0
        for t in self._filter_by_period(
                self._safe_read(BANK_TRANSACTIONS_CSV, BANK_TRANSACTIONS_HEADERS),
                "date", period):
            if t.get("transaction_type") in ("Withdrawal", "Cash Out", "Bank Charge"):
                bank_out += self._tof(t.get("amount", 0))
        row("Financing", "Bank withdrawals / charges", bank_out, "out")

        total_out = total_exp + ap_paid + sal_paid + bank_out
        self._cf_tree.insert("", "end", values=("", "Total Cash Out", f"{total_out:,.2f}"), tags=("total",))

        # ── NET CASH FLOW ──
        net_cf = total_in - total_out
        self._cf_tree.insert("", "end", values=("", "", ""), tags=())
        tag = "in" if net_cf >= 0 else "out"
        self._cf_tree.insert("", "end",
                             values=("", "NET CASH FLOW", f"{net_cf:,.2f}"), tags=(tag, "total"))

    def _refresh_bs(self):
        for iid in self._bs_tree.get_children():
            self._bs_tree.delete(iid)

        def row(section, item, amt, tag=""):
            self._bs_tree.insert("", "end", values=(section, item, f"{amt:,.2f}"), tags=(tag,))

        # ── ASSETS ──
        self._bs_tree.insert("", "end", values=("ASSETS", "", ""), tags=("header",))
        bank_bal = self._get_bank_balance()
        cash_bal = self._get_cash_in_hand()
        cust_out = self._get_customer_outstanding()
        inv_val  = self._get_inventory_value()

        row("Current Assets", "Bank Balance",               bank_bal, "asset")
        row("Current Assets", "Cash in Hand",               cash_bal, "asset")
        row("Current Assets", "Accounts Receivable (AR)",   cust_out, "asset")
        row("Current Assets", "Inventory / Stock Value",    inv_val,  "asset")
        total_assets = bank_bal + cash_bal + cust_out + inv_val
        self._bs_tree.insert("", "end", values=("", "Total Assets", f"{total_assets:,.2f}"), tags=("total",))

        # ── LIABILITIES ──
        self._bs_tree.insert("", "end", values=("", "", ""), tags=())
        self._bs_tree.insert("", "end", values=("LIABILITIES", "", ""), tags=("header",))
        vendor_out  = self._get_vendor_outstanding()
        salary_pay  = self._get_salary_payable()
        gst_pay     = self._get_gst_payable()

        row("Current Liabilities", "Accounts Payable (AP)",  vendor_out, "liab")
        row("Current Liabilities", "Salary Payable",         salary_pay, "liab")
        row("Current Liabilities", "GST Payable",            gst_pay,   "liab")
        total_liab = vendor_out + salary_pay + gst_pay
        self._bs_tree.insert("", "end", values=("", "Total Liabilities", f"{total_liab:,.2f}"), tags=("total",))

        # ── CAPITAL / EQUITY ──
        self._bs_tree.insert("", "end", values=("", "", ""), tags=())
        self._bs_tree.insert("", "end", values=("EQUITY", "", ""), tags=("header",))
        equity = total_assets - total_liab
        row("Owner's Equity", "Net Company Capital", equity, "asset")
        self._bs_tree.insert("", "end", values=("", "Total Equity", f"{equity:,.2f}"), tags=("total",))

    def _refresh_outstanding(self):
        # Customer outstanding
        for iid in self._cust_tree.get_children():
            self._cust_tree.delete(iid)
        ar_rows = self._safe_read(AR_CSV, AR_HEADERS)
        cust_total = 0.0
        for r in ar_rows:
            pend = self._tof(r.get("pending_amount", 0))
            if pend <= 0 or r.get("status") == "Paid":
                continue
            cust_total += pend
            st = r.get("status", "Unpaid")
            tag = "overdue" if st == "Overdue" else ("partial" if st == "Partial" else "")
            self._cust_tree.insert("", "end", tags=(tag,), values=(
                r.get("customer_name",""), r.get("invoice_number",""),
                r.get("invoice_date",""), r.get("due_date",""),
                f"{self._tof(r.get('invoice_amount',0)):,.2f}",
                f"{self._tof(r.get('amount_received',0)):,.2f}",
                f"{pend:,.2f}", st,
            ))
        self._out_cust_var.set(f"₹{cust_total:,.2f}")

        # Vendor outstanding
        for iid in self._vend_tree.get_children():
            self._vend_tree.delete(iid)
        ap_rows = self._safe_read(AP_CSV, AP_HEADERS)
        vend_total = 0.0
        for r in ap_rows:
            pend = self._tof(r.get("pending_amount", 0))
            if pend <= 0 or r.get("status") == "Paid":
                continue
            vend_total += pend
            st = r.get("status", "Unpaid")
            tag = "overdue" if st == "Overdue" else ("partial" if st == "Partial" else "")
            self._vend_tree.insert("", "end", tags=(tag,), values=(
                r.get("payee_name",""), r.get("bill_number",""),
                r.get("bill_date",""), r.get("due_date",""),
                f"{self._tof(r.get('bill_amount',0)):,.2f}",
                f"{self._tof(r.get('amount_paid',0)):,.2f}",
                f"{pend:,.2f}", st,
            ))
        self._out_vendor_var.set(f"₹{vend_total:,.2f}")

    # ====================================================== EXCEL EXPORT ==
    def export_all(self):
        path = filedialog.asksaveasfilename(
            parent=self, title="Save Financial Reports",
            defaultextension=".xlsx",
            initialfile=f"Financial_Reports_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            wb = Workbook()

            def styled_sheet(title, color):
                ws = wb.create_sheet(title)
                return ws

            def write_tree(ws, tree, header_color):
                col_count = len(tree["columns"])
                for col_idx in range(1, col_count + 1):
                    ws.column_dimensions[get_column_letter(col_idx)].width = 28
                for item in tree.get_children():
                    vals = tree.item(item, "values")
                    tags = tree.item(item, "tags")
                    ws.append(list(vals))
                    row = ws[ws.max_row]
                    if "header" in tags or "total" in tags:
                        for cell in row:
                            cell.font = Font(bold=True)
                    if "header" in tags:
                        for cell in row:
                            cell.fill = PatternFill("solid", fgColor=header_color)

            # Sheet 1 — Dashboard
            ws1 = wb.active; ws1.title = "Dashboard"
            ws1.append(["Metric", "Value"])
            for cell in ws1[1]: cell.font = Font(bold=True)
            for key, var in self._dash_vars.items():
                ws1.append([key.replace("_"," ").title(), var.get()])
            ws1.column_dimensions["A"].width = 28
            ws1.column_dimensions["B"].width = 18

            # Sheet 2 — P&L
            ws2 = wb.create_sheet("Profit & Loss")
            ws2.append(["Category", "Description", "Amount (₹)"])
            for cell in ws2[1]: cell.font = Font(bold=True)
            write_tree(ws2, self._pl_tree, "D5E8D4")

            # Sheet 3 — Cash Flow
            ws3 = wb.create_sheet("Cash Flow")
            ws3.append(["Section", "Description", "Amount (₹)"])
            for cell in ws3[1]: cell.font = Font(bold=True)
            write_tree(ws3, self._cf_tree, "D5E8D4")

            # Sheet 4 — Balance Sheet
            ws4 = wb.create_sheet("Balance Sheet")
            ws4.append(["Section", "Item", "Amount (₹)"])
            for cell in ws4[1]: cell.font = Font(bold=True)
            write_tree(ws4, self._bs_tree, "D5E8D4")

            # Sheet 5 — Customer Outstanding
            ws5 = wb.create_sheet("Customer Outstanding")
            ws5.append(["Customer","Invoice No.","Invoice Date","Due Date",
                         "Invoice Amt","Received","Pending","Status"])
            for cell in ws5[1]: cell.font = Font(bold=True)
            for item in self._cust_tree.get_children():
                ws5.append(list(self._cust_tree.item(item, "values")))

            # Sheet 6 — Vendor Outstanding
            ws6 = wb.create_sheet("Vendor Outstanding")
            ws6.append(["Vendor","Bill No.","Bill Date","Due Date",
                         "Bill Amt","Paid","Pending","Status"])
            for cell in ws6[1]: cell.font = Font(bold=True)
            for item in self._vend_tree.get_children():
                ws6.append(list(self._vend_tree.item(item, "values")))

            # Remove default empty sheet if present
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]

            wb.save(path)
            messagebox.showinfo("Exported",
                f"All financial reports exported to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", str(e), parent=self)


class InventoryStockView(ttk.Frame):
    """📦 Inventory / Stock — track raw materials, purchased/used qty, finished goods, stock value."""

    COLS = (
        "item_code", "item_name", "category", "unit",
        "opening_qty", "purchased_qty", "used_qty", "finished_qty",
        "current_stock", "unit_cost", "stock_value", "reorder_level", "notes",
    )
    HEADINGS = (
        "Item Code", "Item Name", "Category", "Unit",
        "Opening Qty", "Purchased", "Used", "Finished",
        "Current Stock", "Unit Cost (₹)", "Stock Value (₹)", "Reorder Level", "Notes",
    )

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_inventory_file()
        self._editing_index = None
        self._sort_reverse = {c: False for c in self.COLS}
        self._build()
        self.load_entries()

    def _build(self):
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="📦 Inventory / Stock",
                  font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel", command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh", command=self.load_entries).pack(side="right", padx=(0, 6))

        # Summary cards
        summary = ttk.LabelFrame(self, text=" Stock Summary ", padding=10)
        summary.pack(fill="x", pady=(0, 10))
        self._total_items_var  = tk.StringVar(value="0")
        self._total_value_var  = tk.StringVar(value="₹0.00")
        self._low_stock_var    = tk.StringVar(value="0")
        for col, (lbl, var, fg) in enumerate((
            ("Total Items",        self._total_items_var,  None),
            ("Total Stock Value",  self._total_value_var,  "navy"),
            ("Low Stock Alerts",   self._low_stock_var,    "red"),
        )):
            card = ttk.Frame(summary, padding=(12, 4))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            lbl_w = ttk.Label(card, textvariable=var, font=("Helvetica", 14, "bold"))
            if fg:
                lbl_w.configure(foreground=fg)
            lbl_w.pack(anchor="w")
            summary.columnconfigure(col, weight=1)

        # Form — Row 0
        form = ttk.LabelFrame(self, text=" Add / Edit Item ", padding=10)
        form.pack(fill="x", pady=(0, 10))

        fields_r0 = [
            ("Item Code:", "_f_code", 10),
            ("Item Name:", "_f_name", 22),
            ("Category:",  "_f_cat",  None),   # combobox
            ("Unit:",      "_f_unit", 8),
        ]
        for col, (lbl, attr, width) in enumerate(fields_r0):
            ttk.Label(form, text=lbl).grid(row=0, column=col*2, sticky="w", padx=4, pady=4)
            if attr == "_f_cat":
                self._f_cat_var = tk.StringVar()
                w = ttk.Combobox(form, textvariable=self._f_cat_var,
                                 values=list(INVENTORY_CATEGORIES), width=18)
                setattr(self, attr, w)
            else:
                w = ttk.Entry(form, width=width)
                setattr(self, attr, w)
            w.grid(row=0, column=col*2+1, sticky="ew", padx=4, pady=4)

        # Form — Row 1: quantities
        qty_labels = [
            ("Opening Qty:", "_f_open", 8),
            ("Purchased Qty:", "_f_purch", 8),
            ("Used Qty:", "_f_used", 8),
            ("Finished Qty:", "_f_finish", 8),
            ("Unit Cost (₹):", "_f_cost", 10),
            ("Reorder Level:", "_f_reorder", 8),
        ]
        for col, (lbl, attr, width) in enumerate(qty_labels):
            ttk.Label(form, text=lbl).grid(row=1, column=col*2, sticky="w", padx=4, pady=4)
            w = ttk.Entry(form, width=width)
            setattr(self, attr, w)
            w.grid(row=1, column=col*2+1, sticky="ew", padx=4, pady=4)

        # Form — Row 2: notes + buttons
        ttk.Label(form, text="Notes:").grid(row=2, column=0, sticky="w", padx=4, pady=4)
        self._f_notes = ttk.Entry(form, width=55)
        self._f_notes.grid(row=2, column=1, columnspan=9, sticky="ew", padx=4, pady=4)
        btn_row = ttk.Frame(form)
        btn_row.grid(row=2, column=10, columnspan=2, sticky="e", padx=4)
        ttk.Button(btn_row, text="💾 Save",   command=self.save_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="🗑 Delete", command=self.delete_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="✖ Clear",  command=self.clear_form).pack(side="left")
        for c in range(12):
            form.columnconfigure(c, weight=1)

        # Treeview
        tbox = ttk.LabelFrame(self, text=" Stock Register ", padding=8)
        tbox.pack(fill="both", expand=True)
        self._tree = ttk.Treeview(tbox, columns=self.COLS, show="headings", selectmode="browse")
        cw = {"item_name": 170, "category": 120, "unit": 55,
              "opening_qty": 75, "purchased_qty": 78, "used_qty": 65,
              "finished_qty": 75, "current_stock": 85,
              "unit_cost": 95, "stock_value": 100, "reorder_level": 90, "notes": 160}
        right = {"opening_qty","purchased_qty","used_qty","finished_qty",
                 "current_stock","unit_cost","stock_value","reorder_level"}
        for col, hdg in zip(self.COLS, self.HEADINGS):
            self._tree.heading(col, text=hdg, command=lambda c=col: self._sort_by(c))
            self._tree.column(col, width=cw.get(col, 80),
                              anchor="e" if col in right else "w")
        self._tree.tag_configure("low", foreground="red")
        vsb = ttk.Scrollbar(tbox, orient="vertical", command=self._tree.yview)
        hsb = ttk.Scrollbar(tbox, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tbox.rowconfigure(0, weight=1); tbox.columnconfigure(0, weight=1)
        self._tree.bind("<<TreeviewSelect>>", self._on_select)

    def _read_all(self):
        ensure_inventory_file()
        with open(INVENTORY_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_all(self, rows):
        with open(INVENTORY_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=INVENTORY_HEADERS)
            writer.writeheader(); writer.writerows(rows)

    @staticmethod
    def _tof(v):
        try: return float(str(v).replace(",", "") or 0)
        except ValueError: return 0.0

    def load_entries(self):
        for iid in self._tree.get_children():
            self._tree.delete(iid)
        rows = self._read_all()
        total_val = 0.0; low_count = 0
        for idx, row in enumerate(rows):
            op = self._tof(row.get("opening_qty", 0))
            pu = self._tof(row.get("purchased_qty", 0))
            us = self._tof(row.get("used_qty", 0))
            fi = self._tof(row.get("finished_qty", 0))
            cur = op + pu - us - fi
            cost = self._tof(row.get("unit_cost", 0))
            val = cur * cost
            total_val += val
            reorder = self._tof(row.get("reorder_level", 0))
            tag = ("low",) if reorder > 0 and cur <= reorder else ()
            if tag: low_count += 1
            self._tree.insert("", "end", iid=str(idx), tags=tag, values=(
                row.get("item_code",""), row.get("item_name",""),
                row.get("category",""), row.get("unit",""),
                f"{op:g}", f"{pu:g}", f"{us:g}", f"{fi:g}",
                f"{cur:g}", f"{cost:,.2f}", f"{val:,.2f}",
                f"{reorder:g}", row.get("notes",""),
            ))
        self._total_items_var.set(str(len(rows)))
        self._total_value_var.set(f"₹{total_val:,.2f}")
        self._low_stock_var.set(str(low_count))
        self._editing_index = None

    def _on_select(self, _e=None):
        sel = self._tree.selection()
        if not sel: return
        idx = int(sel[0]); rows = self._read_all()
        if idx >= len(rows): return
        r = rows[idx]; self._editing_index = idx
        for attr, key in (("_f_code","item_code"),("_f_name","item_name"),
                          ("_f_unit","unit"),("_f_open","opening_qty"),
                          ("_f_purch","purchased_qty"),("_f_used","used_qty"),
                          ("_f_finish","finished_qty"),("_f_cost","unit_cost"),
                          ("_f_reorder","reorder_level"),("_f_notes","notes")):
            w = getattr(self, attr); w.delete(0, tk.END); w.insert(0, r.get(key,""))
        self._f_cat_var.set(r.get("category",""))

    def clear_form(self):
        self._editing_index = None
        self._f_cat_var.set("")
        for attr in ("_f_code","_f_name","_f_unit","_f_open","_f_purch",
                     "_f_used","_f_finish","_f_cost","_f_reorder","_f_notes"):
            getattr(self, attr).delete(0, tk.END)
        self._tree.selection_remove(self._tree.selection())

    def save_entry(self):
        name = self._f_name.get().strip()
        if not name:
            messagebox.showwarning("Missing Field", "Item name is required.", parent=self); return
        op = self._tof(self._f_open.get())
        pu = self._tof(self._f_purch.get())
        us = self._tof(self._f_used.get())
        fi = self._tof(self._f_finish.get())
        cost = self._tof(self._f_cost.get())
        cur = op + pu - us - fi
        new_row = {
            "item_code": self._f_code.get().strip(),
            "item_name": name, "category": self._f_cat_var.get(),
            "unit": self._f_unit.get().strip(),
            "opening_qty": f"{op:g}", "purchased_qty": f"{pu:g}",
            "used_qty": f"{us:g}", "finished_qty": f"{fi:g}",
            "current_stock": f"{cur:g}", "unit_cost": f"{cost:.2f}",
            "stock_value": f"{cur * cost:.2f}",
            "reorder_level": self._f_reorder.get().strip(),
            "notes": self._f_notes.get().strip(),
        }
        rows = self._read_all()
        if self._editing_index is not None and self._editing_index < len(rows):
            rows[self._editing_index] = new_row; msg = "Item updated."
        else:
            rows.append(new_row); msg = "Item saved."
        self._write_all(rows)
        messagebox.showinfo("Saved", msg, parent=self); self.load_entries()

    def delete_entry(self):
        if self._editing_index is None:
            messagebox.showwarning("Select Row", "Select an item to delete.", parent=self); return
        rows = self._read_all()
        if self._editing_index >= len(rows): return
        r = rows[self._editing_index]
        if messagebox.askyesno("Confirm Delete", f"Delete '{r.get('item_name','')}'?", parent=self):
            rows.pop(self._editing_index); self._write_all(rows); self.load_entries()

    def _sort_by(self, col):
        rows = self._read_all(); reverse = self._sort_reverse[col]
        numeric = {"opening_qty","purchased_qty","used_qty","finished_qty",
                   "current_stock","unit_cost","stock_value","reorder_level"}
        rows.sort(key=lambda r: self._tof(r.get(col,0)) if col in numeric else r.get(col,"").lower(),
                  reverse=reverse)
        self._sort_reverse[col] = not reverse; self._write_all(rows); self.load_entries()

    def export_excel(self):
        rows = self._read_all()
        if not rows:
            messagebox.showinfo("No Data", "No inventory to export.", parent=self); return
        path = filedialog.asksaveasfilename(
            parent=self, title="Save Inventory Report",
            defaultextension=".xlsx",
            initialfile=f"Inventory_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook","*.xlsx"),("All files","*.*")])
        if not path: return
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook(); ws = wb.active; ws.title = "Inventory"
            ws.append(list(self.HEADINGS))
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1A5276")
                cell.alignment = Alignment(horizontal="center")
            total_val = 0.0
            for row in rows:
                op = self._tof(row.get("opening_qty",0))
                pu = self._tof(row.get("purchased_qty",0))
                us = self._tof(row.get("used_qty",0))
                fi = self._tof(row.get("finished_qty",0))
                cost = self._tof(row.get("unit_cost",0))
                cur = op + pu - us - fi; val = cur * cost; total_val += val
                ws.append([row.get("item_code",""), row.get("item_name",""),
                           row.get("category",""), row.get("unit",""),
                           op, pu, us, fi, cur, cost, val,
                           self._tof(row.get("reorder_level",0)), row.get("notes","")])
            ws.append(["","","","","","","","","","TOTAL", total_val, "",""])
            for cell in ws[ws.max_row]: cell.font = Font(bold=True)
            for col, w in zip("ABCDEFGHIJKLM", [10,22,14,6,9,9,9,9,9,12,14,10,22]):
                ws.column_dimensions[col].width = w
            for rc in ws.iter_rows(min_row=2, min_col=10, max_col=11):
                for cell in rc: cell.number_format = '#,##0.00'
            wb.save(path)
            messagebox.showinfo("Exported", f"Inventory saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", str(e), parent=self)


class EmployeeSalaryView(ttk.Frame):
    """👨‍💼 Employee Salary — track monthly salary, attendance, advance, and payment status."""

    COLS = (
        "month_year", "employee_name", "designation", "monthly_salary",
        "working_days", "present_days", "advance", "deductions",
        "net_payable", "amount_paid", "balance", "status", "notes",
    )
    HEADINGS = (
        "Month / Year", "Employee", "Designation", "Monthly Salary (₹)",
        "Working Days", "Present Days", "Advance (₹)", "Deductions (₹)",
        "Net Payable (₹)", "Paid (₹)", "Balance (₹)", "Status", "Notes",
    )
    STATUS_COLORS = {"Paid": "#27AE60", "Unpaid": "#E74C3C",
                     "Partial": "#F39C12", "Hold": "#8E44AD"}

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_salary_file()
        self._editing_index = None
        self._sort_reverse = {c: False for c in self.COLS}
        self._filter_month = tk.StringVar(value="")
        self._build()
        self.load_entries()

    def _build(self):
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="👨‍💼 Employee Salary",
                  font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel", command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh", command=self.load_entries).pack(side="right", padx=(0, 6))

        # Summary
        summary = ttk.LabelFrame(self, text=" Salary Summary ", padding=10)
        summary.pack(fill="x", pady=(0, 10))
        self._total_payable_var = tk.StringVar(value="₹0.00")
        self._total_paid_var    = tk.StringVar(value="₹0.00")
        self._total_balance_var = tk.StringVar(value="₹0.00")
        self._unpaid_count_var  = tk.StringVar(value="0")
        for col, (lbl, var, fg) in enumerate((
            ("Total Net Payable", self._total_payable_var, None),
            ("Total Paid",        self._total_paid_var,    "green"),
            ("Total Balance",     self._total_balance_var, "red"),
            ("Unpaid / Partial",  self._unpaid_count_var,  "#E74C3C"),
        )):
            card = ttk.Frame(summary, padding=(12, 4))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            lbl_w = ttk.Label(card, textvariable=var, font=("Helvetica", 14, "bold"))
            if fg: lbl_w.configure(foreground=fg)
            lbl_w.pack(anchor="w")
            summary.columnconfigure(col, weight=1)

        # Filter
        fbar = ttk.Frame(self)
        fbar.pack(fill="x", pady=(0, 6))
        ttk.Label(fbar, text="Filter Month (YYYY-MM):").pack(side="left", padx=(0, 4))
        ttk.Entry(fbar, textvariable=self._filter_month, width=10).pack(side="left")
        ttk.Button(fbar, text="Apply", command=self.load_entries).pack(side="left", padx=4)
        ttk.Button(fbar, text="Clear", command=lambda: [self._filter_month.set(""), self.load_entries()]).pack(side="left")

        # Form
        form = ttk.LabelFrame(self, text=" Add / Edit Entry ", padding=10)
        form.pack(fill="x", pady=(0, 10))

        ttk.Label(form, text="Month/Year:").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self._f_month = ttk.Entry(form, width=10)
        self._f_month.insert(0, datetime.now().strftime("%Y-%m"))
        self._f_month.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Employee:").grid(row=0, column=2, sticky="w", padx=4, pady=4)
        self._emp_var = tk.StringVar()
        self._cmb_emp = ttk.Combobox(form, textvariable=self._emp_var,
                                     values=self._load_employees(), width=22)
        self._cmb_emp.grid(row=0, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Designation:").grid(row=0, column=4, sticky="w", padx=4, pady=4)
        self._f_desig = ttk.Entry(form, width=18)
        self._f_desig.grid(row=0, column=5, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Monthly Salary:").grid(row=0, column=6, sticky="w", padx=4, pady=4)
        self._f_salary = ttk.Entry(form, width=12)
        self._f_salary.grid(row=0, column=7, sticky="ew", padx=4, pady=4)

        num_fields = [
            ("Working Days:", "_f_wdays", 6),
            ("Present Days:", "_f_pdays", 6),
            ("Advance (₹):",  "_f_adv",   10),
            ("Deductions (₹):", "_f_ded",  10),
            ("Amount Paid (₹):", "_f_paid", 10),
        ]
        for col, (lbl, attr, w) in enumerate(num_fields):
            ttk.Label(form, text=lbl).grid(row=1, column=col*2, sticky="w", padx=4, pady=4)
            ent = ttk.Entry(form, width=w)
            ent.grid(row=1, column=col*2+1, sticky="ew", padx=4, pady=4)
            setattr(self, attr, ent)

        ttk.Label(form, text="Status:").grid(row=1, column=10, sticky="w", padx=4, pady=4)
        self._status_var = tk.StringVar(value="Unpaid")
        ttk.Combobox(form, textvariable=self._status_var,
                     values=list(SALARY_STATUS_OPTIONS), state="readonly", width=10
                     ).grid(row=1, column=11, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Notes:").grid(row=2, column=0, sticky="w", padx=4, pady=4)
        self._f_notes = ttk.Entry(form, width=55)
        self._f_notes.grid(row=2, column=1, columnspan=9, sticky="ew", padx=4, pady=4)
        btn_row = ttk.Frame(form)
        btn_row.grid(row=2, column=10, columnspan=2, sticky="e", padx=4)
        ttk.Button(btn_row, text="💾 Save",   command=self.save_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="🗑 Delete", command=self.delete_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="✖ Clear",  command=self.clear_form).pack(side="left")
        for c in range(12): form.columnconfigure(c, weight=1)

        # Treeview
        tbox = ttk.LabelFrame(self, text=" Salary Register ", padding=8)
        tbox.pack(fill="both", expand=True)
        self._tree = ttk.Treeview(tbox, columns=self.COLS, show="headings", selectmode="browse")
        cw = {"month_year":90,"employee_name":160,"designation":120,
              "monthly_salary":120,"working_days":80,"present_days":80,
              "advance":90,"deductions":90,"net_payable":110,
              "amount_paid":90,"balance":90,"status":75,"notes":160}
        right = {"monthly_salary","advance","deductions","net_payable","amount_paid","balance"}
        for col, hdg in zip(self.COLS, self.HEADINGS):
            self._tree.heading(col, text=hdg, command=lambda c=col: self._sort_by(c))
            self._tree.column(col, width=cw.get(col,80),
                              anchor="e" if col in right else "w")
        for st, fg in self.STATUS_COLORS.items():
            self._tree.tag_configure(st, foreground=fg)
        vsb = ttk.Scrollbar(tbox, orient="vertical", command=self._tree.yview)
        hsb = ttk.Scrollbar(tbox, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns"); hsb.grid(row=1, column=0, sticky="ew")
        tbox.rowconfigure(0, weight=1); tbox.columnconfigure(0, weight=1)
        self._tree.bind("<<TreeviewSelect>>", self._on_select)

    @staticmethod
    def _load_employees():
        hr_csv = os.path.join(config.HR_DIR, "employee record.csv")
        if not os.path.exists(hr_csv):
            legacy_hr_csv = os.path.join(config.CSV_DIR, "employee record.csv")
            if os.path.exists(legacy_hr_csv):
                hr_csv = legacy_hr_csv
            else:
                return []

        names = []
        with open(hr_csv, newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    n = (row.get("name") or row.get("employee_name") or
                         row.get("full_name") or "").strip()
                    if n and n not in names: names.append(n)
        return sorted(names)

    def _read_all(self):
        ensure_salary_file()
        with open(SALARY_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_all(self, rows):
        with open(SALARY_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=SALARY_HEADERS)
            writer.writeheader(); writer.writerows(rows)

    @staticmethod
    def _tof(v):
        try: return float(str(v).replace(",","") or 0)
        except ValueError: return 0.0

    def _calc_net(self, salary, wdays, pdays, adv, ded):
        if wdays > 0:
            per_day = salary / wdays
            earned  = per_day * pdays
        else:
            earned = salary
        return max(earned - adv - ded, 0.0)

    def load_entries(self):
        for iid in self._tree.get_children(): self._tree.delete(iid)
        rows = self._read_all()
        fmonth = self._filter_month.get().strip()
        total_payable = total_paid = total_balance = unpaid_count = 0.0
        for idx, row in enumerate(rows):
            if fmonth and not row.get("month_year","").startswith(fmonth):
                continue
            net  = self._tof(row.get("net_payable", 0))
            paid = self._tof(row.get("amount_paid", 0))
            bal  = self._tof(row.get("balance", 0))
            st   = row.get("status", "Unpaid")
            total_payable += net; total_paid += paid; total_balance += bal
            if st in ("Unpaid", "Partial"): unpaid_count += 1
            self._tree.insert("", "end", iid=str(idx), tags=(st,), values=(
                row.get("month_year",""), row.get("employee_name",""),
                row.get("designation",""), f"{self._tof(row.get('monthly_salary',0)):,.2f}",
                row.get("working_days",""), row.get("present_days",""),
                f"{self._tof(row.get('advance',0)):,.2f}",
                f"{self._tof(row.get('deductions',0)):,.2f}",
                f"{net:,.2f}", f"{paid:,.2f}", f"{bal:,.2f}", st, row.get("notes",""),
            ))
        self._total_payable_var.set(f"₹{total_payable:,.2f}")
        self._total_paid_var.set(f"₹{total_paid:,.2f}")
        self._total_balance_var.set(f"₹{total_balance:,.2f}")
        self._unpaid_count_var.set(str(int(unpaid_count)))
        self._editing_index = None

    def _on_select(self, _e=None):
        sel = self._tree.selection()
        if not sel: return
        idx = int(sel[0]); rows = self._read_all()
        if idx >= len(rows): return
        r = rows[idx]; self._editing_index = idx
        self._f_month.delete(0, tk.END); self._f_month.insert(0, r.get("month_year",""))
        self._emp_var.set(r.get("employee_name",""))
        self._f_desig.delete(0, tk.END); self._f_desig.insert(0, r.get("designation",""))
        for attr, key in (("_f_salary","monthly_salary"),("_f_wdays","working_days"),
                          ("_f_pdays","present_days"),("_f_adv","advance"),
                          ("_f_ded","deductions"),("_f_paid","amount_paid"),
                          ("_f_notes","notes")):
            w = getattr(self, attr); w.delete(0, tk.END); w.insert(0, r.get(key,""))
        self._status_var.set(r.get("status","Unpaid"))

    def clear_form(self):
        self._editing_index = None
        self._f_month.delete(0, tk.END)
        self._f_month.insert(0, datetime.now().strftime("%Y-%m"))
        self._emp_var.set(""); self._status_var.set("Unpaid")
        for attr in ("_f_desig","_f_salary","_f_wdays","_f_pdays",
                     "_f_adv","_f_ded","_f_paid","_f_notes"):
            getattr(self, attr).delete(0, tk.END)
        self._tree.selection_remove(self._tree.selection())

    def save_entry(self):
        emp = self._emp_var.get().strip()
        if not emp:
            messagebox.showwarning("Missing Field", "Employee name is required.", parent=self); return
        sal   = self._tof(self._f_salary.get())
        wdays = self._tof(self._f_wdays.get()) or 26
        pdays = self._tof(self._f_pdays.get()) or wdays
        adv   = self._tof(self._f_adv.get())
        ded   = self._tof(self._f_ded.get())
        paid  = self._tof(self._f_paid.get())
        net   = self._calc_net(sal, wdays, pdays, adv, ded)
        bal   = max(net - paid, 0.0)
        status = self._status_var.get()
        new_row = {
            "month_year": self._f_month.get().strip(),
            "employee_name": emp, "designation": self._f_desig.get().strip(),
            "monthly_salary": f"{sal:.2f}", "working_days": f"{wdays:g}",
            "present_days": f"{pdays:g}", "advance": f"{adv:.2f}",
            "deductions": f"{ded:.2f}", "net_payable": f"{net:.2f}",
            "amount_paid": f"{paid:.2f}", "balance": f"{bal:.2f}",
            "status": status, "notes": self._f_notes.get().strip(),
        }
        rows = self._read_all()
        if self._editing_index is not None and self._editing_index < len(rows):
            rows[self._editing_index] = new_row; msg = "Entry updated."
        else:
            rows.append(new_row); msg = "Entry saved."
        self._write_all(rows)
        messagebox.showinfo("Saved", msg, parent=self); self.load_entries()

    def delete_entry(self):
        if self._editing_index is None:
            messagebox.showwarning("Select Row", "Select a row to delete.", parent=self); return
        rows = self._read_all()
        if self._editing_index >= len(rows): return
        r = rows[self._editing_index]
        if messagebox.askyesno("Confirm Delete",
            f"Delete salary record for {r.get('employee_name','')} ({r.get('month_year','')})?",
            parent=self):
            rows.pop(self._editing_index); self._write_all(rows); self.load_entries()

    def _sort_by(self, col):
        rows = self._read_all(); reverse = self._sort_reverse[col]
        numeric = {"monthly_salary","advance","deductions","net_payable","amount_paid","balance"}
        rows.sort(key=lambda r: self._tof(r.get(col,0)) if col in numeric else r.get(col,"").lower(),
                  reverse=reverse)
        self._sort_reverse[col] = not reverse; self._write_all(rows); self.load_entries()

    def export_excel(self):
        rows = self._read_all()
        if not rows:
            messagebox.showinfo("No Data", "No salary records to export.", parent=self); return
        path = filedialog.asksaveasfilename(
            parent=self, title="Save Salary Report", defaultextension=".xlsx",
            initialfile=f"Salary_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook","*.xlsx"),("All files","*.*")])
        if not path: return
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook(); ws = wb.active; ws.title = "Salary"
            ws.append(list(self.HEADINGS))
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2E4057")
                cell.alignment = Alignment(horizontal="center")
            tot_net = tot_paid = tot_bal = 0.0
            for row in rows:
                net  = self._tof(row.get("net_payable",0))
                paid = self._tof(row.get("amount_paid",0))
                bal  = self._tof(row.get("balance",0))
                tot_net += net; tot_paid += paid; tot_bal += bal
                ws.append([
                    row.get("month_year",""), row.get("employee_name",""),
                    row.get("designation",""), self._tof(row.get("monthly_salary",0)),
                    self._tof(row.get("working_days",0)), self._tof(row.get("present_days",0)),
                    self._tof(row.get("advance",0)), self._tof(row.get("deductions",0)),
                    net, paid, bal, row.get("status",""), row.get("notes",""),
                ])
            ws.append(["","","","","","","","","TOTAL", tot_paid, tot_bal,"",""])
            for cell in ws[ws.max_row]: cell.font = Font(bold=True)
            for col, w in zip("ABCDEFGHIJKLM", [10,22,18,14,10,10,12,12,14,12,12,10,22]):
                ws.column_dimensions[col].width = w
            for rc in ws.iter_rows(min_row=2, min_col=4, max_col=11):
                for cell in rc: cell.number_format = '#,##0.00'
            wb.save(path)
            messagebox.showinfo("Exported", f"Salary report saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", str(e), parent=self)


class GSTTaxView(ttk.Frame):
    """🧾 GST & Tax — track sales GST, purchase GST/ITC, GST payable, returns, TDS, income tax."""

    COLS = (
        "date", "entry_type", "reference_number", "party_name",
        "taxable_amount", "cgst", "sgst", "igst", "total_gst",
        "tds_amount", "period", "status", "notes",
    )
    HEADINGS = (
        "Date", "Entry Type", "Reference No.", "Party Name",
        "Taxable Amount", "CGST (₹)", "SGST (₹)", "IGST (₹)", "Total GST (₹)",
        "TDS (₹)", "Period", "Status", "Notes",
    )
    STATUS_COLORS = {"Filed": "#27AE60", "Paid": "#2980B9",
                     "Pending": "#F39C12", "Overdue": "#E74C3C"}

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_gst_file()
        self._editing_index = None
        self._sort_reverse = {c: False for c in self.COLS}
        self._filter_type = tk.StringVar(value="All")
        self._build()
        self.load_entries()

    def _build(self):
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="🧾 GST & Tax",
                  font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel", command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh", command=self.load_entries).pack(side="right", padx=(0, 6))

        # Summary
        summary = ttk.LabelFrame(self, text=" Tax Summary ", padding=10)
        summary.pack(fill="x", pady=(0, 10))
        self._sales_gst_var    = tk.StringVar(value="₹0.00")
        self._purchase_gst_var = tk.StringVar(value="₹0.00")
        self._gst_payable_var  = tk.StringVar(value="₹0.00")
        self._tds_var          = tk.StringVar(value="₹0.00")
        for col, (lbl, var, fg) in enumerate((
            ("Sales GST Collected",    self._sales_gst_var,    None),
            ("Purchase GST / ITC",     self._purchase_gst_var, "green"),
            ("Net GST Payable",        self._gst_payable_var,  "red"),
            ("TDS",                    self._tds_var,          "#8E44AD"),
        )):
            card = ttk.Frame(summary, padding=(12, 4))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            lbl_w = ttk.Label(card, textvariable=var, font=("Helvetica", 14, "bold"))
            if fg: lbl_w.configure(foreground=fg)
            lbl_w.pack(anchor="w")
            summary.columnconfigure(col, weight=1)

        # Filter
        fbar = ttk.Frame(self)
        fbar.pack(fill="x", pady=(0, 6))
        ttk.Label(fbar, text="Filter by Type:").pack(side="left", padx=(0, 4))
        types_for_filter = ["All"] + list(GST_ENTRY_TYPES)
        type_combo = ttk.Combobox(fbar, textvariable=self._filter_type,
                                  values=types_for_filter, state="readonly", width=28)
        type_combo.pack(side="left")
        type_combo.bind("<<ComboboxSelected>>", lambda _: self.load_entries())

        # Form
        form = ttk.LabelFrame(self, text=" Add / Edit Entry ", padding=10)
        form.pack(fill="x", pady=(0, 10))

        ttk.Label(form, text="Date:").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self._f_date = DateEntry(form, date_pattern="yyyy-mm-dd", width=13)
        self._f_date.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Entry Type:").grid(row=0, column=2, sticky="w", padx=4, pady=4)
        self._type_var = tk.StringVar(value="Sales GST")
        ttk.Combobox(form, textvariable=self._type_var,
                     values=list(GST_ENTRY_TYPES), state="readonly", width=22
                     ).grid(row=0, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Reference No.:").grid(row=0, column=4, sticky="w", padx=4, pady=4)
        self._f_ref = ttk.Entry(form, width=14)
        self._f_ref.grid(row=0, column=5, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Party Name:").grid(row=0, column=6, sticky="w", padx=4, pady=4)
        self._f_party = ttk.Entry(form, width=20)
        self._f_party.grid(row=0, column=7, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Period:").grid(row=0, column=8, sticky="w", padx=4, pady=4)
        self._f_period = ttk.Entry(form, width=10)
        self._f_period.grid(row=0, column=9, sticky="ew", padx=4, pady=4)

        num_fields = [
            ("Taxable Amount:", "_f_taxable", 12),
            ("CGST (₹):",       "_f_cgst",   10),
            ("SGST (₹):",       "_f_sgst",   10),
            ("IGST (₹):",       "_f_igst",   10),
            ("TDS (₹):",        "_f_tds",    10),
        ]
        for col, (lbl, attr, w) in enumerate(num_fields):
            ttk.Label(form, text=lbl).grid(row=1, column=col*2, sticky="w", padx=4, pady=4)
            ent = ttk.Entry(form, width=w)
            ent.grid(row=1, column=col*2+1, sticky="ew", padx=4, pady=4)
            setattr(self, attr, ent)

        ttk.Label(form, text="Status:").grid(row=1, column=10, sticky="w", padx=4, pady=4)
        self._status_var = tk.StringVar(value="Pending")
        ttk.Combobox(form, textvariable=self._status_var,
                     values=list(GST_STATUS_OPTIONS), state="readonly", width=10
                     ).grid(row=1, column=11, sticky="ew", padx=4, pady=4)

        ttk.Label(form, text="Notes:").grid(row=2, column=0, sticky="w", padx=4, pady=4)
        self._f_notes = ttk.Entry(form, width=55)
        self._f_notes.grid(row=2, column=1, columnspan=9, sticky="ew", padx=4, pady=4)
        btn_row = ttk.Frame(form)
        btn_row.grid(row=2, column=10, columnspan=2, sticky="e", padx=4)
        ttk.Button(btn_row, text="💾 Save",   command=self.save_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="🗑 Delete", command=self.delete_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="✖ Clear",  command=self.clear_form).pack(side="left")
        for c in range(12): form.columnconfigure(c, weight=1)

        # Treeview
        tbox = ttk.LabelFrame(self, text=" GST & Tax Register ", padding=8)
        tbox.pack(fill="both", expand=True)
        self._tree = ttk.Treeview(tbox, columns=self.COLS, show="headings", selectmode="browse")
        cw = {"date":90,"entry_type":150,"reference_number":110,"party_name":160,
              "taxable_amount":115,"cgst":90,"sgst":90,"igst":90,"total_gst":105,
              "tds_amount":90,"period":80,"status":75,"notes":160}
        right = {"taxable_amount","cgst","sgst","igst","total_gst","tds_amount"}
        for col, hdg in zip(self.COLS, self.HEADINGS):
            self._tree.heading(col, text=hdg, command=lambda c=col: self._sort_by(c))
            self._tree.column(col, width=cw.get(col,90),
                              anchor="e" if col in right else "w")
        for st, fg in self.STATUS_COLORS.items():
            self._tree.tag_configure(st, foreground=fg)
        vsb = ttk.Scrollbar(tbox, orient="vertical", command=self._tree.yview)
        hsb = ttk.Scrollbar(tbox, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns"); hsb.grid(row=1, column=0, sticky="ew")
        tbox.rowconfigure(0, weight=1); tbox.columnconfigure(0, weight=1)
        self._tree.bind("<<TreeviewSelect>>", self._on_select)

    def _read_all(self):
        ensure_gst_file()
        with open(GST_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_all(self, rows):
        with open(GST_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=GST_HEADERS)
            writer.writeheader(); writer.writerows(rows)

    @staticmethod
    def _tof(v):
        try: return float(str(v).replace(",","") or 0)
        except ValueError: return 0.0

    def load_entries(self):
        for iid in self._tree.get_children(): self._tree.delete(iid)
        rows = self._read_all()
        ftype = self._filter_type.get()
        sales_gst = purchase_gst = tds_total = 0.0
        for idx, row in enumerate(rows):
            et = row.get("entry_type","")
            if ftype != "All" and et != ftype: continue
            cgst = self._tof(row.get("cgst",0))
            sgst = self._tof(row.get("sgst",0))
            igst = self._tof(row.get("igst",0))
            total_gst = cgst + sgst + igst
            tds = self._tof(row.get("tds_amount",0))
            tds_total += tds
            if et == "Sales GST": sales_gst += total_gst
            elif "Purchase" in et or "ITC" in et: purchase_gst += total_gst
            st = row.get("status","Pending")
            self._tree.insert("", "end", iid=str(idx), tags=(st,), values=(
                row.get("date",""), et, row.get("reference_number",""),
                row.get("party_name",""),
                f"{self._tof(row.get('taxable_amount',0)):,.2f}",
                f"{cgst:,.2f}", f"{sgst:,.2f}", f"{igst:,.2f}", f"{total_gst:,.2f}",
                f"{tds:,.2f}", row.get("period",""), st, row.get("notes",""),
            ))
        net_payable = max(sales_gst - purchase_gst, 0.0)
        self._sales_gst_var.set(f"₹{sales_gst:,.2f}")
        self._purchase_gst_var.set(f"₹{purchase_gst:,.2f}")
        self._gst_payable_var.set(f"₹{net_payable:,.2f}")
        self._tds_var.set(f"₹{tds_total:,.2f}")
        self._editing_index = None

    def _on_select(self, _e=None):
        sel = self._tree.selection()
        if not sel: return
        idx = int(sel[0]); rows = self._read_all()
        if idx >= len(rows): return
        r = rows[idx]; self._editing_index = idx
        try: self._f_date.set_date(r.get("date", date.today().isoformat()))
        except Exception: pass
        self._type_var.set(r.get("entry_type",""))
        self._f_ref.delete(0, tk.END); self._f_ref.insert(0, r.get("reference_number",""))
        self._f_party.delete(0, tk.END); self._f_party.insert(0, r.get("party_name",""))
        self._f_period.delete(0, tk.END); self._f_period.insert(0, r.get("period",""))
        for attr, key in (("_f_taxable","taxable_amount"),("_f_cgst","cgst"),
                          ("_f_sgst","sgst"),("_f_igst","igst"),
                          ("_f_tds","tds_amount"),("_f_notes","notes")):
            w = getattr(self, attr); w.delete(0, tk.END); w.insert(0, r.get(key,""))
        self._status_var.set(r.get("status","Pending"))

    def clear_form(self):
        self._editing_index = None
        try: self._f_date.set_date(date.today())
        except Exception: pass
        self._type_var.set("Sales GST"); self._status_var.set("Pending")
        for attr in ("_f_ref","_f_party","_f_period","_f_taxable",
                     "_f_cgst","_f_sgst","_f_igst","_f_tds","_f_notes"):
            getattr(self, attr).delete(0, tk.END)
        self._tree.selection_remove(self._tree.selection())

    def save_entry(self):
        et = self._type_var.get().strip()
        taxable = self._tof(self._f_taxable.get())
        cgst    = self._tof(self._f_cgst.get())
        sgst    = self._tof(self._f_sgst.get())
        igst    = self._tof(self._f_igst.get())
        tds     = self._tof(self._f_tds.get())
        total_gst = cgst + sgst + igst
        new_row = {
            "date": self._f_date.get(), "entry_type": et,
            "reference_number": self._f_ref.get().strip(),
            "party_name": self._f_party.get().strip(),
            "taxable_amount": f"{taxable:.2f}",
            "cgst": f"{cgst:.2f}", "sgst": f"{sgst:.2f}", "igst": f"{igst:.2f}",
            "total_gst": f"{total_gst:.2f}", "tds_amount": f"{tds:.2f}",
            "period": self._f_period.get().strip(),
            "status": self._status_var.get(), "notes": self._f_notes.get().strip(),
        }
        rows = self._read_all()
        if self._editing_index is not None and self._editing_index < len(rows):
            rows[self._editing_index] = new_row; msg = "Entry updated."
        else:
            rows.append(new_row); msg = "Entry saved."
        self._write_all(rows)
        messagebox.showinfo("Saved", msg, parent=self); self.load_entries()

    def delete_entry(self):
        if self._editing_index is None:
            messagebox.showwarning("Select Row", "Select a row to delete.", parent=self); return
        rows = self._read_all()
        if self._editing_index >= len(rows): return
        r = rows[self._editing_index]
        if messagebox.askyesno("Confirm Delete",
            f"Delete {r.get('entry_type','')} entry {r.get('reference_number','')}?",
            parent=self):
            rows.pop(self._editing_index); self._write_all(rows); self.load_entries()

    def _sort_by(self, col):
        rows = self._read_all(); reverse = self._sort_reverse[col]
        numeric = {"taxable_amount","cgst","sgst","igst","total_gst","tds_amount"}
        rows.sort(key=lambda r: self._tof(r.get(col,0)) if col in numeric else r.get(col,"").lower(),
                  reverse=reverse)
        self._sort_reverse[col] = not reverse; self._write_all(rows); self.load_entries()

    def export_excel(self):
        rows = self._read_all()
        if not rows:
            messagebox.showinfo("No Data", "No GST/Tax records to export.", parent=self); return
        path = filedialog.asksaveasfilename(
            parent=self, title="Save GST & Tax Report", defaultextension=".xlsx",
            initialfile=f"GST_Tax_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook","*.xlsx"),("All files","*.*")])
        if not path: return
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook(); ws = wb.active; ws.title = "GST & Tax"
            ws.append(list(self.HEADINGS))
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1E8449")
                cell.alignment = Alignment(horizontal="center")
            tot_sales = tot_itc = tot_tds = 0.0
            for row in rows:
                cgst = self._tof(row.get("cgst",0))
                sgst = self._tof(row.get("sgst",0))
                igst = self._tof(row.get("igst",0))
                tgst = cgst + sgst + igst
                tds  = self._tof(row.get("tds_amount",0))
                et   = row.get("entry_type","")
                if et == "Sales GST": tot_sales += tgst
                elif "Purchase" in et or "ITC" in et: tot_itc += tgst
                tot_tds += tds
                ws.append([
                    row.get("date",""), et, row.get("reference_number",""),
                    row.get("party_name",""),
                    self._tof(row.get("taxable_amount",0)),
                    cgst, sgst, igst, tgst, tds,
                    row.get("period",""), row.get("status",""), row.get("notes",""),
                ])
            ws.append(["","","","","","","","","Sales GST", tot_sales, "","",""])
            ws.append(["","","","","","","","","ITC",        tot_itc,  "","",""])
            ws.append(["","","","","","","","","Net Payable", max(tot_sales-tot_itc,0), "","",""])
            ws.append(["","","","","","","","","TDS Total",  tot_tds,  "","",""])
            for cell in ws[ws.max_row]: cell.font = Font(bold=True)
            for col, w in zip("ABCDEFGHIJKLM", [12,22,14,22,14,10,10,10,12,10,10,10,22]):
                ws.column_dimensions[col].width = w
            for rc in ws.iter_rows(min_row=2, min_col=5, max_col=10):
                for cell in rc: cell.number_format = '#,##0.00'
            wb.save(path)
            messagebox.showinfo("Exported", f"GST & Tax report saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", str(e), parent=self)


class AccountsPayableView(ttk.Frame):
    """📤 Accounts Payable — track supplier bills, vendor payments, loans/EMI, and other liabilities."""

    COLS = (
        "bill_date", "due_date", "payee_name", "bill_number",
        "payment_type", "bill_amount", "amount_paid", "pending_amount",
        "status", "notes",
    )
    HEADINGS = (
        "Bill Date", "Due Date", "Payee / Vendor", "Bill No.",
        "Type", "Bill Amount", "Paid", "Pending",
        "Status", "Notes",
    )
    STATUS_COLORS = {
        "Paid":     "#27AE60",
        "Partial":  "#F39C12",
        "Unpaid":   "#E74C3C",
        "Overdue":  "#8E44AD",
        "Disputed": "#2C3E50",
    }

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_ap_file()
        self._editing_index = None
        self._sort_reverse  = {c: False for c in self.COLS}
        self._filter_status = tk.StringVar(value="All")
        self._build()
        self.load_entries()

    # ------------------------------------------------------------------ UI --
    def _build(self):
        # ── Header ───────────────────────────────────────────────────────
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="📤 Accounts Payable",
                  font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel",
                   command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh",
                   command=self.load_entries).pack(side="right", padx=(0, 6))

        # ── Summary cards ─────────────────────────────────────────────────
        summary = ttk.LabelFrame(self, text=" Liability Summary ", padding=10)
        summary.pack(fill="x", pady=(0, 10))
        self._total_billed_var   = tk.StringVar(value="₹0.00")
        self._total_paid_var     = tk.StringVar(value="₹0.00")
        self._total_pending_var  = tk.StringVar(value="₹0.00")
        self._overdue_count_var  = tk.StringVar(value="0")
        for col, (lbl, var, fg) in enumerate((
            ("Total Billed",       self._total_billed_var,  None),
            ("Total Paid",         self._total_paid_var,    "green"),
            ("Total Outstanding",  self._total_pending_var, "red"),
            ("Overdue Bills",      self._overdue_count_var, "#8E44AD"),
        )):
            card = ttk.Frame(summary, padding=(12, 4))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            lbl_val = ttk.Label(card, textvariable=var, font=("Helvetica", 14, "bold"))
            if fg:
                lbl_val.configure(foreground=fg)
            lbl_val.pack(anchor="w")
            summary.columnconfigure(col, weight=1)

        # ── Filter bar ────────────────────────────────────────────────────
        filter_bar = ttk.Frame(self)
        filter_bar.pack(fill="x", pady=(0, 6))
        ttk.Label(filter_bar, text="Filter by Status:").pack(side="left", padx=(0, 6))
        for status in ("All",) + AP_STATUS_OPTIONS:
            ttk.Radiobutton(
                filter_bar, text=status,
                variable=self._filter_status, value=status,
                command=self.load_entries,
            ).pack(side="left", padx=3)

        # ── Entry form ────────────────────────────────────────────────────
        form_box = ttk.LabelFrame(self, text=" Add / Edit Entry ", padding=10)
        form_box.pack(fill="x", pady=(0, 10))

        # Row 0
        ttk.Label(form_box, text="Bill Date:").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self._ent_bill_date = DateEntry(form_box, date_pattern="yyyy-mm-dd", width=13)
        self._ent_bill_date.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Due Date:").grid(row=0, column=2, sticky="w", padx=4, pady=4)
        self._ent_due_date = DateEntry(form_box, date_pattern="yyyy-mm-dd", width=13)
        self._ent_due_date.grid(row=0, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Payee / Vendor:").grid(row=0, column=4, sticky="w", padx=4, pady=4)
        self._payee_var = tk.StringVar()
        self._cmb_payee = ttk.Combobox(
            form_box, textvariable=self._payee_var,
            values=self._load_vendor_names(), width=24,
        )
        self._cmb_payee.grid(row=0, column=5, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Bill No.:").grid(row=0, column=6, sticky="w", padx=4, pady=4)
        self._ent_bill_no = ttk.Entry(form_box, width=14)
        self._ent_bill_no.grid(row=0, column=7, sticky="ew", padx=4, pady=4)

        # Row 1
        ttk.Label(form_box, text="Payment Type:").grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self._type_var = tk.StringVar(value="Supplier Payment")
        ttk.Combobox(
            form_box, textvariable=self._type_var,
            values=list(AP_PAYMENT_TYPES), state="readonly", width=18,
        ).grid(row=1, column=1, columnspan=2, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Bill Amount:").grid(row=1, column=3, sticky="w", padx=4, pady=4)
        self._ent_bill_amt = ttk.Entry(form_box, width=13)
        self._ent_bill_amt.grid(row=1, column=4, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Amount Paid:").grid(row=1, column=5, sticky="w", padx=4, pady=4)
        self._ent_paid = ttk.Entry(form_box, width=13)
        self._ent_paid.grid(row=1, column=6, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Status:").grid(row=1, column=7, sticky="w", padx=4, pady=4)
        self._status_var = tk.StringVar(value="Unpaid")
        ttk.Combobox(
            form_box, textvariable=self._status_var,
            values=list(AP_STATUS_OPTIONS), state="readonly", width=12,
        ).grid(row=1, column=8, sticky="ew", padx=4, pady=4)

        # Row 2 — notes + buttons
        ttk.Label(form_box, text="Notes:").grid(row=2, column=0, sticky="w", padx=4, pady=4)
        self._ent_notes = ttk.Entry(form_box, width=55)
        self._ent_notes.grid(row=2, column=1, columnspan=6, sticky="ew", padx=4, pady=4)

        btn_row = ttk.Frame(form_box)
        btn_row.grid(row=2, column=7, columnspan=2, sticky="e", padx=4, pady=4)
        ttk.Button(btn_row, text="💾 Save",   command=self.save_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="🗑 Delete", command=self.delete_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="✖ Clear",  command=self.clear_form).pack(side="left")

        for c in range(9):
            form_box.columnconfigure(c, weight=1)

        # ── Treeview ──────────────────────────────────────────────────────
        table_box = ttk.LabelFrame(self, text=" Payable Register ", padding=8)
        table_box.pack(fill="both", expand=True)

        self._tree = ttk.Treeview(table_box, columns=self.COLS,
                                  show="headings", selectmode="browse")
        col_widths = {
            "bill_date": 90, "due_date": 90, "payee_name": 180,
            "bill_number": 100, "payment_type": 130,
            "bill_amount": 110, "amount_paid": 100,
            "pending_amount": 100, "status": 80, "notes": 180,
        }
        for col, heading in zip(self.COLS, self.HEADINGS):
            self._tree.heading(col, text=heading,
                               command=lambda c=col: self._sort_by(c))
            self._tree.column(
                col, width=col_widths.get(col, 100),
                anchor="e" if col in ("bill_amount", "amount_paid", "pending_amount") else "w",
            )
        for status, color in self.STATUS_COLORS.items():
            self._tree.tag_configure(status, foreground=color)

        vsb = ttk.Scrollbar(table_box, orient="vertical",   command=self._tree.yview)
        hsb = ttk.Scrollbar(table_box, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_box.rowconfigure(0, weight=1)
        table_box.columnconfigure(0, weight=1)
        self._tree.bind("<<TreeviewSelect>>", self._on_row_select)

    # -------------------------------------------------------- data helpers --
    @staticmethod
    def _load_vendor_names():
        """Pull vendor names from vendors.csv."""
        vendors_csv = os.path.join(config.CSV_DIR, "vendors.csv")
        names = []
        if os.path.exists(vendors_csv):
            with open(vendors_csv, newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    name = (row.get("vendor_name") or row.get("name") or
                            row.get("company_name") or "").strip()
                    if name and name not in names:
                        names.append(name)
        return sorted(names)

    def _read_all(self):
        ensure_ap_file()
        with open(AP_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_all(self, rows):
        with open(AP_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=AP_HEADERS)
            writer.writeheader()
            writer.writerows(rows)

    @staticmethod
    def _to_float(val):
        try:
            return float(str(val).replace(",", "") or 0)
        except ValueError:
            return 0.0

    def _auto_status(self, bill_amt, paid_amt, due_date_str):
        pending = bill_amt - paid_amt
        if pending <= 0:
            return "Paid"
        try:
            due = datetime.strptime(due_date_str, "%Y-%m-%d").date()
            if due < date.today():
                return "Overdue"
        except Exception:
            pass
        return "Partial" if paid_amt > 0 else "Unpaid"

    # ------------------------------------------------------------ actions --
    def load_entries(self):
        for iid in self._tree.get_children():
            self._tree.delete(iid)

        rows = self._read_all()
        filter_val = self._filter_status.get()
        total_billed = total_paid = total_pending = 0.0
        overdue_count = 0

        for idx, row in enumerate(rows):
            bill_amt = self._to_float(row.get("bill_amount", 0))
            paid_amt = self._to_float(row.get("amount_paid", 0))
            pending  = max(bill_amt - paid_amt, 0.0)
            status   = row.get("status", "Unpaid")

            total_billed  += bill_amt
            total_paid    += paid_amt
            total_pending += pending
            if status == "Overdue":
                overdue_count += 1

            if filter_val != "All" and status != filter_val:
                continue

            self._tree.insert("", "end", iid=str(idx), tags=(status,), values=(
                row.get("bill_date", ""),
                row.get("due_date", ""),
                row.get("payee_name", ""),
                row.get("bill_number", ""),
                row.get("payment_type", ""),
                f"{bill_amt:,.2f}",
                f"{paid_amt:,.2f}",
                f"{pending:,.2f}",
                status,
                row.get("notes", ""),
            ))

        self._total_billed_var.set(f"₹{total_billed:,.2f}")
        self._total_paid_var.set(f"₹{total_paid:,.2f}")
        self._total_pending_var.set(f"₹{total_pending:,.2f}")
        self._overdue_count_var.set(str(overdue_count))
        self._editing_index = None

    def _on_row_select(self, _event=None):
        sel = self._tree.selection()
        if not sel:
            return
        idx = int(sel[0])
        rows = self._read_all()
        if idx >= len(rows):
            return
        row = rows[idx]
        self._editing_index = idx
        try:
            self._ent_bill_date.set_date(row.get("bill_date", date.today().isoformat()))
        except Exception:
            pass
        try:
            self._ent_due_date.set_date(row.get("due_date", date.today().isoformat()))
        except Exception:
            pass
        self._payee_var.set(row.get("payee_name", ""))
        self._ent_bill_no.delete(0, tk.END)
        self._ent_bill_no.insert(0, row.get("bill_number", ""))
        self._type_var.set(row.get("payment_type", "Supplier Payment"))
        self._ent_bill_amt.delete(0, tk.END)
        self._ent_bill_amt.insert(0, row.get("bill_amount", ""))
        self._ent_paid.delete(0, tk.END)
        self._ent_paid.insert(0, row.get("amount_paid", ""))
        self._status_var.set(row.get("status", "Unpaid"))
        self._ent_notes.delete(0, tk.END)
        self._ent_notes.insert(0, row.get("notes", ""))

    def clear_form(self):
        self._editing_index = None
        try:
            self._ent_bill_date.set_date(date.today())
            self._ent_due_date.set_date(date.today())
        except Exception:
            pass
        self._payee_var.set("")
        self._type_var.set("Supplier Payment")
        self._status_var.set("Unpaid")
        for w in (self._ent_bill_no, self._ent_bill_amt,
                  self._ent_paid, self._ent_notes):
            w.delete(0, tk.END)
        self._tree.selection_remove(self._tree.selection())

    def save_entry(self):
        payee    = self._payee_var.get().strip()
        bill_no  = self._ent_bill_no.get().strip()
        bill_str = self._ent_bill_amt.get().strip()
        paid_str = self._ent_paid.get().strip()
        due_str  = self._ent_due_date.get()

        if not payee:
            messagebox.showwarning("Missing Field", "Payee / Vendor name is required.", parent=self)
            return
        if not bill_no:
            messagebox.showwarning("Missing Field", "Bill number is required.", parent=self)
            return
        try:
            bill_amt = float(bill_str.replace(",", "") or 0)
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Bill amount must be a number.", parent=self)
            return
        try:
            paid_amt = float(paid_str.replace(",", "") or 0)
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Amount paid must be a number.", parent=self)
            return

        pending = max(bill_amt - paid_amt, 0.0)
        status  = self._status_var.get() or self._auto_status(bill_amt, paid_amt, due_str)

        new_row = {
            "bill_date":    self._ent_bill_date.get(),
            "due_date":     due_str,
            "payee_name":   payee,
            "bill_number":  bill_no,
            "payment_type": self._type_var.get(),
            "bill_amount":  f"{bill_amt:.2f}",
            "amount_paid":  f"{paid_amt:.2f}",
            "pending_amount": f"{pending:.2f}",
            "status":       status,
            "notes":        self._ent_notes.get().strip(),
        }

        rows = self._read_all()
        if self._editing_index is not None and self._editing_index < len(rows):
            rows[self._editing_index] = new_row
            msg = "Entry updated."
        else:
            rows.append(new_row)
            msg = "Entry saved."

        self._write_all(rows)
        messagebox.showinfo("Saved", msg, parent=self)
        self.load_entries()

    def delete_entry(self):
        if self._editing_index is None:
            messagebox.showwarning("Select Row", "Select a row to delete.", parent=self)
            return
        rows = self._read_all()
        if self._editing_index >= len(rows):
            return
        row = rows[self._editing_index]
        if messagebox.askyesno(
            "Confirm Delete",
            f"Delete bill {row.get('bill_number','')} for {row.get('payee_name','')}?",
            parent=self,
        ):
            rows.pop(self._editing_index)
            self._write_all(rows)
            self.load_entries()

    def _sort_by(self, col):
        rows = self._read_all()
        reverse = self._sort_reverse[col]
        numeric = {"bill_amount", "amount_paid", "pending_amount"}

        def key(r):
            v = r.get(col, "")
            if col in numeric:
                try:
                    return float(str(v).replace(",", "") or 0)
                except ValueError:
                    return 0.0
            return v.lower()

        rows.sort(key=key, reverse=reverse)
        self._sort_reverse[col] = not reverse
        self._write_all(rows)
        self.load_entries()

    # ──────────────────────────────────────────────────── Excel export ───
    def export_excel(self):
        rows = self._read_all()
        if not rows:
            messagebox.showinfo("No Data", "No payable entries to export.", parent=self)
            return

        path = filedialog.asksaveasfilename(
            parent=self, title="Save Accounts Payable Report",
            defaultextension=".xlsx",
            initialfile=f"Accounts_Payable_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return

        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            wb = Workbook()
            ws = wb.active
            ws.title = "Payables"

            ws.append(list(self.HEADINGS))
            for cell in ws[1]:
                cell.font      = Font(bold=True, color="FFFFFF")
                cell.fill      = PatternFill("solid", fgColor="922B21")
                cell.alignment = Alignment(horizontal="center")

            total_billed = total_paid = total_pending = 0.0
            type_totals  = {}
            for row in rows:
                bill = self._to_float(row.get("bill_amount", 0))
                paid = self._to_float(row.get("amount_paid", 0))
                pend = max(bill - paid, 0.0)
                total_billed  += bill
                total_paid    += paid
                total_pending += pend
                pt = row.get("payment_type", "Other")
                type_totals[pt] = type_totals.get(pt, 0.0) + pend

                ws.append([
                    row.get("bill_date", ""), row.get("due_date", ""),
                    row.get("payee_name", ""), row.get("bill_number", ""),
                    row.get("payment_type", ""),
                    bill, paid, pend,
                    row.get("status", ""), row.get("notes", ""),
                ])

            # Totals row
            ws.append(["", "", "", "", "TOTAL",
                        total_billed, total_paid, total_pending, "", ""])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)

            for col, w in zip("ABCDEFGHIJ", [12, 12, 26, 14, 18, 14, 12, 12, 10, 26]):
                ws.column_dimensions[get_column_letter(ord(col) - 64)].width = w
            for row_cells in ws.iter_rows(min_row=2, min_col=6, max_col=8):
                for cell in row_cells:
                    cell.number_format = '#,##0.00'

            # Summary by type
            ws2 = wb.create_sheet("Summary by Type")
            ws2.append(["Payment Type", "Outstanding (₹)"])
            for cell in ws2[1]:
                cell.font = Font(bold=True)
            for pt, amt in sorted(type_totals.items(), key=lambda x: -x[1]):
                ws2.append([pt, amt])
            ws2.append(["TOTAL OUTSTANDING", total_pending])
            ws2.column_dimensions["A"].width = 22
            ws2.column_dimensions["B"].width = 18

            wb.save(path)
            messagebox.showinfo("Exported", f"Accounts Payable report saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", f"Could not save Excel file:\n{e}", parent=self)


class AccountsReceivableView(ttk.Frame):
    """📥 Accounts Receivable — track money customers owe you."""

    COLS = (
        "invoice_date", "due_date", "customer_name", "invoice_number",
        "product_service", "invoice_amount", "amount_received", "pending_amount",
        "status", "notes",
    )
    HEADINGS = (
        "Invoice Date", "Due Date", "Customer", "Invoice No.",
        "Product / Service", "Invoice Amount", "Received", "Pending",
        "Status", "Notes",
    )

    STATUS_COLORS = {
        "Paid":     "#27AE60",
        "Partial":  "#F39C12",
        "Unpaid":   "#E74C3C",
        "Overdue":  "#8E44AD",
        "Disputed": "#2C3E50",
    }

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_ar_file()
        self._editing_index = None
        self._sort_reverse = {c: False for c in self.COLS}
        self._filter_status = tk.StringVar(value="All")
        self._build()
        self.load_entries()

    # ------------------------------------------------------------------ UI --
    def _build(self):
        # ── Header ───────────────────────────────────────────────────────
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="📥 Accounts Receivable",
                  font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel",
                   command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh",
                   command=self.load_entries).pack(side="right", padx=(0, 6))

        # ── Summary cards ─────────────────────────────────────────────────
        summary = ttk.LabelFrame(self, text=" Cash Flow Summary ", padding=10)
        summary.pack(fill="x", pady=(0, 10))
        self._total_invoiced_var  = tk.StringVar(value="₹0.00")
        self._total_received_var  = tk.StringVar(value="₹0.00")
        self._total_pending_var   = tk.StringVar(value="₹0.00")
        self._overdue_count_var   = tk.StringVar(value="0")
        for col, (lbl, var, fg) in enumerate((
            ("Total Invoiced",     self._total_invoiced_var, None),
            ("Total Received",     self._total_received_var, "green"),
            ("Total Outstanding",  self._total_pending_var,  "red"),
            ("Overdue Invoices",   self._overdue_count_var,  "#8E44AD"),
        )):
            card = ttk.Frame(summary, padding=(12, 4))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            lbl_val = ttk.Label(card, textvariable=var,
                                font=("Helvetica", 14, "bold"))
            if fg:
                lbl_val.configure(foreground=fg)
            lbl_val.pack(anchor="w")
            summary.columnconfigure(col, weight=1)

        # ── Filter bar ────────────────────────────────────────────────────
        filter_bar = ttk.Frame(self)
        filter_bar.pack(fill="x", pady=(0, 6))
        ttk.Label(filter_bar, text="Filter by Status:").pack(side="left", padx=(0, 6))
        for status in ("All",) + AR_STATUS_OPTIONS:
            ttk.Radiobutton(
                filter_bar, text=status,
                variable=self._filter_status, value=status,
                command=self.load_entries,
            ).pack(side="left", padx=3)

        # ── Entry form ────────────────────────────────────────────────────
        form_box = ttk.LabelFrame(self, text=" Add / Edit Entry ", padding=10)
        form_box.pack(fill="x", pady=(0, 10))

        # Row 0
        ttk.Label(form_box, text="Invoice Date:").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self._ent_inv_date = DateEntry(form_box, date_pattern="yyyy-mm-dd", width=13)
        self._ent_inv_date.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Due Date:").grid(row=0, column=2, sticky="w", padx=4, pady=4)
        self._ent_due_date = DateEntry(form_box, date_pattern="yyyy-mm-dd", width=13)
        self._ent_due_date.grid(row=0, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Customer:").grid(row=0, column=4, sticky="w", padx=4, pady=4)
        self._customer_var = tk.StringVar()
        self._cmb_customer = ttk.Combobox(
            form_box, textvariable=self._customer_var,
            values=load_customer_names(), width=24,
        )
        self._cmb_customer.grid(row=0, column=5, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Invoice No.:").grid(row=0, column=6, sticky="w", padx=4, pady=4)
        self._ent_inv_no = ttk.Entry(form_box, width=14)
        self._ent_inv_no.grid(row=0, column=7, sticky="ew", padx=4, pady=4)

        # Row 1
        ttk.Label(form_box, text="Product / Service:").grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self._ent_product = ttk.Entry(form_box, width=28)
        self._ent_product.grid(row=1, column=1, columnspan=2, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Invoice Amount:").grid(row=1, column=3, sticky="w", padx=4, pady=4)
        self._ent_inv_amt = ttk.Entry(form_box, width=13)
        self._ent_inv_amt.grid(row=1, column=4, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Amount Received:").grid(row=1, column=5, sticky="w", padx=4, pady=4)
        self._ent_received = ttk.Entry(form_box, width=13)
        self._ent_received.grid(row=1, column=6, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Status:").grid(row=1, column=7, sticky="w", padx=4, pady=4)
        self._status_var = tk.StringVar(value="Unpaid")
        ttk.Combobox(
            form_box, textvariable=self._status_var,
            values=list(AR_STATUS_OPTIONS), state="readonly", width=12,
        ).grid(row=1, column=8, sticky="ew", padx=4, pady=4)

        # Row 2 — notes + action buttons
        ttk.Label(form_box, text="Notes:").grid(row=2, column=0, sticky="w", padx=4, pady=4)
        self._ent_notes = ttk.Entry(form_box, width=55)
        self._ent_notes.grid(row=2, column=1, columnspan=6, sticky="ew", padx=4, pady=4)

        btn_row = ttk.Frame(form_box)
        btn_row.grid(row=2, column=7, columnspan=2, sticky="e", padx=4, pady=4)
        ttk.Button(btn_row, text="💾 Save",   command=self.save_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="🗑 Delete", command=self.delete_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="✖ Clear",  command=self.clear_form).pack(side="left")

        for c in range(9):
            form_box.columnconfigure(c, weight=1)

        # ── Treeview ──────────────────────────────────────────────────────
        table_box = ttk.LabelFrame(self, text=" Receivable Register ", padding=8)
        table_box.pack(fill="both", expand=True)

        self._tree = ttk.Treeview(table_box, columns=self.COLS,
                                  show="headings", selectmode="browse")
        col_widths = {
            "invoice_date": 95, "due_date": 90, "customer_name": 180,
            "invoice_number": 105, "product_service": 180,
            "invoice_amount": 115, "amount_received": 100,
            "pending_amount": 100, "status": 80, "notes": 180,
        }
        for col, heading in zip(self.COLS, self.HEADINGS):
            self._tree.heading(col, text=heading,
                               command=lambda c=col: self._sort_by(c))
            self._tree.column(
                col, width=col_widths.get(col, 100),
                anchor="e" if col in ("invoice_amount", "amount_received", "pending_amount") else "w",
            )

        # Tag colors per status
        for status, color in self.STATUS_COLORS.items():
            self._tree.tag_configure(status, foreground=color)

        vsb = ttk.Scrollbar(table_box, orient="vertical",   command=self._tree.yview)
        hsb = ttk.Scrollbar(table_box, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_box.rowconfigure(0, weight=1)
        table_box.columnconfigure(0, weight=1)
        self._tree.bind("<<TreeviewSelect>>", self._on_row_select)

    # -------------------------------------------------------- data helpers --
    def _read_all(self):
        ensure_ar_file()
        with open(AR_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_all(self, rows):
        with open(AR_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=AR_HEADERS)
            writer.writeheader()
            writer.writerows(rows)

    @staticmethod
    def _to_float(val):
        try:
            return float(str(val).replace(",", "") or 0)
        except ValueError:
            return 0.0

    def _auto_status(self, inv_amt, rec_amt, due_date_str):
        """Auto-suggest status based on amounts and due date."""
        pending = inv_amt - rec_amt
        if pending <= 0:
            return "Paid"
        try:
            due = datetime.strptime(due_date_str, "%Y-%m-%d").date()
            if due < date.today():
                return "Overdue"
        except Exception:
            pass
        if rec_amt > 0:
            return "Partial"
        return "Unpaid"

    # ------------------------------------------------------------ actions --
    def load_entries(self):
        for iid in self._tree.get_children():
            self._tree.delete(iid)

        rows = self._read_all()
        filter_val = self._filter_status.get()

        total_inv = total_rec = total_pend = 0.0
        overdue_count = 0

        for idx, row in enumerate(rows):
            inv_amt = self._to_float(row.get("invoice_amount", 0))
            rec_amt = self._to_float(row.get("amount_received", 0))
            pending = max(inv_amt - rec_amt, 0.0)
            status  = row.get("status", "Unpaid")

            total_inv  += inv_amt
            total_rec  += rec_amt
            total_pend += pending
            if status == "Overdue":
                overdue_count += 1

            if filter_val != "All" and status != filter_val:
                continue

            self._tree.insert("", "end", iid=str(idx),
                              tags=(status,),
                              values=(
                                  row.get("invoice_date", ""),
                                  row.get("due_date", ""),
                                  row.get("customer_name", ""),
                                  row.get("invoice_number", ""),
                                  row.get("product_service", ""),
                                  f"{inv_amt:,.2f}",
                                  f"{rec_amt:,.2f}",
                                  f"{pending:,.2f}",
                                  status,
                                  row.get("notes", ""),
                              ))

        self._total_invoiced_var.set(f"₹{total_inv:,.2f}")
        self._total_received_var.set(f"₹{total_rec:,.2f}")
        self._total_pending_var.set(f"₹{total_pend:,.2f}")
        self._overdue_count_var.set(str(overdue_count))
        self._editing_index = None

    def _on_row_select(self, _event=None):
        sel = self._tree.selection()
        if not sel:
            return
        idx = int(sel[0])
        rows = self._read_all()
        if idx >= len(rows):
            return
        row = rows[idx]
        self._editing_index = idx
        try:
            self._ent_inv_date.set_date(row.get("invoice_date", date.today().isoformat()))
        except Exception:
            pass
        try:
            self._ent_due_date.set_date(row.get("due_date", date.today().isoformat()))
        except Exception:
            pass
        self._customer_var.set(row.get("customer_name", ""))
        self._ent_inv_no.delete(0, tk.END)
        self._ent_inv_no.insert(0, row.get("invoice_number", ""))
        self._ent_product.delete(0, tk.END)
        self._ent_product.insert(0, row.get("product_service", ""))
        self._ent_inv_amt.delete(0, tk.END)
        self._ent_inv_amt.insert(0, row.get("invoice_amount", ""))
        self._ent_received.delete(0, tk.END)
        self._ent_received.insert(0, row.get("amount_received", ""))
        self._status_var.set(row.get("status", "Unpaid"))
        self._ent_notes.delete(0, tk.END)
        self._ent_notes.insert(0, row.get("notes", ""))

    def clear_form(self):
        self._editing_index = None
        try:
            self._ent_inv_date.set_date(date.today())
            self._ent_due_date.set_date(date.today())
        except Exception:
            pass
        self._customer_var.set("")
        self._status_var.set("Unpaid")
        for w in (self._ent_inv_no, self._ent_product,
                  self._ent_inv_amt, self._ent_received, self._ent_notes):
            w.delete(0, tk.END)
        self._tree.selection_remove(self._tree.selection())

    def save_entry(self):
        customer  = self._customer_var.get().strip()
        inv_no    = self._ent_inv_no.get().strip()
        inv_str   = self._ent_inv_amt.get().strip()
        rec_str   = self._ent_received.get().strip()
        due_str   = self._ent_due_date.get()

        if not customer:
            messagebox.showwarning("Missing Field", "Customer name is required.", parent=self)
            return
        if not inv_no:
            messagebox.showwarning("Missing Field", "Invoice number is required.", parent=self)
            return
        try:
            inv_amt = float(inv_str.replace(",", "") or 0)
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Invoice amount must be a number.", parent=self)
            return
        try:
            rec_amt = float(rec_str.replace(",", "") or 0)
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Amount received must be a number.", parent=self)
            return

        pending = max(inv_amt - rec_amt, 0.0)
        status  = self._status_var.get() or self._auto_status(inv_amt, rec_amt, due_str)

        new_row = {
            "invoice_date":    self._ent_inv_date.get(),
            "due_date":        due_str,
            "customer_name":   customer,
            "invoice_number":  inv_no,
            "product_service": self._ent_product.get().strip(),
            "invoice_amount":  f"{inv_amt:.2f}",
            "amount_received": f"{rec_amt:.2f}",
            "pending_amount":  f"{pending:.2f}",
            "status":          status,
            "notes":           self._ent_notes.get().strip(),
        }

        rows = self._read_all()
        if self._editing_index is not None and self._editing_index < len(rows):
            rows[self._editing_index] = new_row
            msg = "Entry updated."
        else:
            rows.append(new_row)
            msg = "Entry saved."

        self._write_all(rows)
        messagebox.showinfo("Saved", msg, parent=self)
        self.load_entries()

    def delete_entry(self):
        if self._editing_index is None:
            messagebox.showwarning("Select Row", "Select a row to delete.", parent=self)
            return
        rows = self._read_all()
        if self._editing_index >= len(rows):
            return
        row = rows[self._editing_index]
        if messagebox.askyesno(
            "Confirm Delete",
            f"Delete invoice {row.get('invoice_number','')} for {row.get('customer_name','')}?",
            parent=self,
        ):
            rows.pop(self._editing_index)
            self._write_all(rows)
            self.load_entries()

    def _sort_by(self, col):
        rows = self._read_all()
        reverse = self._sort_reverse[col]
        numeric = {"invoice_amount", "amount_received", "pending_amount"}

        def key(r):
            v = r.get(col, "")
            if col in numeric:
                try:
                    return float(str(v).replace(",", "") or 0)
                except ValueError:
                    return 0.0
            return v.lower()

        rows.sort(key=key, reverse=reverse)
        self._sort_reverse[col] = not reverse
        self._write_all(rows)
        self.load_entries()

    # ──────────────────────────────────────────────────── Excel export ───
    def export_excel(self):
        rows = self._read_all()
        if not rows:
            messagebox.showinfo("No Data", "No receivable entries to export.", parent=self)
            return

        path = filedialog.asksaveasfilename(
            parent=self, title="Save Accounts Receivable Report",
            defaultextension=".xlsx",
            initialfile=f"Accounts_Receivable_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return

        try:
            from openpyxl.styles import Font, PatternFill, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
            wb  = Workbook()
            ws  = wb.active
            ws.title = "Receivables"

            ws.append(list(self.HEADINGS))
            for cell in ws[1]:
                cell.font      = Font(bold=True, color="FFFFFF")
                cell.fill      = PatternFill("solid", fgColor="1F618D")
                cell.alignment = Alignment(horizontal="center")

            total_inv = total_rec = total_pend = 0.0
            status_totals = {}
            for row in rows:
                inv = self._to_float(row.get("invoice_amount", 0))
                rec = self._to_float(row.get("amount_received", 0))
                pend = max(inv - rec, 0.0)
                total_inv += inv; total_rec += rec; total_pend += pend
                st = row.get("status", "Unpaid")
                status_totals[st] = status_totals.get(st, 0.0) + pend

                ws.append([
                    row.get("invoice_date", ""), row.get("due_date", ""),
                    row.get("customer_name", ""), row.get("invoice_number", ""),
                    row.get("product_service", ""),
                    inv, rec, pend,
                    st, row.get("notes", ""),
                ])

            # Totals row
            ws.append(["", "", "", "", "TOTAL", total_inv, total_rec, total_pend, "", ""])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)

            # Col widths
            for col, w in zip("ABCDEFGHIJ", [12, 12, 28, 14, 30, 14, 12, 12, 10, 28]):
                ws.column_dimensions[get_column_letter(ord(col) - 64)].width = w
            for row_cells in ws.iter_rows(min_row=2, min_col=6, max_col=8):
                for cell in row_cells:
                    cell.number_format = '#,##0.00'

            # Summary sheet
            ws2 = wb.create_sheet("Summary by Status")
            ws2.append(["Status", "Outstanding (₹)"])
            for cell in ws2[1]:
                cell.font = Font(bold=True)
            for st, amt in sorted(status_totals.items(), key=lambda x: -x[1]):
                ws2.append([st, amt])
            ws2.append(["TOTAL OUTSTANDING", total_pend])
            ws2["B2"].number_format = '#,##0.00'
            ws2.column_dimensions["A"].width = 20
            ws2.column_dimensions["B"].width = 18

            wb.save(path)
            messagebox.showinfo("Exported", f"Accounts Receivable report saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", f"Could not save Excel file:\n{e}", parent=self)


class BankCashView(ttk.Frame):
    """Bank & Cash — manage accounts, balances, deposits, withdrawals, and bank charges."""

    TXN_COLS     = ("date", "account_name", "transaction_type", "amount", "description", "reference", "notes")
    TXN_HEADINGS = ("Date", "Account", "Type", "Amount (₹)", "Description", "Reference", "Notes")
    ACC_COLS     = ("account_name", "bank_name", "account_number", "opening_balance", "current_balance")
    ACC_HEADINGS = ("Account Name", "Bank", "Account Number", "Opening Balance", "Current Balance")

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_bank_files()
        self._txn_editing_index = None
        self._sort_reverse = {c: False for c in self.TXN_COLS}
        self._build()
        self.load_all()

    # ------------------------------------------------------------------ UI --
    def _build(self):
        # ── Header ───────────────────────────────────────────────────────
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="🏦 Bank & Cash", font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel", command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh",        command=self.load_all).pack(side="right", padx=(0, 6))

        # ── Notebook: Accounts | Transactions ────────────────────────────
        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True)

        # ── TAB 1 — Accounts ─────────────────────────────────────────────
        acc_tab = ttk.Frame(nb, padding=10)
        nb.add(acc_tab, text="  🏦 Accounts & Balances  ")

        # Summary row
        self._summary_frame = ttk.Frame(acc_tab)
        self._summary_frame.pack(fill="x", pady=(0, 10))

        # Add / Edit account form
        acc_form = ttk.LabelFrame(acc_tab, text=" Add / Edit Account ", padding=10)
        acc_form.pack(fill="x", pady=(0, 8))

        labels_entries = [
            ("Account Name:",    "_acc_name"),
            ("Bank Name:",       "_acc_bank"),
            ("Account Number:",  "_acc_number"),
            ("Opening Balance:", "_acc_opening"),
        ]
        for col, (lbl, attr) in enumerate(labels_entries):
            ttk.Label(acc_form, text=lbl).grid(row=0, column=col*2, sticky="w", padx=4, pady=4)
            ent = ttk.Entry(acc_form, width=18)
            ent.grid(row=0, column=col*2+1, sticky="ew", padx=4, pady=4)
            setattr(self, attr, ent)

        ttk.Label(acc_form, text="Notes:").grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self._acc_notes = ttk.Entry(acc_form, width=60)
        self._acc_notes.grid(row=1, column=1, columnspan=6, sticky="ew", padx=4, pady=4)

        acc_btns = ttk.Frame(acc_form)
        acc_btns.grid(row=1, column=7, sticky="e", padx=4)
        ttk.Button(acc_btns, text="💾 Save Account",   command=self.save_account).pack(side="left", padx=(0, 4))
        ttk.Button(acc_btns, text="🗑 Delete Account", command=self.delete_account).pack(side="left", padx=(0, 4))
        ttk.Button(acc_btns, text="✖ Clear",           command=self.clear_acc_form).pack(side="left")

        for c in range(8):
            acc_form.columnconfigure(c, weight=1)

        # Accounts treeview
        acc_tree_box = ttk.LabelFrame(acc_tab, text=" Accounts ", padding=6)
        acc_tree_box.pack(fill="both", expand=True)
        self._acc_tree = ttk.Treeview(acc_tree_box, columns=self.ACC_COLS,
                                      show="headings", selectmode="browse", height=6)
        acc_widths = {"account_name": 160, "bank_name": 140, "account_number": 150,
                      "opening_balance": 120, "current_balance": 130}
        for col, heading in zip(self.ACC_COLS, self.ACC_HEADINGS):
            self._acc_tree.heading(col, text=heading)
            self._acc_tree.column(col, width=acc_widths.get(col, 120),
                                  anchor="e" if "balance" in col else "w")
        acc_vsb = ttk.Scrollbar(acc_tree_box, orient="vertical", command=self._acc_tree.yview)
        self._acc_tree.configure(yscrollcommand=acc_vsb.set)
        self._acc_tree.pack(side="left", fill="both", expand=True)
        acc_vsb.pack(side="right", fill="y")
        self._acc_tree.bind("<<TreeviewSelect>>", self._on_acc_select)

        # ── TAB 2 — Transactions ─────────────────────────────────────────
        txn_tab = ttk.Frame(nb, padding=10)
        nb.add(txn_tab, text="  💳 Transactions  ")

        # Transaction form
        txn_form = ttk.LabelFrame(txn_tab, text=" Add / Edit Transaction ", padding=10)
        txn_form.pack(fill="x", pady=(0, 8))

        # Row 0
        ttk.Label(txn_form, text="Date:").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self._txn_date = DateEntry(txn_form, date_pattern="yyyy-mm-dd", width=14)
        self._txn_date.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(txn_form, text="Account:").grid(row=0, column=2, sticky="w", padx=4, pady=4)
        self._txn_acc_var = tk.StringVar()
        self._txn_acc_combo = ttk.Combobox(txn_form, textvariable=self._txn_acc_var,
                                           state="readonly", width=20)
        self._txn_acc_combo.grid(row=0, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(txn_form, text="Type:").grid(row=0, column=4, sticky="w", padx=4, pady=4)
        self._txn_type_var = tk.StringVar(value="Deposit")
        ttk.Combobox(txn_form, textvariable=self._txn_type_var,
                     values=list(BANK_CASH_TRANSACTION_TYPES), state="readonly", width=14
                     ).grid(row=0, column=5, sticky="ew", padx=4, pady=4)

        ttk.Label(txn_form, text="Amount (₹):").grid(row=0, column=6, sticky="w", padx=4, pady=4)
        self._txn_amount = ttk.Entry(txn_form, width=14)
        self._txn_amount.grid(row=0, column=7, sticky="ew", padx=4, pady=4)

        # Row 1
        ttk.Label(txn_form, text="Description:").grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self._txn_desc = ttk.Entry(txn_form, width=28)
        self._txn_desc.grid(row=1, column=1, columnspan=2, sticky="ew", padx=4, pady=4)

        ttk.Label(txn_form, text="Reference:").grid(row=1, column=3, sticky="w", padx=4, pady=4)
        self._txn_ref = ttk.Entry(txn_form, width=18)
        self._txn_ref.grid(row=1, column=4, sticky="ew", padx=4, pady=4)

        ttk.Label(txn_form, text="Notes:").grid(row=1, column=5, sticky="w", padx=4, pady=4)
        self._txn_notes = ttk.Entry(txn_form, width=24)
        self._txn_notes.grid(row=1, column=6, sticky="ew", padx=4, pady=4)

        txn_btns = ttk.Frame(txn_form)
        txn_btns.grid(row=1, column=7, sticky="e", padx=4)
        ttk.Button(txn_btns, text="💾 Save",   command=self.save_transaction).pack(side="left", padx=(0, 4))
        ttk.Button(txn_btns, text="🗑 Delete", command=self.delete_transaction).pack(side="left", padx=(0, 4))
        ttk.Button(txn_btns, text="✖ Clear",  command=self.clear_txn_form).pack(side="left")

        for c in range(8):
            txn_form.columnconfigure(c, weight=1)

        # Transactions treeview
        txn_tree_box = ttk.LabelFrame(txn_tab, text=" Transaction Register ", padding=6)
        txn_tree_box.pack(fill="both", expand=True)
        self._txn_tree = ttk.Treeview(txn_tree_box, columns=self.TXN_COLS,
                                      show="headings", selectmode="browse")
        txn_widths = {"date": 90, "account_name": 150, "transaction_type": 110,
                      "amount": 110, "description": 180, "reference": 110, "notes": 180}
        for col, heading in zip(self.TXN_COLS, self.TXN_HEADINGS):
            self._txn_tree.heading(col, text=heading,
                                   command=lambda c=col: self._sort_by(c))
            self._txn_tree.column(col, width=txn_widths.get(col, 110),
                                  anchor="e" if col == "amount" else "w")
        txn_vsb = ttk.Scrollbar(txn_tree_box, orient="vertical",   command=self._txn_tree.yview)
        txn_hsb = ttk.Scrollbar(txn_tree_box, orient="horizontal", command=self._txn_tree.xview)
        self._txn_tree.configure(yscrollcommand=txn_vsb.set, xscrollcommand=txn_hsb.set)
        self._txn_tree.grid(row=0, column=0, sticky="nsew")
        txn_vsb.grid(row=0, column=1, sticky="ns")
        txn_hsb.grid(row=1, column=0, sticky="ew")
        txn_tree_box.rowconfigure(0, weight=1)
        txn_tree_box.columnconfigure(0, weight=1)
        self._txn_tree.bind("<<TreeviewSelect>>", self._on_txn_select)

    # ─────────────────────────────────────────────── account data helpers ──
    def _read_accounts(self):
        ensure_bank_files()
        with open(BANK_ACCOUNTS_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_accounts(self, rows):
        with open(BANK_ACCOUNTS_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=BANK_ACCOUNTS_HEADERS)
            writer.writeheader()
            writer.writerows(rows)

    def _read_transactions(self):
        ensure_bank_files()
        with open(BANK_TRANSACTIONS_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_transactions(self, rows):
        with open(BANK_TRANSACTIONS_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=BANK_TRANSACTIONS_HEADERS)
            writer.writeheader()
            writer.writerows(rows)

    def _compute_balance(self, account_name):
        """Opening balance ± transactions for a given account."""
        accounts = self._read_accounts()
        acc = next((a for a in accounts if a["account_name"] == account_name), None)
        try:
            balance = float((acc or {}).get("opening_balance", "0").replace(",", "") or 0)
        except ValueError:
            balance = 0.0
        credit_types = {"Deposit", "Cash In"}
        debit_types  = {"Withdrawal", "Cash Out", "Bank Charge"}
        for txn in self._read_transactions():
            if txn.get("account_name") != account_name:
                continue
            try:
                amt = float(str(txn.get("amount", "0")).replace(",", "") or 0)
            except ValueError:
                amt = 0.0
            t = txn.get("transaction_type", "")
            if t in credit_types:
                balance += amt
            elif t in debit_types:
                balance -= amt
        return balance

    # ─────────────────────────────────────────────────── load / refresh ───
    def load_all(self):
        self._load_accounts()
        self._load_transactions()
        self._refresh_account_combo()

    def _load_accounts(self):
        # Summary cards
        for w in self._summary_frame.winfo_children():
            w.destroy()
        accounts = self._read_accounts()
        total_balance = 0.0
        for idx, acc in enumerate(accounts):
            bal = self._compute_balance(acc["account_name"])
            total_balance += bal
            card = ttk.LabelFrame(self._summary_frame,
                                  text=f" {acc['account_name']} ({acc['bank_name']}) ",
                                  padding=(10, 4))
            card.grid(row=0, column=idx, sticky="ew", padx=6)
            ttk.Label(card, text="Current Balance", font=("Helvetica", 8)).pack(anchor="w")
            colour = "green" if bal >= 0 else "red"
            ttk.Label(card, text=f"₹{bal:,.2f}",
                      font=("Helvetica", 13, "bold"),
                      foreground=colour).pack(anchor="w")
            self._summary_frame.columnconfigure(idx, weight=1)

        # Total card
        total_card = ttk.LabelFrame(self._summary_frame, text=" Total Balance ", padding=(10, 4))
        total_card.grid(row=0, column=len(accounts), sticky="ew", padx=6)
        ttk.Label(total_card, text="All Accounts", font=("Helvetica", 8)).pack(anchor="w")
        ttk.Label(total_card, text=f"₹{total_balance:,.2f}",
                  font=("Helvetica", 13, "bold"),
                  foreground="green" if total_balance >= 0 else "red").pack(anchor="w")
        self._summary_frame.columnconfigure(len(accounts), weight=1)

        # Treeview
        for iid in self._acc_tree.get_children():
            self._acc_tree.delete(iid)
        for idx, acc in enumerate(accounts):
            bal = self._compute_balance(acc["account_name"])
            self._acc_tree.insert("", "end", iid=str(idx), values=(
                acc.get("account_name", ""),
                acc.get("bank_name", ""),
                acc.get("account_number", ""),
                f"{float(acc.get('opening_balance', '0') or 0):,.2f}",
                f"{bal:,.2f}",
            ))

    def _load_transactions(self):
        for iid in self._txn_tree.get_children():
            self._txn_tree.delete(iid)
        for idx, row in enumerate(self._read_transactions()):
            try:
                amt = float(str(row.get("amount", "0")).replace(",", "") or 0)
            except ValueError:
                amt = 0.0
            self._txn_tree.insert("", "end", iid=str(idx), values=(
                row.get("date", ""),
                row.get("account_name", ""),
                row.get("transaction_type", ""),
                f"{amt:,.2f}",
                row.get("description", ""),
                row.get("reference", ""),
                row.get("notes", ""),
            ))

    def _refresh_account_combo(self):
        accounts = self._read_accounts()
        names = [a["account_name"] for a in accounts]
        self._txn_acc_combo["values"] = names
        if names and not self._txn_acc_var.get():
            self._txn_acc_var.set(names[0])

    # ──────────────────────────────────────────────────── account actions ──
    def _on_acc_select(self, _event=None):
        sel = self._acc_tree.selection()
        if not sel:
            return
        accounts = self._read_accounts()
        idx = int(sel[0])
        if idx >= len(accounts):
            return
        acc = accounts[idx]
        self._acc_editing_index = idx
        for attr, key in (("_acc_name", "account_name"), ("_acc_bank", "bank_name"),
                          ("_acc_number", "account_number"), ("_acc_opening", "opening_balance"),
                          ("_acc_notes", "notes")):
            w = getattr(self, attr)
            w.delete(0, tk.END)
            w.insert(0, acc.get(key, ""))

    def clear_acc_form(self):
        self._acc_editing_index = None
        for attr in ("_acc_name", "_acc_bank", "_acc_number", "_acc_opening", "_acc_notes"):
            getattr(self, attr).delete(0, tk.END)
        self._acc_tree.selection_remove(self._acc_tree.selection())

    def save_account(self):
        name = self._acc_name.get().strip()
        if not name:
            messagebox.showwarning("Missing Field", "Account name is required.", parent=self)
            return
        opening = self._acc_opening.get().strip()
        try:
            float(opening.replace(",", "") or 0)
        except ValueError:
            messagebox.showwarning("Invalid Balance", "Opening balance must be a number.", parent=self)
            return
        new_row = {
            "account_name":    name,
            "account_number":  self._acc_number.get().strip(),
            "bank_name":       self._acc_bank.get().strip(),
            "opening_balance": opening or "0",
            "notes":           self._acc_notes.get().strip(),
        }
        accounts = self._read_accounts()
        idx = getattr(self, "_acc_editing_index", None)
        if idx is not None and idx < len(accounts):
            accounts[idx] = new_row
            msg = "Account updated."
        else:
            if any(a["account_name"] == name for a in accounts):
                messagebox.showwarning("Duplicate", f"Account '{name}' already exists.", parent=self)
                return
            accounts.append(new_row)
            msg = "Account added."
        self._write_accounts(accounts)
        messagebox.showinfo("Saved", msg, parent=self)
        self.clear_acc_form()
        self.load_all()

    def delete_account(self):
        idx = getattr(self, "_acc_editing_index", None)
        if idx is None:
            messagebox.showwarning("Select Row", "Select an account to delete.", parent=self)
            return
        accounts = self._read_accounts()
        if idx >= len(accounts):
            return
        name = accounts[idx]["account_name"]
        if messagebox.askyesno("Confirm Delete", f"Delete account '{name}'?\nAll its transactions will remain.", parent=self):
            accounts.pop(idx)
            self._write_accounts(accounts)
            self.clear_acc_form()
            self.load_all()

    # ────────────────────────────────────────────── transaction actions ───
    def _on_txn_select(self, _event=None):
        sel = self._txn_tree.selection()
        if not sel:
            return
        idx = int(sel[0])
        txns = self._read_transactions()
        if idx >= len(txns):
            return
        row = txns[idx]
        self._txn_editing_index = idx
        try:
            self._txn_date.set_date(row.get("date", date.today().isoformat()))
        except Exception:
            pass
        self._txn_acc_var.set(row.get("account_name", ""))
        self._txn_type_var.set(row.get("transaction_type", "Deposit"))
        self._txn_amount.delete(0, tk.END)
        self._txn_amount.insert(0, row.get("amount", ""))
        self._txn_desc.delete(0, tk.END)
        self._txn_desc.insert(0, row.get("description", ""))
        self._txn_ref.delete(0, tk.END)
        self._txn_ref.insert(0, row.get("reference", ""))
        self._txn_notes.delete(0, tk.END)
        self._txn_notes.insert(0, row.get("notes", ""))

    def clear_txn_form(self):
        self._txn_editing_index = None
        try:
            self._txn_date.set_date(date.today())
        except Exception:
            pass
        self._txn_type_var.set("Deposit")
        for w in (self._txn_amount, self._txn_desc, self._txn_ref, self._txn_notes):
            w.delete(0, tk.END)
        self._txn_tree.selection_remove(self._txn_tree.selection())

    def save_transaction(self):
        acc = self._txn_acc_var.get().strip()
        txn_type = self._txn_type_var.get().strip()
        amt_str  = self._txn_amount.get().strip()
        if not acc:
            messagebox.showwarning("Missing Field", "Please select an account.", parent=self)
            return
        try:
            amt = float(amt_str.replace(",", "") or 0)
            if amt <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Please enter a valid amount greater than 0.", parent=self)
            return
        new_row = {
            "date":             self._txn_date.get(),
            "account_name":     acc,
            "transaction_type": txn_type,
            "amount":           f"{amt:.2f}",
            "description":      self._txn_desc.get().strip(),
            "reference":        self._txn_ref.get().strip(),
            "notes":            self._txn_notes.get().strip(),
        }
        txns = self._read_transactions()
        if self._txn_editing_index is not None and self._txn_editing_index < len(txns):
            txns[self._txn_editing_index] = new_row
            msg = "Transaction updated."
        else:
            txns.append(new_row)
            msg = "Transaction saved."
        self._write_transactions(txns)
        messagebox.showinfo("Saved", msg, parent=self)
        self.load_all()

    def delete_transaction(self):
        if self._txn_editing_index is None:
            messagebox.showwarning("Select Row", "Select a transaction to delete.", parent=self)
            return
        txns = self._read_transactions()
        if self._txn_editing_index >= len(txns):
            return
        row = txns[self._txn_editing_index]
        if messagebox.askyesno("Confirm Delete",
                               f"Delete {row.get('transaction_type','')} — ₹{row.get('amount','')}?",
                               parent=self):
            txns.pop(self._txn_editing_index)
            self._write_transactions(txns)
            self.load_all()

    def _sort_by(self, col):
        txns = self._read_transactions()
        reverse = self._sort_reverse[col]
        def key(r):
            v = r.get(col, "")
            if col == "amount":
                try:
                    return float(str(v).replace(",", "") or 0)
                except ValueError:
                    return 0.0
            return v.lower()
        txns.sort(key=key, reverse=reverse)
        self._sort_reverse[col] = not reverse
        self._write_transactions(txns)
        self._load_transactions()

    # ──────────────────────────────────────────────────── Excel export ───
    def export_excel(self):
        txns = self._read_transactions()
        accounts = self._read_accounts()
        if not txns and not accounts:
            messagebox.showinfo("No Data", "No data to export.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="Save Bank & Cash Report",
            defaultextension=".xlsx",
            initialfile=f"Bank_Cash_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook()

            # Sheet 1 — Account balances
            ws1 = wb.active
            ws1.title = "Account Balances"
            ws1.append(["Account Name", "Bank", "Account Number", "Opening Balance", "Current Balance"])
            for cell in ws1[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1A5276")
                cell.alignment = Alignment(horizontal="center")
            for acc in accounts:
                bal = self._compute_balance(acc["account_name"])
                op = float(acc.get("opening_balance", "0") or 0)
                ws1.append([acc["account_name"], acc["bank_name"],
                             acc["account_number"], op, bal])
            for col, w in zip("ABCDE", [22, 18, 20, 16, 16]):
                ws1.column_dimensions[col].width = w
            for row_cells in ws1.iter_rows(min_row=2, min_col=4, max_col=5):
                for cell in row_cells:
                    cell.number_format = '#,##0.00'

            # Sheet 2 — Transactions
            ws2 = wb.create_sheet("Transactions")
            ws2.append(list(self.TXN_HEADINGS))
            for cell in ws2[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="117A65")
                cell.alignment = Alignment(horizontal="center")
            for txn in txns:
                try:
                    amt = float(str(txn.get("amount", "0")).replace(",", "") or 0)
                except ValueError:
                    amt = 0.0
                ws2.append([txn.get("date",""), txn.get("account_name",""),
                             txn.get("transaction_type",""), amt,
                             txn.get("description",""), txn.get("reference",""),
                             txn.get("notes","")])
            for col, w in zip("ABCDEFG", [12, 22, 14, 14, 25, 16, 25]):
                ws2.column_dimensions[col].width = w
            for row_cells in ws2.iter_rows(min_row=2, min_col=4, max_col=4):
                for cell in row_cells:
                    cell.number_format = '#,##0.00'

            wb.save(path)
            messagebox.showinfo("Exported", f"Bank & Cash report saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", f"Could not save Excel file:\n{e}", parent=self)


class ExpensesView(ttk.Frame):
    """Expenses register — track all company expenses by category."""

    COLS     = ("date", "category", "description", "amount", "payment_method", "paid_to", "notes")
    HEADINGS = ("Date", "Category", "Description", "Amount (₹)", "Payment Method", "Paid To", "Notes")
    PAYMENT_METHODS = ("Cash", "Bank Transfer", "Cheque", "UPI", "Credit Card", "Other")

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_expenses_file()
        self._editing_index = None
        self._sort_reverse = {c: False for c in self.COLS}
        self._build()
        self.load_entries()

    # ------------------------------------------------------------------ UI --
    def _build(self):
        # ── Header ───────────────────────────────────────────────────────
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 10))
        ttk.Label(header, text="💸 Expenses", font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="⬇ Download Excel", command=self.export_excel).pack(side="right")
        ttk.Button(header, text="🔄 Refresh", command=self.load_entries).pack(side="right", padx=(0, 6))

        # ── Summary cards ─────────────────────────────────────────────────
        summary = ttk.LabelFrame(self, text=" Summary ", padding=10)
        summary.pack(fill="x", pady=(0, 10))
        self._total_var    = tk.StringVar(value="₹0.00")
        self._this_month_var = tk.StringVar(value="₹0.00")
        self._top_cat_var  = tk.StringVar(value="—")
        for col, (lbl, var) in enumerate((
            ("Total Expenses",      self._total_var),
            ("This Month",         self._this_month_var),
            ("Top Category",       self._top_cat_var),
        )):
            card = ttk.Frame(summary, padding=(12, 4))
            card.grid(row=0, column=col, sticky="ew", padx=8)
            ttk.Label(card, text=lbl, font=("Helvetica", 9)).pack(anchor="w")
            ttk.Label(card, textvariable=var, font=("Helvetica", 14, "bold")).pack(anchor="w")
            summary.columnconfigure(col, weight=1)

        # ── Category breakdown bar ────────────────────────────────────────
        cat_box = ttk.LabelFrame(self, text=" Expenses by Category ", padding=8)
        cat_box.pack(fill="x", pady=(0, 10))
        self._cat_frame = ttk.Frame(cat_box)
        self._cat_frame.pack(fill="x")

        # ── Entry form ────────────────────────────────────────────────────
        form_box = ttk.LabelFrame(self, text=" Add / Edit Expense ", padding=10)
        form_box.pack(fill="x", pady=(0, 10))

        # Row 0
        ttk.Label(form_box, text="Date:").grid(row=0, column=0, sticky="w", padx=4, pady=4)
        self._ent_date = DateEntry(form_box, date_pattern="yyyy-mm-dd", width=14)
        self._ent_date.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Category:").grid(row=0, column=2, sticky="w", padx=4, pady=4)
        self._cat_var = tk.StringVar()
        ttk.Combobox(form_box, textvariable=self._cat_var,
                     values=list(EXPENSE_CATEGORIES), state="readonly", width=22
                     ).grid(row=0, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Amount (₹):").grid(row=0, column=4, sticky="w", padx=4, pady=4)
        self._ent_amount = ttk.Entry(form_box, width=14)
        self._ent_amount.grid(row=0, column=5, sticky="ew", padx=4, pady=4)

        # Row 1
        ttk.Label(form_box, text="Description:").grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self._ent_desc = ttk.Entry(form_box, width=30)
        self._ent_desc.grid(row=1, column=1, columnspan=2, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Payment Method:").grid(row=1, column=3, sticky="w", padx=4, pady=4)
        self._pay_var = tk.StringVar(value="Cash")
        ttk.Combobox(form_box, textvariable=self._pay_var,
                     values=list(self.PAYMENT_METHODS), state="readonly", width=14
                     ).grid(row=1, column=4, sticky="ew", padx=4, pady=4)

        ttk.Label(form_box, text="Paid To:").grid(row=1, column=5, sticky="w", padx=4, pady=4)
        self._ent_paid_to = ttk.Entry(form_box, width=18)
        self._ent_paid_to.grid(row=1, column=6, sticky="ew", padx=4, pady=4)

        # Row 2
        ttk.Label(form_box, text="Notes:").grid(row=2, column=0, sticky="w", padx=4, pady=4)
        self._ent_notes = ttk.Entry(form_box, width=50)
        self._ent_notes.grid(row=2, column=1, columnspan=4, sticky="ew", padx=4, pady=4)

        btn_row = ttk.Frame(form_box)
        btn_row.grid(row=2, column=5, columnspan=2, sticky="e", padx=4, pady=4)
        ttk.Button(btn_row, text="💾 Save",   command=self.save_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="🗑 Delete", command=self.delete_entry).pack(side="left", padx=(0, 4))
        ttk.Button(btn_row, text="✖ Clear",  command=self.clear_form).pack(side="left")

        for c in range(7):
            form_box.columnconfigure(c, weight=1)

        # ── Treeview ──────────────────────────────────────────────────────
        table_box = ttk.LabelFrame(self, text=" Expense Register ", padding=8)
        table_box.pack(fill="both", expand=True)

        self._tree = ttk.Treeview(table_box, columns=self.COLS, show="headings", selectmode="browse")
        col_widths = {
            "date": 95, "category": 150, "description": 200,
            "amount": 110, "payment_method": 120, "paid_to": 150, "notes": 200,
        }
        for col, heading in zip(self.COLS, self.HEADINGS):
            self._tree.heading(col, text=heading, command=lambda c=col: self._sort_by(c))
            self._tree.column(col, width=col_widths.get(col, 110),
                              anchor="e" if col == "amount" else "w")

        vsb = ttk.Scrollbar(table_box, orient="vertical",   command=self._tree.yview)
        hsb = ttk.Scrollbar(table_box, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self._tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_box.rowconfigure(0, weight=1)
        table_box.columnconfigure(0, weight=1)

        self._tree.bind("<<TreeviewSelect>>", self._on_row_select)

    # -------------------------------------------------------- data helpers --
    def _read_all(self):
        ensure_expenses_file()
        with open(EXPENSES_CSV, newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    def _write_all(self, rows):
        with open(EXPENSES_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=EXPENSES_HEADERS)
            writer.writeheader()
            writer.writerows(rows)

    # ------------------------------------------------------------ actions --
    def load_entries(self):
        for iid in self._tree.get_children():
            self._tree.delete(iid)

        rows = self._read_all()
        total = 0.0
        this_month = 0.0
        cat_totals = {}
        current_ym = datetime.now().strftime("%Y-%m")

        for idx, row in enumerate(rows):
            try:
                amt = float(str(row.get("amount", "0")).replace(",", "") or 0)
            except ValueError:
                amt = 0.0
            total += amt
            cat = row.get("category", "Other")
            cat_totals[cat] = cat_totals.get(cat, 0.0) + amt
            if str(row.get("date", "")).startswith(current_ym):
                this_month += amt

            self._tree.insert("", "end", iid=str(idx), values=(
                row.get("date", ""),
                cat,
                row.get("description", ""),
                f"{amt:,.2f}",
                row.get("payment_method", ""),
                row.get("paid_to", ""),
                row.get("notes", ""),
            ))

        self._total_var.set(f"₹{total:,.2f}")
        self._this_month_var.set(f"₹{this_month:,.2f}")
        top_cat = max(cat_totals, key=cat_totals.get) if cat_totals else "—"
        self._top_cat_var.set(top_cat)
        self._refresh_category_bars(cat_totals, total)
        self._editing_index = None

    def _refresh_category_bars(self, cat_totals, grand_total):
        for w in self._cat_frame.winfo_children():
            w.destroy()
        if not cat_totals:
            return
        for col, (cat, amt) in enumerate(sorted(cat_totals.items(), key=lambda x: -x[1])):
            pct = (amt / grand_total * 100) if grand_total else 0
            f = ttk.Frame(self._cat_frame)
            f.grid(row=0, column=col, padx=6, pady=2, sticky="n")
            ttk.Label(f, text=cat, font=("Helvetica", 8)).pack()
            ttk.Label(f, text=f"₹{amt:,.0f}", font=("Helvetica", 9, "bold")).pack()
            ttk.Label(f, text=f"{pct:.1f}%", font=("Helvetica", 8),
                      foreground="#555").pack()
        for c in range(len(cat_totals)):
            self._cat_frame.columnconfigure(c, weight=1)

    def _on_row_select(self, _event=None):
        sel = self._tree.selection()
        if not sel:
            return
        idx = int(sel[0])
        rows = self._read_all()
        if idx >= len(rows):
            return
        row = rows[idx]
        self._editing_index = idx
        try:
            self._ent_date.set_date(row.get("date", date.today().isoformat()))
        except Exception:
            pass
        self._cat_var.set(row.get("category", ""))
        self._ent_amount.delete(0, tk.END)
        self._ent_amount.insert(0, row.get("amount", ""))
        self._ent_desc.delete(0, tk.END)
        self._ent_desc.insert(0, row.get("description", ""))
        self._pay_var.set(row.get("payment_method", "Cash"))
        self._ent_paid_to.delete(0, tk.END)
        self._ent_paid_to.insert(0, row.get("paid_to", ""))
        self._ent_notes.delete(0, tk.END)
        self._ent_notes.insert(0, row.get("notes", ""))

    def clear_form(self):
        self._editing_index = None
        try:
            self._ent_date.set_date(date.today())
        except Exception:
            pass
        self._cat_var.set("")
        self._pay_var.set("Cash")
        for w in (self._ent_amount, self._ent_desc,
                  self._ent_paid_to, self._ent_notes):
            w.delete(0, tk.END)
        self._tree.selection_remove(self._tree.selection())

    def save_entry(self):
        cat    = self._cat_var.get().strip()
        amt_str = self._ent_amount.get().strip()
        desc   = self._ent_desc.get().strip()

        if not cat:
            messagebox.showwarning("Missing Field", "Please select a category.", parent=self)
            return
        try:
            amt = float(amt_str.replace(",", "") or 0)
            if amt <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Invalid Amount", "Please enter a valid amount greater than 0.", parent=self)
            return

        new_row = {
            "date":           self._ent_date.get(),
            "category":       cat,
            "description":    desc,
            "amount":         f"{amt:.2f}",
            "payment_method": self._pay_var.get(),
            "paid_to":        self._ent_paid_to.get().strip(),
            "notes":          self._ent_notes.get().strip(),
        }

        rows = self._read_all()
        if self._editing_index is not None and self._editing_index < len(rows):
            rows[self._editing_index] = new_row
            msg = "Expense updated successfully."
        else:
            rows.append(new_row)
            msg = "Expense saved successfully."

        self._write_all(rows)
        messagebox.showinfo("Saved", msg, parent=self)
        self.load_entries()

    def delete_entry(self):
        if self._editing_index is None:
            messagebox.showwarning("Select Row", "Select a row to delete.", parent=self)
            return
        rows = self._read_all()
        if self._editing_index >= len(rows):
            return
        row = rows[self._editing_index]
        if messagebox.askyesno("Confirm Delete",
                               f"Delete expense: {row.get('category','')} — ₹{row.get('amount','')}?",
                               parent=self):
            rows.pop(self._editing_index)
            self._write_all(rows)
            self.load_entries()

    def _sort_by(self, col):
        rows = self._read_all()
        reverse = self._sort_reverse[col]

        def key(r):
            v = r.get(col, "")
            if col == "amount":
                try:
                    return float(str(v).replace(",", "") or 0)
                except ValueError:
                    return 0.0
            return v.lower()

        rows.sort(key=key, reverse=reverse)
        self._sort_reverse[col] = not reverse
        self._write_all(rows)
        self.load_entries()

    def export_excel(self):
        rows = self._read_all()
        if not rows:
            messagebox.showinfo("No Data", "There are no expenses to export.", parent=self)
            return

        path = filedialog.asksaveasfilename(
            parent=self,
            title="Save Expenses Report",
            defaultextension=".xlsx",
            initialfile=f"Expenses_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return

        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook()
            ws = wb.active
            ws.title = "Expenses"

            # Header row
            ws.append(list(self.HEADINGS))
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="C0392B")
                cell.alignment = Alignment(horizontal="center")

            # Data rows
            grand_total = 0.0
            cat_totals = {}
            for row in rows:
                try:
                    amt = float(str(row.get("amount", "0")).replace(",", "") or 0)
                except ValueError:
                    amt = 0.0
                grand_total += amt
                cat = row.get("category", "Other")
                cat_totals[cat] = cat_totals.get(cat, 0.0) + amt
                ws.append([
                    row.get("date", ""), cat,
                    row.get("description", ""), amt,
                    row.get("payment_method", ""),
                    row.get("paid_to", ""),
                    row.get("notes", ""),
                ])

            # Totals row
            ws.append(["", "", "TOTAL", grand_total, "", "", ""])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
            ws.cell(ws.max_row, 4).number_format = '#,##0.00'

            # Category summary sheet
            ws2 = wb.create_sheet("By Category")
            ws2.append(["Category", "Total (₹)"])
            for cell in ws2[1]:
                cell.font = Font(bold=True)
            for cat, amt in sorted(cat_totals.items(), key=lambda x: -x[1]):
                ws2.append([cat, amt])
            ws2.append(["TOTAL", grand_total])
            ws2["A1"].font = Font(bold=True)

            # Column widths
            for col, width in zip("ABCDEFG", [12, 20, 30, 14, 16, 20, 30]):
                ws.column_dimensions[col].width = width
                ws2.column_dimensions[col].width = width

            for row_cells in ws.iter_rows(min_row=2, min_col=4, max_col=4):
                for cell in row_cells:
                    cell.number_format = '#,##0.00'

            wb.save(path)
            messagebox.showinfo("Exported", f"Expenses report saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", f"Could not save Excel file:\n{e}", parent=self)


class AccountsFinanceView(ttk.Frame):
    """Accounts & Finance — landing page with module buttons."""

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=20)
        self.user_data = user_data or {}
        self.navigator = navigator
        self._build()

    def _build(self):
        # Title
        ttk.Label(
            self,
            text="💰 Accounts & Finance",
            font=("Helvetica", 16, "bold"),
        ).pack(pady=(0, 6))

        ttk.Label(
            self,
            text="Select a module to proceed:",
            font=("Helvetica", 11),
        ).pack(pady=(0, 20))

        btn_frame = ttk.Frame(self)
        btn_frame.pack()

        ttk.Button(
            btn_frame,
            text="💰 Sales / Income",
            width=32,
            command=self._open_sales_income,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="💸 Expenses",
            width=32,
            command=self._open_expenses,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="🏦 Bank & Cash",
            width=32,
            command=self._open_bank_cash,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="📥 Accounts Receivable",
            width=32,
            command=self._open_accounts_receivable,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="📤 Accounts Payable",
            width=32,
            command=self._open_accounts_payable,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="📦 Inventory / Stock",
            width=32,
            command=self._open_inventory,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="👨‍💼 Employee Salary",
            width=32,
            command=self._open_salary,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="🧾 GST & Tax",
            width=32,
            command=self._open_gst,
        ).pack(pady=8)

        ttk.Button(
            btn_frame,
            text="📊 Financial Reports",
            width=32,
            command=self._open_financial_reports,
        ).pack(pady=8)

    def _open_sales_income(self):
        self.navigator.show_sales_income_page()

    def _open_expenses(self):
        self.navigator.show_expenses_page()

    def _open_bank_cash(self):
        self.navigator.show_bank_cash_page()

    def _open_accounts_receivable(self):
        self.navigator.show_accounts_receivable_page()

    def _open_accounts_payable(self):
        self.navigator.show_accounts_payable_page()

    def _open_inventory(self):
        self.navigator.show_inventory_page()

    def _open_salary(self):
        self.navigator.show_employee_salary_page()

    def _open_gst(self):
        self.navigator.show_gst_tax_page()

    def _open_financial_reports(self):
        self.navigator.show_financial_reports_page()


class PurchaseProcurementView(ttk.Frame):
    """Purchase / Procurement — inward entries, outward documents, vendor registration, cheque details."""

    def __init__(self, parent, user_data=None, navigator=None):
        super().__init__(parent, padding=14)
        self.user_data = user_data or {}
        self.navigator = navigator
        ensure_inward_file()
        self._build()
        self.load_entries()

    def _build(self):
        header = ttk.Frame(self)
        header.pack(fill="x", pady=(0, 12))
        ttk.Label(header, text="Purchase / Procurement", font=("Helvetica", 16, "bold")).pack(side="left")
        ttk.Button(header, text="+ Inward Entry", command=self.open_inward_entry).pack(side="right")
        ttk.Button(header, text="+ Outward Document", command=self.open_outward_window).pack(side="right", padx=(0, 8))
        ttk.Button(header, text="+ Vendor Registration", command=self.open_vendor_registration).pack(side="right", padx=(0, 8))
        ttk.Button(header, text="Cheque Details", command=self.open_cheque_details).pack(side="right", padx=(0, 8))

        summary = ttk.LabelFrame(self, text=" Inward Purchase Summary ", padding=12)
        summary.pack(fill="x", pady=(0, 12))
        self.total_entries_var = tk.StringVar(value="0")
        self.total_amount_var  = tk.StringVar(value="0.00")
        self.pending_amount_var = tk.StringVar(value="0.00")
        for index, (label, variable) in enumerate((
            ("Product Lines",        self.total_entries_var),
            ("Total Inward Amount",  self.total_amount_var),
            ("Pending Amount",       self.pending_amount_var),
        )):
            card = ttk.Frame(summary, padding=(12, 3))
            card.grid(row=0, column=index, sticky="ew", padx=8)
            ttk.Label(card, text=label).pack(anchor="w")
            ttk.Label(card, textvariable=variable, font=("Helvetica", 14, "bold")).pack(anchor="w")
            summary.columnconfigure(index, weight=1)

        pending_box = ttk.LabelFrame(self, text=" Pending Invoices ", padding=8)
        pending_box.pack(fill="x", pady=(0, 12))
        pending_columns = ("supplier", "invoice", "department", "total", "paid", "balance")
        self.pending_tree = ttk.Treeview(pending_box, columns=pending_columns, show="headings", height=4, selectmode="browse")
        for column, heading in zip(pending_columns, ("Vendor", "Invoice No.", "Department", "Invoice Total", "Paid", "Pending Balance")):
            self.pending_tree.heading(column, text=heading)
            self.pending_tree.column(column, width=175 if column == "supplier" else 120,
                                     anchor="e" if column in ("total", "paid", "balance") else "w")
        self.pending_tree.pack(side="left", fill="x", expand=True)
        pending_actions = ttk.Frame(pending_box)
        pending_actions.pack(side="right", fill="y", padx=(8, 0))
        ttk.Button(pending_actions, text="Pay Selected\nInvoice", command=self.pay_selected_invoice).pack(fill="x")

        table_box = ttk.LabelFrame(self, text=" Inward Entry Register ", padding=8)
        table_box.pack(fill="both", expand=True)
        columns  = ("date", "supplier", "invoice_no", "department", "item_description", "quantity", "rate", "amount", "payment_status")
        headings = ("Date", "Supplier", "Invoice No.", "Department", "Product / Material", "Qty", "Rate", "Amount", "Status")
        self.tree = ttk.Treeview(table_box, columns=columns, show="headings", selectmode="browse")
        for column, heading in zip(columns, headings):
            self.tree.heading(column, text=heading)
            self.tree.column(column, width=180 if column in ("supplier", "item_description") else 115,
                             anchor="e" if column in ("quantity", "rate", "amount") else "w")
        scrollbar = ttk.Scrollbar(table_box, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        footer = ttk.Frame(self)
        footer.pack(fill="x", pady=(10, 0))
        ttk.Button(footer, text="Refresh", command=self.load_entries).pack(side="right")

    def open_inward_entry(self):
        self.navigator.show_inward_entry_page()

    def open_outward_window(self):
        self.navigator.show_outward_document_page()

    def open_vendor_registration(self):
        self.navigator.show_vendor_registration_page()

    def open_cheque_details(self):
        self.navigator.show_cheque_details_page()

    def pay_selected_invoice(self):
        selected = self.pending_tree.selection()
        if not selected:
            messagebox.showwarning("Select Invoice", "Select a pending invoice first.", parent=self)
            return
        invoice = next(
            (item for item in get_invoice_balances()
             if f"{item['supplier']}|{item['invoice_no']}" == selected[0]),
            None,
        )
        if invoice:
            InvoicePaymentDialog(self, invoice, self.load_entries)

    def load_entries(self):
        for row_id in self.tree.get_children():
            self.tree.delete(row_id)
        total = 0.0
        count = 0
        with open(INWARD_CSV, newline="", encoding="utf-8") as file:
            for row in csv.DictReader(file):
                count += 1
                try:
                    amount = float(row.get("amount", 0) or 0)
                except ValueError:
                    amount = 0.0
                total += amount
                self.tree.insert("", "end", values=(
                    row.get("date", ""), row.get("supplier", ""), row.get("invoice_no", ""),
                    row.get("department", ""), row.get("item_description", ""),
                    row.get("quantity", ""), row.get("rate", ""),
                    f"{amount:.2f}", row.get("payment_status", ""),
                ))
        self.total_entries_var.set(str(count))
        self.total_amount_var.set(f"{total:,.2f}")

        for row_id in self.pending_tree.get_children():
            self.pending_tree.delete(row_id)
        pending_total = 0.0
        for invoice in get_invoice_balances():
            if invoice["balance"] <= 0:
                continue
            pending_total += invoice["balance"]
            item_id = f"{invoice['supplier']}|{invoice['invoice_no']}"
            self.pending_tree.insert("", "end", iid=item_id, values=(
                invoice["supplier"], invoice["invoice_no"], invoice["department"],
                f"{invoice['total']:,.2f}", f"{invoice['paid']:,.2f}", f"{invoice['balance']:,.2f}",
            ))
        self.pending_amount_var.set(f"{pending_total:,.2f}")
