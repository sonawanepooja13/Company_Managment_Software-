"""Customer-wise quotation storage.

Every customer gets a dedicated CSV file inside ``csv_data/quotation``
(for example ``csv_data/quotation/Sunny_Hapse.csv``) so that all quotations
belonging to one customer stay together and can be shared with that customer.

Each row of a file is one quoted item; the quotation header fields (quotation
number, dates, status, totals ...) repeat on every row of the same quotation,
which keeps the file readable and spreadsheet friendly.

This module holds the data layer only - the Tkinter screens live in
``sales_marketing_window.py``.
"""

import csv
import json
import os
import re
from datetime import date, datetime, timedelta

import config

QUOTATION_DIR = getattr(config, "QUOTATION_DIR", os.path.join(config.CSV_DIR, "quotation"))
EXPORT_DIR = os.path.join(QUOTATION_DIR, "exports")

# Company bank details (printed on every quotation) and the reusable
# Terms & Conditions / Additional Notes library.
BANK_DETAILS_CSV = os.path.join(QUOTATION_DIR, "bank_details.csv")
TERMS_LIBRARY_CSV = os.path.join(QUOTATION_DIR, "terms_library.csv")

BANK_DETAIL_FIELDS = (
    "bank_name",
    "account_name",
    "account_number",
    "ifsc_code",
    "branch",
    "upi_id",
)
BANK_DETAIL_LABELS = {
    "bank_name": "Bank Name",
    "account_name": "Account Name",
    "account_number": "Account No.",
    "ifsc_code": "IFSC Code",
    "branch": "Branch",
    "upi_id": "UPI ID",
}
TERMS_LIBRARY_HEADERS = ["title", "detail"]

STATUSES = ("Draft", "Sent", "Accepted", "Rejected", "Converted")
DEFAULT_TAX_PERCENT = 18.0
DEFAULT_VALIDITY_DAYS = 15

QUOTATION_HEADERS = [
    "quotation_no",
    "quotation_date",
    "valid_until",
    "status",
    "customer_name",
    "customer_contact_person",
    "customer_email",
    "customer_phone",
    "customer_gstin",
    "customer_address",
    "customer_type",
    "challan_no",
    "challan_date",
    "lr_no",
    "delivery_mode",
    "rev_charge",
    "ship_to",
    "distance_for_eway_bill",
    "place_of_supply",
    "item",
    "description",
    "hsn_code",
    "capacity_unit",
    "quantity",
    "unit_price",
    "discount_percent",
    "discount_amount",
    "amount",
    "subtotal",
    "discount_total",
    "tax_percent",
    "tax_amount",
    "grand_total",
    "terms_json",
    "notes",
    "created_at",
    "created_by",
]

CUSTOMERS_DETAILED_CSV = getattr(
    config, "CUSTOMERS_DETAILED_CSV", os.path.join(config.CSV_DIR, "customers_detailed.csv")
)
SIMPLE_CUSTOMERS_CSV = getattr(
    config, "CUSTOMERS_CSV", os.path.join(config.CSV_DIR, "customers.csv")
)


def ensure_quotation_dir():
    """Make sure csv_data/quotation exists and return its path."""
    os.makedirs(QUOTATION_DIR, exist_ok=True)
    return QUOTATION_DIR


def safe_customer_stem(customer_name):
    """Turn a customer name into a safe CSV file name (without extension)."""
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "_", (customer_name or "").strip())
    stem = re.sub(r"[\s_]+", "_", stem).strip("._-")
    return stem or "Unnamed_Customer"


def customer_quotation_file(customer_name):
    """Return the customer-wise CSV path for a customer name."""
    ensure_quotation_dir()
    return os.path.join(QUOTATION_DIR, f"{safe_customer_stem(customer_name)}.csv")


def customer_quotation_files():
    """Return all customer-wise quotation CSV files."""
    ensure_quotation_dir()
    return sorted(
        os.path.join(QUOTATION_DIR, name)
        for name in os.listdir(QUOTATION_DIR)
        if name.lower().endswith(".csv")
    )


def load_customers():
    """Customer names from the CRM data (falls back to the simple list)."""
    names = set()
    for path in (CUSTOMERS_DETAILED_CSV, SIMPLE_CUSTOMERS_CSV):
        if not os.path.exists(path):
            continue
        try:
            with open(path, newline="", encoding="utf-8-sig") as file:
                for row in csv.DictReader(file):
                    name = (row.get("company_name") or row.get("customer_name") or "").strip()
                    if name:
                        names.add(name)
        except OSError:
            continue
    return sorted(names, key=str.lower)


def _customer_detail_value(row, *keys):
    """Return the first non-empty customer detail from a CRM row."""
    for key in keys:
        value = str(row.get(key) or "").strip()
        if value:
            return value
    return ""


def _company_key(value):
    """Normalise a company name so lookups ignore case and extra spaces."""
    return " ".join(str(value or "").split()).casefold()


def load_customer_details(customer_name):
    """Return saved CRM details for one company, or an empty dictionary.

    If a company has more than one CRM row, the first non-empty value for
    each field is retained. This supports duplicate company/contact records
    without leaving available details blank.

    The CRM keeps the company tax id in its ``GST Number`` column, which is
    exactly what the quotation card shows as ``GSTIN / PAN`` - the CRM GST
    data is therefore returned as ``customer_gstin``.
    """
    wanted = _company_key(customer_name)
    if not wanted or not os.path.exists(CUSTOMERS_DETAILED_CSV):
        return {}

    details = {
        "customer_contact_person": "",
        "customer_email": "",
        "customer_phone": "",
        "customer_gstin": "",
        "customer_address": "",
    }
    field_sources = {
        "customer_contact_person": ("contact_person", "contact", "primary_contact"),
        "customer_email": ("email", "customer_email", "email_address"),
        "customer_phone": ("contact_number", "phone", "mobile", "contact"),
        "customer_gstin": ("gst_number", "gstin", "gst", "pan"),
        "customer_address": ("address", "customer_address"),
    }
    matched = False
    try:
        with open(CUSTOMERS_DETAILED_CSV, newline="", encoding="utf-8-sig") as file:
            for row in csv.DictReader(file):
                company = row.get("company_name") or row.get("customer_name") or ""
                if _company_key(company) != wanted:
                    continue
                matched = True
                for field, source_keys in field_sources.items():
                    if not details[field]:
                        details[field] = _customer_detail_value(row, *source_keys)
    except OSError:
        return {}
    return details if matched else {}


def load_customer_gstin(customer_name):
    """Return the saved CRM GST data for one company (the GSTIN / PAN value).

    The CRM ``GST Number`` column and the quotation ``GSTIN / PAN`` field hold
    the same tax id, so the CRM GST data can be shown directly in GSTIN / PAN.
    An empty string is returned when the company is unknown or has no GST.
    """
    return load_customer_details(customer_name).get("customer_gstin", "")


def _price_list_path():
    for path in (
        os.path.join(config.SCRIPT_DIR, "price_list_clean.csv"),
        os.path.join(config.CSV_DIR, "price_list_clean.csv"),
    ):
        if os.path.exists(path):
            return path
    return os.path.join(config.SCRIPT_DIR, "price_list_clean.csv")


def _clean_rate(row):
    """Net rate for a price list row (DP when available, else List Price)."""
    for key in ("DP", "List Price"):
        raw = str(row.get(key) or "").replace(",", "").replace("₹", "").strip()
        if not raw:
            continue
        try:
            value = float(raw)
        except ValueError:
            continue
        if value > 0:
            return value
    return 0.0


def load_products():
    """Products / services that can be quoted, taken from the price list."""
    products = []
    path = _price_list_path()
    if not os.path.exists(path):
        return products
    with open(path, newline="", encoding="utf-8-sig") as file:
        for row in csv.DictReader(file):
            item = (row.get("Item Name") or "").strip()
            if not item:
                continue
            capacity = (row.get("CAPACITY") or "").strip()
            unit = (row.get("Unit") or "").strip()
            supplier = (row.get("Supplier") or "").strip()
            label = " | ".join(part for part in (item, capacity, unit, supplier) if part)
            products.append(
                {
                    "label": f"{label}  (Rs {_clean_rate(row):.2f})",
                    "item": item,
                    "capacity_unit": " ".join(part for part in (capacity, unit) if part),
                    "rate": _clean_rate(row),
                }
            )
    return products


def read_rows(path):
    """Read one quotation CSV file (missing file -> empty list)."""
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        return [dict(row) for row in reader]


def write_rows(path, rows):
    """(Re)write a quotation CSV file with the canonical header row."""
    ensure_quotation_dir()
    with open(path, "w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=QUOTATION_HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in QUOTATION_HEADERS})


def try_float(value, default=0.0):
    try:
        return float(str(value).replace(",", "").replace("₹", "").strip())
    except (TypeError, ValueError):
        return default


def calculate_totals(items, tax_percent=DEFAULT_TAX_PERCENT):
    """Totals for a list of item dicts, applying every item's discount.

    Returns a dict with subtotal (before discount), discount_total, net_total,
    tax_percent, tax_amount and grand_total.
    """
    subtotal = 0.0
    discount_total = 0.0
    for item in items:
        quantity, unit_price, _percent, discount, _amount = line_amounts(item)
        subtotal += quantity * unit_price
        discount_total += discount
    net_total = subtotal - discount_total
    tax_percent = try_float(tax_percent, 0.0)
    tax_amount = net_total * tax_percent / 100.0
    return {
        "subtotal": round(subtotal, 2),
        "discount_total": round(discount_total, 2),
        "net_total": round(net_total, 2),
        "tax_percent": tax_percent,
        "tax_amount": round(tax_amount, 2),
        "grand_total": round(net_total + tax_amount, 2),
    }


def line_amounts(item):
    """Return (quantity, unit_price, discount_percent, discount_amount, amount)."""
    quantity = try_float(item.get("quantity"), 0.0)
    unit_price = try_float(item.get("unit_price"), 0.0)
    gross = quantity * unit_price
    percent = min(max(try_float(item.get("discount_percent"), 0.0), 0.0), 100.0)
    discount = min(max(try_float(item.get("discount_amount"), gross * percent / 100.0), 0.0), gross)
    return quantity, unit_price, percent, discount, gross - discount


def next_quotation_number():
    """Next free quotation number across every customer file (QT-0001, ...)."""
    highest = 0
    for path in customer_quotation_files():
        for row in read_rows(path):
            match = re.search(r"(\d+)\s*$", (row.get("quotation_no") or "").strip())
            if match:
                highest = max(highest, int(match.group(1)))
    return f"QT-{highest + 1:04d}"




def build_quotation_rows(
    quotation_no,
    customer_name,
    items,
    customer_email="",
    customer_phone="",
    customer_gstin="",
    customer_address="",
    customer_contact_person="",
    customer_type="Quotation",
    challan_no="",
    challan_date="",
    lr_no="",
    delivery_mode="",
    rev_charge="No",
    ship_to="",
    distance_for_eway_bill="",
    place_of_supply="",
    quotation_date=None,
    valid_until=None,
    status="Draft",
    tax_percent=DEFAULT_TAX_PERCENT,
    terms=None,
    notes="",
    created_at=None,
    created_by="",
):
    """Build the CSV rows (one per item) of a single quotation.

    Returns (rows, totals).  Every row carries all header fields (customer,
    GSTIN/PAN, status, totals, terms ...) so the customer file stays readable
    and spreadsheet friendly.
    """
    quotation_date = (quotation_date or date.today().isoformat()).strip()
    if not valid_until:
        try:
            base = datetime.strptime(quotation_date, "%Y-%m-%d").date()
        except ValueError:
            base = date.today()
        valid_until = (base + timedelta(days=DEFAULT_VALIDITY_DAYS)).isoformat()

    totals = calculate_totals(items, tax_percent)
    clean_terms = [
        {"title": (term.get("title") or "").strip(), "detail": (term.get("detail") or "").strip()}
        for term in (terms or [])
        if (term.get("title") or "").strip() or (term.get("detail") or "").strip()
    ]
    terms_json = json.dumps(clean_terms, ensure_ascii=False)
    created_at = created_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    rows = []
    for item in items:
        quantity, unit_price, percent, discount, amount = line_amounts(item)
        rows.append(
            {
                "quotation_no": quotation_no,
                "quotation_date": quotation_date,
                "valid_until": (valid_until or "").strip(),
                "status": (status or "Draft").strip() or "Draft",
                "customer_name": customer_name,
                "customer_email": (customer_email or "").strip(),
                "customer_phone": (customer_phone or "").strip(),
                "customer_gstin": (customer_gstin or "").strip(),
                "customer_address": (customer_address or "").strip(),
                "customer_contact_person": (customer_contact_person or "").strip(),
                "customer_type": (customer_type or "Quotation").strip(),
                "challan_no": (challan_no or "").strip(),
                "challan_date": (challan_date or "").strip(),
                "lr_no": (lr_no or "").strip(),
                "delivery_mode": (delivery_mode or "").strip(),
                "rev_charge": (rev_charge or "No").strip(),
                "ship_to": (ship_to or "").strip(),
                "distance_for_eway_bill": (distance_for_eway_bill or "").strip(),
                "place_of_supply": (place_of_supply or "").strip(),
                "item": (item.get("item") or "").strip(),
                "description": (item.get("description") or "").strip(),
                "hsn_code": (item.get("hsn_code") or "").strip(),
                "capacity_unit": (item.get("capacity_unit") or "").strip(),
                "quantity": f"{quantity:g}",
                "unit_price": f"{unit_price:.2f}",
                "discount_percent": f"{percent:g}",
                "discount_amount": f"{discount:.2f}",
                "amount": f"{amount:.2f}",
                "subtotal": f"{totals['subtotal']:.2f}",
                "discount_total": f"{totals['discount_total']:.2f}",
                "tax_percent": f"{totals['tax_percent']:g}",
                "tax_amount": f"{totals['tax_amount']:.2f}",
                "grand_total": f"{totals['grand_total']:.2f}",
                "terms_json": terms_json,
                "notes": (notes or "").strip(),
                "created_at": created_at,
                "created_by": (created_by or "").strip(),
            }
        )
    return rows, totals


def save_quotation(
    customer_name,
    items,
    customer_email="",
    customer_phone="",
    customer_gstin="",
    customer_address="",
    customer_contact_person="",
    customer_type="Quotation",
    challan_no="",
    challan_date="",
    lr_no="",
    delivery_mode="",
    rev_charge="No",
    ship_to="",
    distance_for_eway_bill="",
    place_of_supply="",
    quotation_date=None,
    valid_until=None,
    status="Draft",
    tax_percent=DEFAULT_TAX_PERCENT,
    terms=None,
    notes="",
    quotation_no=None,
    created_by="",
):
    """Append a quotation to the customer's own CSV file.

    ``items`` is a list of dicts with the keys item, description, hsn_code,
    capacity_unit, quantity, unit_price and discount_percent; amounts are
    calculated here.  ``terms`` is a list of {"title", "detail"} notes.

    Returns a dict with quotation_no, customer_file, rows_written and totals.
    """
    customer_name = (customer_name or "").strip()
    if not customer_name:
        raise ValueError("Customer name is required.")
    if not items:
        raise ValueError("Add at least one item to the quotation.")

    quotation_no = (quotation_no or "").strip() or next_quotation_number()
    path = customer_quotation_file(customer_name)
    rows, totals = build_quotation_rows(
        quotation_no=quotation_no,
        customer_name=customer_name,
        items=items,
        customer_email=customer_email,
        customer_phone=customer_phone,
        customer_gstin=customer_gstin,
        customer_address=customer_address,
        customer_contact_person=customer_contact_person,
        customer_type=customer_type,
        challan_no=challan_no,
        challan_date=challan_date,
        lr_no=lr_no,
        delivery_mode=delivery_mode,
        rev_charge=rev_charge,
        ship_to=ship_to,
        distance_for_eway_bill=distance_for_eway_bill,
        place_of_supply=place_of_supply,
        quotation_date=quotation_date,
        valid_until=valid_until,
        status=status,
        tax_percent=tax_percent,
        terms=terms,
        notes=notes,
        created_by=created_by,
    )

    write_rows(path, read_rows(path) + rows)

    return {
        "quotation_no": quotation_no,
        "customer_file": path,
        "rows_written": len(rows),
        "subtotal": totals["subtotal"],
        "discount_total": totals["discount_total"],
        "tax_amount": totals["tax_amount"],
        "grand_total": totals["grand_total"],
    }


def file_with_quotation(quotation_no):
    """Path of the customer file that holds the given quotation (or None)."""
    number = (quotation_no or "").strip()
    if not number:
        return None
    for path in customer_quotation_files():
        for row in read_rows(path):
            if (row.get("quotation_no") or "").strip() == number:
                return path
    return None


def update_quotation(quotation_no, **fields):
    """Replace an existing quotation in place (used by Edit Quotation).

    Accepts the same keyword fields as :func:`save_quotation`.  When the
    customer name changes, the quotation is moved to the new customer file.
    """
    number = (quotation_no or "").strip()
    if not number:
        raise ValueError("A quotation number is required to update a quotation.")

    old_path = file_with_quotation(number)
    if old_path is None:
        raise ValueError(f"Quotation {number} was not found.")

    old_rows = read_rows(old_path)
    first_old = next(
        row for row in old_rows if (row.get("quotation_no") or "").strip() == number
    )
    remaining = [
        row for row in old_rows if (row.get("quotation_no") or "").strip() != number
    ]

    customer_name = (fields.get("customer_name") or "").strip()
    if not customer_name:
        raise ValueError("Customer name is required.")
    items = fields.get("items") or []
    if not items:
        raise ValueError("Add at least one item to the quotation.")

    new_path = customer_quotation_file(customer_name)
    rows, totals = build_quotation_rows(
        quotation_no=number,
        customer_name=customer_name,
        items=items,
        customer_email=fields.get("customer_email", first_old.get("customer_email", "")),
        customer_phone=fields.get("customer_phone", first_old.get("customer_phone", "")),
        customer_gstin=fields.get("customer_gstin", first_old.get("customer_gstin", "")),
        customer_address=fields.get("customer_address", first_old.get("customer_address", "")),
        customer_contact_person=fields.get("customer_contact_person", first_old.get("customer_contact_person", "")),
        customer_type=fields.get("customer_type", first_old.get("customer_type", "Quotation")),
        challan_no=fields.get("challan_no", first_old.get("challan_no", "")),
        challan_date=fields.get("challan_date", first_old.get("challan_date", "")),
        lr_no=fields.get("lr_no", first_old.get("lr_no", "")),
        delivery_mode=fields.get("delivery_mode", first_old.get("delivery_mode", "")),
        rev_charge=fields.get("rev_charge", first_old.get("rev_charge", "No")),
        ship_to=fields.get("ship_to", first_old.get("ship_to", "")),
        distance_for_eway_bill=fields.get("distance_for_eway_bill", first_old.get("distance_for_eway_bill", "")),
        place_of_supply=fields.get("place_of_supply", first_old.get("place_of_supply", "")),
        quotation_date=fields.get("quotation_date", first_old.get("quotation_date")),
        valid_until=fields.get("valid_until", first_old.get("valid_until")),
        status=fields.get("status", first_old.get("status", "Draft")),
        tax_percent=fields.get("tax_percent", first_old.get("tax_percent", DEFAULT_TAX_PERCENT)),
        terms=fields.get("terms"),
        notes=fields.get("notes", first_old.get("notes", "")),
        created_at=first_old.get("created_at"),
        created_by=fields.get("created_by") or first_old.get("created_by", ""),
    )

    if new_path == old_path:
        write_rows(new_path, remaining + rows)
    else:
        write_rows(old_path, remaining)
        write_rows(new_path, read_rows(new_path) + rows)

    return {
        "quotation_no": number,
        "customer_file": new_path,
        "moved_from": None if new_path == old_path else old_path,
        "rows_written": len(rows),
        "subtotal": totals["subtotal"],
        "discount_total": totals["discount_total"],
        "tax_amount": totals["tax_amount"],
        "grand_total": totals["grand_total"],
    }


def parse_terms(terms_json):
    """Terms & Conditions stored in a CSV cell, as a list of dicts."""
    if not terms_json:
        return []
    try:
        data = json.loads(terms_json)
    except (TypeError, ValueError):
        return []
    if not isinstance(data, list):
        return []
    return [
        {"title": str(term.get("title") or ""), "detail": str(term.get("detail") or "")}
        for term in data
        if isinstance(term, dict)
    ]


def load_bank_details():
    """Company bank details shown on the printed quotation."""
    details = {field: "" for field in BANK_DETAIL_FIELDS}
    rows = read_rows(BANK_DETAILS_CSV)
    if rows:
        for field in BANK_DETAIL_FIELDS:
            details[field] = (rows[0].get(field) or "").strip()
    return details


def save_bank_details(details):
    """Store the company bank details (one row, overwritten each time)."""
    ensure_quotation_dir()
    with open(BANK_DETAILS_CSV, "w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=list(BANK_DETAIL_FIELDS))
        writer.writeheader()
        writer.writerow(
            {field: str(details.get(field) or "").strip() for field in BANK_DETAIL_FIELDS}
        )
    return BANK_DETAILS_CSV


def load_terms_library():
    """Previously saved Terms & Conditions / Additional Notes.

    Used by the "+ Add Note" picker so old terms can be reused.
    """
    terms = []
    seen = set()
    for row in read_rows(TERMS_LIBRARY_CSV):
        title = (row.get("title") or "").strip()
        detail = (row.get("detail") or "").strip()
        if not title and not detail:
            continue
        key = (title.lower(), detail.lower())
        if key in seen:
            continue
        seen.add(key)
        terms.append({"title": title, "detail": detail})
    return terms


def add_terms_to_library(terms):
    """Remember the given terms so they can be reused on the next quotation."""
    existing = load_terms_library()
    known = {(term["title"].lower(), term["detail"].lower()) for term in existing}
    fresh = []
    for term in terms or []:
        title = (term.get("title") or "").strip()
        detail = (term.get("detail") or "").strip()
        if not title and not detail:
            continue
        key = (title.lower(), detail.lower())
        if key in known:
            continue
        known.add(key)
        fresh.append({"title": title, "detail": detail})
    if not fresh:
        return 0
    ensure_quotation_dir()
    is_new = not os.path.exists(TERMS_LIBRARY_CSV) or os.path.getsize(TERMS_LIBRARY_CSV) == 0
    with open(TERMS_LIBRARY_CSV, "a", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=TERMS_LIBRARY_HEADERS)
        if is_new:
            writer.writeheader()
        writer.writerows(fresh)
    return len(fresh)


def list_quotations(customer_name=None, status=None):
    """Summaries of every stored quotation (one entry per quotation number).

    Each summary holds the header fields plus an ``items`` count and the
    ``file`` it was read from, ready for the quotation register table.
    """
    wanted_customer = (customer_name or "").strip().lower()
    wanted_status = (status or "").strip().lower()
    summaries = {}

    for path in customer_quotation_files():
        for row in read_rows(path):
            number = (row.get("quotation_no") or "").strip()
            if not number:
                continue
            name = (row.get("customer_name") or "").strip()
            if wanted_customer and wanted_customer not in name.lower():
                continue
            if wanted_status and (row.get("status") or "").strip().lower() != wanted_status:
                continue
            entry = summaries.get(number)
            if entry is None:
                entry = {
                    "quotation_no": number,
                    "quotation_date": (row.get("quotation_date") or "").strip(),
                    "valid_until": (row.get("valid_until") or "").strip(),
                    "status": (row.get("status") or "").strip(),
                    "customer_name": name,
                    "customer_email": (row.get("customer_email") or "").strip(),
                    "customer_phone": (row.get("customer_phone") or "").strip(),
                    "customer_gstin": (row.get("customer_gstin") or "").strip(),
                    "customer_address": (row.get("customer_address") or "").strip(),
                    "customer_contact_person": (row.get("customer_contact_person") or "").strip(),
                    "customer_type": (row.get("customer_type") or "Quotation").strip(),
                    "challan_no": (row.get("challan_no") or "").strip(),
                    "challan_date": (row.get("challan_date") or "").strip(),
                    "lr_no": (row.get("lr_no") or "").strip(),
                    "delivery_mode": (row.get("delivery_mode") or "").strip(),
                    "rev_charge": (row.get("rev_charge") or "No").strip(),
                    "ship_to": (row.get("ship_to") or "").strip(),
                    "distance_for_eway_bill": (row.get("distance_for_eway_bill") or "").strip(),
                    "place_of_supply": (row.get("place_of_supply") or "").strip(),
                    "notes": (row.get("notes") or "").strip(),
                    "terms": parse_terms(row.get("terms_json")),
                    "subtotal": try_float(row.get("subtotal"), 0.0),
                    "discount_total": try_float(row.get("discount_total"), 0.0),
                    "net_total": round(
                        try_float(row.get("subtotal"), 0.0)
                        - try_float(row.get("discount_total"), 0.0),
                        2,
                    ),
                    "tax_percent": try_float(row.get("tax_percent"), 0.0),
                    "tax_amount": try_float(row.get("tax_amount"), 0.0),
                    "grand_total": try_float(row.get("grand_total"), 0.0),
                    "items": 0,
                    "rows": [],
                    "file": path,
                }
                summaries[number] = entry
            entry["items"] += 1
            entry["rows"].append(row)

    return sorted(
        summaries.values(),
        key=lambda entry: (entry["quotation_date"], entry["quotation_no"]),
        reverse=True,
    )


def get_quotation(quotation_no):
    """Full details (header + item rows) of one quotation, or None."""
    number = (quotation_no or "").strip()
    if not number:
        return None
    for summary in list_quotations():
        if summary["quotation_no"] == number:
            return summary
    return None


def update_status(quotation_no, status):
    """Set the status of a quotation in every file that holds it."""
    number = (quotation_no or "").strip()
    status = (status or "").strip()
    if not number or not status:
        return 0
    changed = 0
    for path in customer_quotation_files():
        rows = read_rows(path)
        touched = False
        for row in rows:
            if (row.get("quotation_no") or "").strip() == number:
                row["status"] = status
                changed += 1
                touched = True
        if touched:
            write_rows(path, rows)
    return changed




def export_quotation(quotation_no):
    """Write one quotation as Excel + printable PDF. Returns their paths."""
    summary = get_quotation(quotation_no)
    if summary is None:
        raise ValueError(f"Quotation {quotation_no} was not found.")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    target = os.path.join(EXPORT_DIR, safe_customer_stem(summary["customer_name"]))
    os.makedirs(target, exist_ok=True)
    base = f"{safe_customer_stem(quotation_no)}_{datetime.now():%Y%m%d_%H%M%S}"
    xlsx_path = os.path.join(target, f"{base}.xlsx")
    pdf_path = os.path.join(target, f"{base}.pdf")

    bank = load_bank_details()
    bank_lines = [
        (BANK_DETAIL_LABELS[field], bank[field]) for field in BANK_DETAIL_FIELDS if bank.get(field)
    ]
    terms = summary.get("terms") or []
    header = [
        ("Quotation No.", summary["quotation_no"]),
        ("Date", summary["quotation_date"]),
        ("Valid Until", summary["valid_until"]),
        ("Status", summary["status"]),
        ("Customer", summary["customer_name"]),
        ("Contact Person", summary.get("customer_contact_person", "")),
        ("Type", summary.get("customer_type", "Quotation")),
        ("Challan No.", summary.get("challan_no", "")),
        ("Challan Date", summary.get("challan_date", "")),
        ("L.R. No.", summary.get("lr_no", "")),
        ("Delivery", summary.get("delivery_mode", "")),
        ("Rev. Charge", summary.get("rev_charge", "No")),
        ("Ship To", summary.get("ship_to", "")),
        ("Distance for e-way bill (km)", summary.get("distance_for_eway_bill", "")),
        ("Place of Supply", summary.get("place_of_supply", "")),
        ("GSTIN / PAN", summary.get("customer_gstin", "")),
        ("Email", summary["customer_email"]),
        ("Phone", summary["customer_phone"]),
        ("Address", summary.get("customer_address", "")),
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quotation"
    sheet.append([summary["quotation_no"]])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    for label, value in header:
        sheet.append([label, value])
    sheet.append([])
    sheet.append(["Item", "HSN Code", "Description", "Capacity / Unit", "Qty", "Rate",
                  "Disc %", "Disc Amt", "Amount"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for row in summary["rows"]:
        sheet.append([
            row.get("item", ""), row.get("hsn_code", ""), row.get("description", ""),
            row.get("capacity_unit", ""), row.get("quantity", ""), row.get("unit_price", ""),
            row.get("discount_percent", ""), row.get("discount_amount", ""), row.get("amount", ""),
        ])
    sheet.append([])
    sheet.append(["", "", "", "", "", "", "", "Subtotal", f"{summary['subtotal']:.2f}"])
    if summary.get("discount_total"):
        sheet.append(["", "", "", "", "", "", "", "Discount", f"-{summary['discount_total']:.2f}"])
        sheet.append(["", "", "", "", "", "", "", "Net Amount", f"{summary['net_total']:.2f}"])
    sheet.append(["", "", "", "", "", "", "", f"Tax {summary['tax_percent']:g}%",
                  f"{summary['tax_amount']:.2f}"])
    sheet.append(["", "", "", "", "", "", "", "Grand Total", f"{summary['grand_total']:.2f}"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    if bank_lines:
        sheet.append([])
        sheet.append(["Bank Details"])
        sheet[sheet.max_row][0].font = Font(bold=True)
        for label, value in bank_lines:
            sheet.append([label, value])
    if terms:
        sheet.append([])
        sheet.append(["Terms & Conditions / Additional Note"])
        sheet[sheet.max_row][0].font = Font(bold=True)
        for index, term in enumerate(terms, 1):
            sheet.append([term.get("title", ""), term.get("detail", "")])
    if summary["notes"]:
        sheet.append([])
        sheet.append(["Notes", summary["notes"]])
    for column in sheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=8) + 2
        sheet.column_dimensions[column[0].column_letter].width = min(max(width, 10), 45)
    workbook.save(xlsx_path)

    styles = getSampleStyleSheet()
    pdf = SimpleDocTemplate(pdf_path, pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm)
    story = [Paragraph("QUOTATION", styles["Title"]), Spacer(1, 4 * mm)]
    for label, value in header:
        story.append(
            Paragraph(f"<b>{escape(str(label))}:</b> {escape(str(value))}", styles["Normal"])
        )
    story.append(Spacer(1, 4 * mm))
    data = [["No.", "Item", "HSN", "Description", "Capacity / Unit", "Qty", "Rate", "Disc %", "Amount"]]
    for index, row in enumerate(summary["rows"], 1):
        data.append([
            str(index), str(row.get("item", "")), str(row.get("hsn_code", "")),
            str(row.get("description", "")),
            str(row.get("capacity_unit", "")), str(row.get("quantity", "")),
            f"{try_float(row.get('unit_price'), 0.0):.2f}",
            str(row.get("discount_percent", "") or "0"),
            f"{try_float(row.get('amount'), 0.0):.2f}",
        ])
    data.append(["", "", "", "", "", "", "", "Subtotal", f"{summary['subtotal']:.2f}"])
    if summary.get("discount_total"):
        data.append(["", "", "", "", "", "", "", "Discount", f"-{summary['discount_total']:.2f}"])
        data.append(["", "", "", "", "", "", "", "Net Amount", f"{summary['net_total']:.2f}"])
    data.append(["", "", "", "", "", "", "", f"Tax {summary['tax_percent']:g}%", f"{summary['tax_amount']:.2f}"])
    data.append(["", "", "", "", "", "", "", "Grand Total", f"{summary['grand_total']:.2f}"])
    table = Table(data, colWidths=(8 * mm, 30 * mm, 16 * mm, 38 * mm, 22 * mm, 12 * mm, 20 * mm, 14 * mm, 24 * mm))
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (7, -4 if summary.get("discount_total") else -3), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ALIGN", (5, 1), (-1, -1), "RIGHT"),
    ]))
    story.extend([table, Spacer(1, 4 * mm)])
    if bank_lines:
        story.append(Paragraph("<b>Bank Details:</b>", styles["Heading3"]))
        for label, value in bank_lines:
            story.append(
                Paragraph(f"<b>{escape(str(label))}:</b> {escape(str(value))}", styles["Normal"])
            )
        story.append(Spacer(1, 2 * mm))
    if terms:
        story.append(Paragraph("<b>Terms &amp; Conditions / Additional Note:</b>", styles["Heading3"]))
        for index, term in enumerate(terms, 1):
            title = escape(str(term.get("title", "")))
            detail = escape(str(term.get("detail", "")))
            story.append(Paragraph(f"<b>{index}. {title}:</b> {detail}", styles["Normal"]))
        story.append(Spacer(1, 2 * mm))
    if summary["notes"]:
        story.append(Paragraph(f"<b>Notes:</b> {escape(str(summary['notes']))}", styles["Normal"]))
    pdf.build(story)

    return {
        "quotation_no": summary["quotation_no"],
        "customer_file": summary["file"],
        "xlsx": xlsx_path,
        "pdf": pdf_path,
    }


def open_path(path):
    """Open a file or folder with the operating system's default program."""
    import subprocess
    import sys

    try:
        if sys.platform == "win32":
            os.startfile(path)  # noqa: S606 - intended shell action
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
        return True
    except OSError:
        return False
