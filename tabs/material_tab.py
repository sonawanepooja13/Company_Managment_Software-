import csv
import os
import re

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl

import bom_engine
import category_order
import config
import views
from csv_product_manager import CSVProductManagerWindow


class MaterialTab(ttk.Frame):
    def __init__(self, parent, price_tab_data=None):
        super().__init__(parent)
        self.current_bom = []
        self.current_labor_cost = 0.0
        self.current_grand_total = 0.0
        self.price_tab_data = price_tab_data or {}
        self.product_price_csv = os.path.join(
            config.CSV_DIR,
            "Product Price Calculator",
            "price_list.csv",
        )
        self.vfd_make_values = self.load_vfd_makes()
        self.switch_gear_make_values = self.load_switch_gear_makes()
        self.mccb_make_values = self.load_mccb_makes()
        self.current_breaker_details = {}

        self.build_ui()
        
        # If data received from Price List, populate the fields after UI is ready
        if self.price_tab_data:
            self.after(50, self.populate_from_price_list)
    
    def update_price_tab_data(self, data):
        """Update price tab data and refresh fields."""
        self.price_tab_data = data
        self.populate_from_price_list()

    def build_ui(self):
        title_frame = ttk.Frame(self)
        title_frame.pack(fill="x", padx=15, pady=10)

        title = ttk.Label(title_frame, text="Material & Labor Calculator", font=("Helvetica", 14, "bold"))
        title.pack(side="left")

        ttk.Button(
            title_frame,
            text="Material Details Editor",
            command=self.open_material_details_editor,
        ).pack(side="right")

        # Create scrollable frame
        canvas = tk.Canvas(self)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
        scrollable_frame = ttk.Frame(canvas)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Mouse wheel scrolling
        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1*(event.delta/120)), "units")
        
        canvas.bind_all("<MouseWheel>", _on_mousewheel)

        form_frame = ttk.LabelFrame(scrollable_frame, text=" Inputs (Yellow Section) ", padding="10")
        form_frame.pack(fill="x", padx=15, pady=5)

        row = 0

        ttk.Label(form_frame, text="Pump Type:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_pump_type_combo = ttk.Combobox(form_frame, values=["3 Phase", "1 Phase"], width=18, state="readonly")
        self.m_pump_type_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_pump_type_combo.set("3 Phase")
        row += 1

        ttk.Label(form_frame, text="Pump Current (A):").grid(row=row, column=0, sticky="w", pady=4)
        self.m_current_entry = ttk.Entry(form_frame, width=20)
        self.m_current_entry.grid(row=row, column=1, sticky="e", pady=4)
        self.m_current_entry.insert(0, "15")
        ttk.Label(form_frame, text="Switch Gear Make:").grid(row=row, column=2, sticky="w", padx=(30, 0), pady=4)
        self.m_switch_gear_make_combo = ttk.Combobox(
            form_frame, values=self.switch_gear_make_values, width=18, state="readonly"
        )
        self.m_switch_gear_make_combo.grid(row=row, column=3, sticky="e", pady=4)
        if self.switch_gear_make_values:
            self.m_switch_gear_make_combo.set(self.switch_gear_make_values[0])
        row += 1

        ttk.Label(form_frame, text="MCCB Make:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_mccb_make_combo = ttk.Combobox(
            form_frame, values=self.mccb_make_values, width=18, state="readonly"
        )
        self.m_mccb_make_combo.grid(row=row, column=1, sticky="e", pady=4)
        if self.mccb_make_values:
            self.m_mccb_make_combo.set(self.mccb_make_values[0])
        ttk.Label(form_frame, text="Breaking Capacity (kA):").grid(row=row, column=2, sticky="w", padx=(30, 0), pady=4)
        self.m_breaking_capacity_combo = ttk.Combobox(
            form_frame, values=["25 kA", "36 kA", "50 kA"], width=18, state="readonly"
        )
        self.m_breaking_capacity_combo.grid(row=row, column=3, sticky="e", pady=4)
        self.m_breaking_capacity_combo.set("36 kA")
        row += 1

        ttk.Label(form_frame, text="Number of Pumps:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_pumps_combo = ttk.Combobox(form_frame, values=["1", "2", "3", "4", "5"], width=18, state="readonly")
        self.m_pumps_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_pumps_combo.set("3")
        row += 1

        ttk.Label(form_frame, text="Number of VFD:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_vfd_combo = ttk.Combobox(form_frame, values=["0", "1", "2", "3", "4"], width=18, state="readonly")
        self.m_vfd_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_vfd_combo.set("1")
        ttk.Label(form_frame, text="VFD Make:").grid(row=row, column=2, sticky="w", padx=(30, 0), pady=4)
        self.m_vfd_make_combo = ttk.Combobox(
            form_frame,
            values=self.vfd_make_values,
            width=18,
            state="readonly",
        )
        self.m_vfd_make_combo.grid(row=row, column=3, sticky="e", pady=4)
        if self.vfd_make_values:
            self.m_vfd_make_combo.set(self.vfd_make_values[0])
        row += 1

        ttk.Label(form_frame, text="Main Incomer Required:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_incomer_combo = ttk.Combobox(form_frame, values=["1 - Yes", "0 - No"], width=18, state="readonly")
        self.m_incomer_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_incomer_combo.set("1 - Yes")
        row += 1

        ttk.Label(form_frame, text="3-Pole Door Mount Switch:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_door_mount_combo = ttk.Combobox(form_frame, values=["1 - Yes", "0 - No"], width=18, state="readonly")
        self.m_door_mount_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_door_mount_combo.set("1 - Yes")
        row += 1

        ttk.Label(form_frame, text="Panel Size Level (1-5):").grid(row=row, column=0, sticky="w", pady=4)
        self.m_panel_size_combo = ttk.Combobox(form_frame, values=["1", "2", "3", "4", "5"], width=18, state="readonly")
        self.m_panel_size_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_panel_size_combo.set("4")
        row += 1

        ttk.Label(form_frame, text="Bypass:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_bypass_combo = ttk.Combobox(form_frame, values=["With Bypass", "Without Bypass"], width=18, state="readonly")
        self.m_bypass_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_bypass_combo.set("With Bypass")
        row += 1

        ttk.Label(form_frame, text="Panel Type:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_panel_type_combo = ttk.Combobox(form_frame, values=["Indoor", "Outdoor"], width=18, state="readonly")
        self.m_panel_type_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_panel_type_combo.set("Indoor")
        row += 1

        ttk.Label(form_frame, text="Panel Class:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_panel_class_combo = ttk.Combobox(form_frame, values=["Industrial", "Domestic"], width=18, state="readonly")
        self.m_panel_class_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_panel_class_combo.set("Industrial")
        self.m_panel_class_combo.bind("<<ComboboxSelected>>", self.update_breaking_capacity_default)
        self.update_breaking_capacity_default()
        row += 1

        ttk.Label(form_frame, text="OLR Requirement:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_olr_combo = ttk.Combobox(form_frame, values=["0 - No", "1 - Yes"], width=18, state="readonly")
        self.m_olr_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_olr_combo.set("0 - No")
        row += 1

        ttk.Label(form_frame, text="Indicator Light Required:").grid(row=row, column=0, sticky="w", pady=4)
        self.m_light_combo = ttk.Combobox(form_frame, values=["0 - No", "1 - Yes"], width=18, state="readonly")
        self.m_light_combo.grid(row=row, column=1, sticky="e", pady=4)
        self.m_light_combo.set("0 - No")
        row += 1

        action_frame = ttk.Frame(scrollable_frame)
        action_frame.pack(pady=5)

        calc_btn = ttk.Button(action_frame, text="Calculate Material List & BOM", command=self.calculate_materials)
        calc_btn.pack(side="left", padx=5)

        auto_calc_btn = ttk.Button(action_frame, text="Auto Calculator", command=self.auto_calculate_bom)
        auto_calc_btn.pack(side="left", padx=5)

        auto_calc2_btn = ttk.Button(action_frame, text="Auto Calculator2", command=self.open_auto_calculator2)
        auto_calc2_btn.pack(side="left", padx=5)

        export_btn = ttk.Button(action_frame, text="💾 Export BOM to CSV", command=self.export_bom_to_csv)
        export_btn.pack(side="left", padx=5)

        output_frame = ttk.LabelFrame(scrollable_frame, text=" Material Breakdown & BOM ", padding="10")
        output_frame.pack(fill="both", expand=True, padx=15, pady=5)

        self.output_text = tk.Text(output_frame, wrap="none", width=85, height=22, font=("Consolas", 9))

        vsb = ttk.Scrollbar(output_frame, orient="vertical", command=self.output_text.yview)
        hsb = ttk.Scrollbar(output_frame, orient="horizontal", command=self.output_text.xview)
        self.output_text.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.output_text.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        output_frame.grid_rowconfigure(0, weight=1)
        output_frame.grid_columnconfigure(0, weight=1)

    def update_breaking_capacity_default(self, _event=None):
        default_capacity = "36 kA" if self.m_panel_class_combo.get().strip().lower() == "industrial" else "25 kA"
        self.m_breaking_capacity_combo.set(default_capacity)

    def populate_from_price_list(self):
        """Populate fields from Price List Search data."""
        if not self.price_tab_data:
            return
        
        # Clear existing data first
        self.output_text.delete("1.0", tk.END)
        
        # Copy Pump Type
        if 'pump_type' in self.price_tab_data:
            self.m_pump_type_combo.set(self.price_tab_data['pump_type'])
        
        # Copy Pump Current
        if 'pump_current' in self.price_tab_data:
            self.m_current_entry.delete(0, tk.END)
            self.m_current_entry.insert(0, str(self.price_tab_data['pump_current']))
        
        # Copy Number of Pumps
        if 'num_pumps' in self.price_tab_data:
            self.m_pumps_combo.set(self.price_tab_data['num_pumps'])
        
        # Copy Number of VFD
        if 'num_vfd' in self.price_tab_data:
            self.m_vfd_combo.set(self.price_tab_data['num_vfd'])

        if 'vfd_make' in self.price_tab_data and self.price_tab_data['vfd_make'] in self.vfd_make_values:
            self.m_vfd_make_combo.set(self.price_tab_data['vfd_make'])

        if 'switch_gear_make' in self.price_tab_data and self.price_tab_data['switch_gear_make'] in self.switch_gear_make_values:
            self.m_switch_gear_make_combo.set(self.price_tab_data['switch_gear_make'])

        if 'mccb_make' in self.price_tab_data and self.price_tab_data['mccb_make'] in self.mccb_make_values:
            self.m_mccb_make_combo.set(self.price_tab_data['mccb_make'])

        if 'panel_class' in self.price_tab_data:
            self.m_panel_class_combo.set(self.price_tab_data['panel_class'])
        self.update_breaking_capacity_default()
        if 'breaking_capacity' in self.price_tab_data:
            capacity = str(self.price_tab_data['breaking_capacity']).strip()
            if capacity in {"25 kA", "36 kA", "50 kA"}:
                self.m_breaking_capacity_combo.set(capacity)
        
        # Copy Main Incomer (convert from Yes/No to 1 - Yes/0 - No)
        if 'main_incomer' in self.price_tab_data:
            incomer_value = self.price_tab_data['main_incomer']
            if incomer_value in ['Yes', 'yes', 'With Bypass', 'with bypass']:
                self.m_incomer_combo.set("1 - Yes")
            else:
                self.m_incomer_combo.set("0 - No")
        
        # Copy OLR Requirement (convert from Yes/No to 1 - Yes/0 - No)
        if 'olr_required' in self.price_tab_data:
            olr_value = self.price_tab_data['olr_required']
            if olr_value in ['Yes', 'yes']:
                self.m_olr_combo.set("1 - Yes")
            else:
                self.m_olr_combo.set("0 - No")
        
        # Copy Indicator Light (convert from Yes/No to 1 - Yes/0 - No)
        if 'indicator_light' in self.price_tab_data:
            light_value = self.price_tab_data['indicator_light']
            if light_value in ['Yes', 'yes']:
                self.m_light_combo.set("1 - Yes")
            else:
                self.m_light_combo.set("0 - No")
        
        # Copy Panel Size (convert from dimensions to level)
        if 'panel_size' in self.price_tab_data:
            size_mapping = {
                '400x300': '1',
                '600x400': '2', 
                '800x600': '3',
                '1000x800': '4',
                '1200x800': '5'
            }
            size_value = self.price_tab_data['panel_size']
            if size_value in size_mapping:
                self.m_panel_size_combo.set(size_mapping[size_value])
            else:
                self.m_panel_size_combo.set('4')  # Default
        
        # Copy Bypass
        if 'bypass' in self.price_tab_data:
            self.m_bypass_combo.set(self.price_tab_data['bypass'])
        
        # Copy Panel Type
        if 'panel_type' in self.price_tab_data:
            self.m_panel_type_combo.set(self.price_tab_data['panel_type'])
        
        # Copy Panel Class
        if 'panel_class' in self.price_tab_data:
            self.m_panel_class_combo.set(self.price_tab_data['panel_class'])
        self.update_breaking_capacity_default()

        # Show message that data has been copied
        self.output_text.insert(tk.END, "✓ Data copied from Price List Search\n")
        self.output_text.insert(tk.END, "✓ Click 'Calculate Material List & BOM' to get the cost breakdown\n")
        self.output_text.insert(tk.END, "✓ Then return to Price List Search to add the calculated price\n")

    def open_product_price_list(self):
        views.CSVViewerWindow(
            self.winfo_toplevel(),
            csv_path=self.product_price_csv,
            title="Product Price List - Material & Labor Calculator",
            edit_callback=self.edit_product_price_list,
        )

    def open_material_details_editor(self):
        editor = tk.Toplevel(self.winfo_toplevel())
        self.material_details_editor = editor
        editor.title("Material Details Editor")
        editor.geometry("520x260")
        editor.minsize(460, 220)
        editor.transient(self.winfo_toplevel())
        editor.grab_set()

        content = ttk.Frame(editor, padding=20)
        content.pack(fill="both", expand=True)

        ttk.Label(
            content,
            text="Material Details Editor",
            font=("Helvetica", 14, "bold"),
        ).pack(pady=(0, 5))
        ttk.Label(
            content,
            text="Manage material and labour pricing data",
        ).pack(pady=(0, 18))

        buttons = ttk.Frame(content)
        buttons.pack(fill="x")
        buttons.columnconfigure(0, weight=1)
        buttons.columnconfigure(1, weight=1)

        actions = [
            ("📊 Open Product Price List", self.open_product_price_list),
            ("✏ Edit CSV List", self.edit_product_price_list),
            ("⬆ Attach Excel", self.import_excel_price_list),
            ("⬇ Download Excel", self.export_excel_price_list),
        ]
        for index, (label, command) in enumerate(actions):
            ttk.Button(buttons, text=label, command=command).grid(
                row=index // 2,
                column=index % 2,
                padx=6,
                pady=6,
                sticky="ew",
            )

    def edit_product_price_list(self):
        parent = getattr(self, "material_details_editor", None)
        if parent is None or not parent.winfo_exists():
            parent = self.winfo_toplevel()
        CSVProductManagerWindow(
            parent,
            default_file=self.product_price_csv,
        )

    def import_excel_price_list(self):
        source_path = filedialog.askopenfilename(
            parent=self.winfo_toplevel(),
            title="Attach Excel Price List",
            filetypes=(
                ("Excel files", "*.xlsx *.xlsm"),
                ("All files", "*.*"),
            ),
        )
        if not source_path:
            return

        try:
            workbook = openpyxl.load_workbook(source_path, data_only=True)
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows or not any(value not in (None, "") for value in rows[0]):
                raise ValueError("The first worksheet does not contain a header row.")

            # Arrange the list category-wise before saving
            rows = self.sort_rows_by_category(rows)

            os.makedirs(os.path.dirname(self.product_price_csv), exist_ok=True)
            with open(self.product_price_csv, "w", newline="", encoding="utf-8-sig") as output:
                csv.writer(output).writerows(rows)
            messagebox.showinfo(
                "Excel Attached",
                f"The Excel price list was imported into:\n{self.product_price_csv}",
                parent=self.winfo_toplevel(),
            )
        except Exception as error:
            messagebox.showerror(
                "Excel Import Failed",
                f"Could not attach this Excel price list:\n{error}",
                parent=self.winfo_toplevel(),
            )

    def export_excel_price_list(self):
        destination = filedialog.asksaveasfilename(
            parent=self.winfo_toplevel(),
            title="Download Product Price List as Excel",
            defaultextension=".xlsx",
            initialfile="price_list.xlsx",
            filetypes=(("Excel workbook", "*.xlsx"),),
        )
        if not destination:
            return

        try:
            with open(self.product_price_csv, newline="", encoding="utf-8-sig") as source:
                rows = list(csv.reader(source))
            # Download the list sorted category-wise
            rows = self.sort_rows_by_category(rows)
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.title = "Price List"
            for row in rows:
                worksheet.append(row)
            workbook.save(destination)
            messagebox.showinfo(
                "Excel Downloaded",
                f"The Excel price list was saved to:\n{destination}",
                parent=self.winfo_toplevel(),
            )
        except Exception as error:
            messagebox.showerror(
                "Excel Export Failed",
                f"Could not download the Excel price list:\n{error}",
                parent=self.winfo_toplevel(),
            )

    @staticmethod
    def sort_rows_by_category(rows):
        """Arrange rows by the Category column, keeping the header row first.

        Uses the shared preferred category order so every upload is arranged
        the same way (MCB, MCCB, RCCB, MPCB, Contactor, ...).
        """
        return category_order.sort_rows_by_category(rows)

    def calculate_materials(self):
        selection_window = tk.Toplevel(self.winfo_toplevel())
        selection_window.title("Material List & BOM - Category Selection")
        selection_window.state("zoomed")
        selection_window.transient(self.winfo_toplevel())
        selection_window.grab_set()

        content = ttk.Frame(selection_window, padding=24)
        content.pack(fill="both", expand=True)
        ttk.Label(
            content,
            text="Material List & BOM",
            font=("Helvetica", 18, "bold"),
        ).pack(pady=(0, 4))
        ttk.Label(
            content,
            text="Select a category and item. Add Category to include more material lines.",
        ).pack(pady=(0, 16))

        material_id = self.get_panel_material_id()
        id_frame = ttk.Frame(content)
        id_frame.pack(fill="x", pady=(0, 12))
        ttk.Label(id_frame, text=f"Unique ID: {material_id}", font=("Helvetica", 10, "bold")).pack(side="left")
        saved_ids = self.get_saved_material_ids()
        saved_id_combo = ttk.Combobox(id_frame, values=saved_ids, state="readonly", width=26)
        saved_id_combo.pack(side="right", padx=(8, 0))

        rows = []

        category_values = self.load_price_categories()
        list_frame = ttk.Frame(content)
        list_frame.pack(fill="both", expand=True)
        list_frame.columnconfigure(1, weight=1)
        list_frame.columnconfigure(2, weight=1)

        def add_category_row():
            row_index = len(rows)
            category_combo = ttk.Combobox(
                list_frame, values=category_values, state="readonly", width=28
            )
            item_combo = ttk.Combobox(list_frame, state="readonly", width=80)
            quantity_entry = ttk.Entry(list_frame, width=10)
            quantity_entry.insert(0, "1")
            remove_button = ttk.Button(
                list_frame,
                text="Remove",
                command=lambda: remove_category_row(row_data),
            )

            ttk.Label(list_frame, text=f"Material {row_index + 1}").grid(
                row=row_index + 1, column=0, sticky="w", padx=(0, 12), pady=8
            )
            category_combo.grid(row=row_index + 1, column=1, sticky="ew", padx=4, pady=8)
            item_combo.grid(row=row_index + 1, column=2, sticky="ew", padx=4, pady=8)
            quantity_entry.grid(row=row_index + 1, column=3, padx=4, pady=8)
            remove_button.grid(row=row_index + 1, column=4, padx=(8, 0), pady=8)

            row_data = {
                "label": list_frame.grid_slaves(row=row_index + 1, column=0)[0],
                "category": category_combo,
                "item": item_combo,
                "quantity": quantity_entry,
                "remove": remove_button,
            }
            rows.append(row_data)

            def update_items(_event=None, current=row_data):
                options = self.load_component_options(current["category"].get())
                current["item"]["values"] = [option["label"] for option in options]
                current["options"] = options
                if options:
                    current["item"].current(0)
                else:
                    current["item"].set("")

            category_combo.bind("<<ComboboxSelected>>", update_items)
            row_data["options"] = []
            if category_values:
                category_combo.current(0)
                update_items()

        def load_material_rows(material_rows):
            for current in list(rows):
                for widget in current.values():
                    if isinstance(widget, tk.Widget):
                        widget.destroy()
            rows.clear()
            for record in material_rows:
                add_category_row()
                current = rows[-1]
                current["category"].set(record.get("category", ""))
                options = self.load_component_options(current["category"].get())
                current["options"] = options
                current["item"]["values"] = [option["label"] for option in options]
                item_label = record.get("item", "")
                if item_label in current["item"]["values"]:
                    current["item"].set(item_label)
                current["quantity"].delete(0, tk.END)
                current["quantity"].insert(0, record.get("quantity", "1"))
            if not rows:
                add_category_row()

        def refresh_saved_ids():
            saved_id_combo["values"] = self.get_saved_material_ids()

        def open_saved_list():
            selected_id = saved_id_combo.get().strip()
            if not selected_id:
                messagebox.showwarning("Select Saved List", "Select a unique ID first.", parent=selection_window)
                return
            loaded_rows = self.load_saved_material_rows(selected_id)
            if loaded_rows is None:
                messagebox.showerror("Load Failed", "The selected material list could not be opened.", parent=selection_window)
                return
            load_material_rows(loaded_rows)

        def save_current_list():
            material_rows = []
            for row_data in rows:
                category = row_data["category"].get().strip()
                item = row_data["item"].get().strip()
                if not category or not item:
                    messagebox.showwarning("Required Selection", "Select a category and item for every material row.", parent=selection_window)
                    return
                material_rows.append({
                    "category": category,
                    "item": item,
                    "quantity": row_data["quantity"].get().strip() or "1",
                })
            if self.save_material_rows(material_id, material_rows):
                refresh_saved_ids()
                messagebox.showinfo("Saved", f"Material list saved as {material_id}.", parent=selection_window)

        def remove_category_row(row_data):
            if len(rows) == 1:
                return
            for widget in row_data.values():
                if isinstance(widget, tk.Widget):
                    widget.destroy()
            rows.remove(row_data)
            for index, current in enumerate(rows):
                current["label"].configure(text=f"Material {index + 1}")
                for widget in current.values():
                    if isinstance(widget, tk.Widget):
                        widget.grid_configure(row=index + 1)

        ttk.Label(list_frame, text="Material").grid(row=0, column=0, sticky="w")
        ttk.Label(list_frame, text="Category").grid(row=0, column=1, sticky="w", padx=4)
        ttk.Label(list_frame, text="Item / Price").grid(row=0, column=2, sticky="w", padx=4)
        ttk.Label(list_frame, text="Qty").grid(row=0, column=3, sticky="w", padx=4)
        existing_rows = self.load_saved_material_rows(material_id)
        if existing_rows:
            load_material_rows(existing_rows)
        else:
            add_category_row()

        actions = ttk.Frame(content)
        actions.pack(fill="x", pady=(18, 0))
        ttk.Button(actions, text="+ Add Category", command=add_category_row).pack(side="left")
        ttk.Button(actions, text="Open Saved List", command=open_saved_list).pack(side="left", padx=8)
        ttk.Button(actions, text="Copy Selected ID", command=open_saved_list).pack(side="left")
        ttk.Button(actions, text="Save Material List", command=save_current_list).pack(side="left")

        def calculate_selected_materials():
            selected_materials = []
            for row_data in rows:
                category = row_data["category"].get().strip()
                item = row_data["item"].get().strip()
                if not category or not item:
                    messagebox.showwarning(
                        "Required Selection",
                        "Select a category and item for every material row.",
                        parent=selection_window,
                    )
                    return
                selected_materials.append(
                    {
                        "category": category,
                        "item": item,
                        "quantity": row_data["quantity"].get().strip() or "1",
                    }
                )
            selection_window.destroy()
            self.calculate_materials_now(selected_materials)

        ttk.Button(
            actions,
            text="⚡ Auto Calculate",
            command=lambda: (selection_window.destroy(), self.calculate_materials_now()),
        ).pack(side="right", padx=(8, 0))

        ttk.Button(
            actions,
            text="Calculate Material List & BOM",
            command=calculate_selected_materials,
        ).pack(side="right")

    def get_panel_input_values(self):
        return {
            "pump_type": self.m_pump_type_combo.get().strip(),
            "pump_current": self.m_current_entry.get().strip(),
            "switch_gear_make": self.m_switch_gear_make_combo.get().strip(),
            "mccb_make": self.m_mccb_make_combo.get().strip(),
            "breaking_capacity": self.m_breaking_capacity_combo.get().strip(),
            "num_pumps": self.m_pumps_combo.get().strip(),
            "num_vfd": self.m_vfd_combo.get().strip(),
            "vfd_make": self.m_vfd_make_combo.get().strip(),
            "main_incomer": self.m_incomer_combo.get().strip(),
            "door_mount": self.m_door_mount_combo.get().strip(),
            "panel_size": self.m_panel_size_combo.get().strip(),
            "bypass": self.m_bypass_combo.get().strip(),
            "panel_type": self.m_panel_type_combo.get().strip(),
            "panel_class": self.m_panel_class_combo.get().strip(),
            "olr": self.m_olr_combo.get().strip(),
            "indicator_light": self.m_light_combo.get().strip(),
        }

    def get_panel_material_id(self):
        values = self.get_panel_input_values()
        def numeric_value(value):
            match = re.search(r"-?\d+(?:\.\d+)?", str(value))
            if not match:
                return "0"
            number = float(match.group())
            return str(int(number)) if number.is_integer() else str(number).replace(".", "p")

        def option_code(value, options):
            normalized = str(value).strip().lower()
            for index, option in enumerate(options, start=1):
                if normalized == str(option).strip().lower():
                    return str(index)
            return "0"

        id_values = [
            f"PT{option_code(values['pump_type'], ['3 Phase', '1 Phase'])}",
            f"PC{numeric_value(values['pump_current'])}A",
            f"NOP{numeric_value(values['num_pumps'])}",
            f"NOVFD{numeric_value(values['num_vfd'])}",
            f"VFM{option_code(values['vfd_make'], self.vfd_make_values)}",
            f"MIR{numeric_value(values['main_incomer'])}",
            f"3PDMS{numeric_value(values['door_mount'])}",
            f"PS{option_code(values['panel_size'], ['1', '2', '3', '4', '5'])}",
            f"BY{option_code(values['bypass'], ['With Bypass', 'Without Bypass'])}",
            f"PNLT{option_code(values['panel_type'], ['Indoor', 'Outdoor'])}",
            f"PCL{option_code(values['panel_class'], ['Industrial', 'Domestic'])}",
            f"OLR{numeric_value(values['olr'])}",
            f"ILR{numeric_value(values['indicator_light'])}",
            f"SGM{option_code(values['switch_gear_make'], self.switch_gear_make_values)}",
            f"MCM{option_code(values['mccb_make'], self.mccb_make_values)}",
            f"BC{option_code(values['breaking_capacity'], ['25 kA', '36 kA', '50 kA'])}",
        ]
        return "AD | " + " | ".join(id_values)

    def get_saved_material_ids(self):
        os.makedirs(config.PANEL_MATERIAL_LIST_DIR, exist_ok=True)
        return sorted(
            os.path.splitext(name)[0]
            for name in os.listdir(config.PANEL_MATERIAL_LIST_DIR)
            if name.lower().endswith(".csv")
        )

    def panel_material_path(self, material_id):
        safe_id = re.sub(r"[^A-Za-z0-9_-]", "", material_id)
        return os.path.join(config.PANEL_MATERIAL_LIST_DIR, f"{safe_id}.csv")

    def save_material_rows(self, material_id, material_rows):
        try:
            with open(self.panel_material_path(material_id), "w", newline="", encoding="utf-8-sig") as file:
                writer = csv.DictWriter(file, fieldnames=["unique_id", "category", "item", "quantity"])
                writer.writeheader()
                for row in material_rows:
                    writer.writerow({"unique_id": material_id, **row})
            return True
        except OSError as error:
            messagebox.showerror("Save Failed", f"Could not save the material list:\n{error}", parent=self.winfo_toplevel())
            return False

    def load_saved_material_rows(self, material_id):
        try:
            with open(self.panel_material_path(material_id), newline="", encoding="utf-8-sig") as file:
                return list(csv.DictReader(file))
        except (OSError, csv.Error):
            return None

    def load_price_categories(self):
        categories = []
        try:
            with open(self.product_price_csv, newline="", encoding="utf-8-sig") as file:
                for row in csv.DictReader(file):
                    category = str(row.get("Category", "")).strip()
                    if category and category.lower() not in {value.lower() for value in categories}:
                        categories.append(category)
        except (OSError, csv.Error):
            pass
        return categories or ["General"]

    def auto_calculate_bom(self):
        self.calculate_materials_now()

    def calculate_materials_now(self, component_selections=None):
        try:
            pump_type = self.m_pump_type_combo.get()
            pump_current = float(self.m_current_entry.get().strip())
            switch_gear_make = self.m_switch_gear_make_combo.get().strip()
            mccb_make = self.m_mccb_make_combo.get().strip()
            breaking_capacity = self.m_breaking_capacity_combo.get().strip()
            num_pumps = int(self.m_pumps_combo.get().strip())
            num_vfd = int(self.m_vfd_combo.get().strip())
            vfd_make = self.m_vfd_make_combo.get().strip()
            incomer_req = 1 if "1" in self.m_incomer_combo.get() else 0
            door_mount_req = 1 if "1" in self.m_door_mount_combo.get() else 0
            panel_size_lvl = int(self.m_panel_size_combo.get().strip())
            olr_req = 1 if "1" in self.m_olr_combo.get() else 0
            light_req = 1 if "1" in self.m_light_combo.get() else 0
            bypass = self.m_bypass_combo.get()
            panel_type = self.m_panel_type_combo.get()
            panel_class = self.m_panel_class_combo.get()
            if pump_current <= 0 or num_pumps <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Input Error", "Pump current and number of pumps must be positive values.")
            return

        pump_hp = (pump_current * 1.732 * 415 * 0.85) / 746
        total_panel_current = pump_current * num_pumps
        controller_type = "AIPCU OR HMI" if num_vfd > 0 else "DOL/STAR-DELTA"
        mcb_mccb_type = "MCCB" if total_panel_current > 63 else "MCB"
        breaker_qty = 1
        raw_breaker_rating = total_panel_current

        fan_qty = 2 if num_vfd > 1 else (1 if num_vfd == 1 else 0)
        filter_qty = fan_qty
        endlock_qty = 4

        self.current_bom = bom_engine.generate_bom(
            controller_type, num_pumps, num_vfd, pump_hp, olr_req, light_req,
            fan_qty, filter_qty, endlock_qty, mcb_mccb_type, breaker_qty,
            raw_breaker_rating, incomer_req, door_mount_req, total_panel_current,
            vfd_make, "3P", panel_class, mccb_make, breaking_capacity,
            switch_gear_make,
        )
        self.current_breaker_details = {
            "total_current": total_panel_current,
            "breaker_type": mcb_mccb_type,
            "panel_class": panel_class,
            "mccb_make": mccb_make,
            "breaking_capacity": breaking_capacity,
        }

        if isinstance(component_selections, dict):
            breaker_category = "MCCB" if mcb_mccb_type == "MCCB" else "MCB"
            selected_breaker = component_selections.get(breaker_category)
            if selected_breaker:
                for item in self.current_bom:
                    if item["Category"] == breaker_category:
                        item.update(
                            Item_Name=selected_breaker["item_name"],
                            **{"SUB Category": selected_breaker["sub_category"]},
                            Capacity=f"{selected_breaker['capacity']}A",
                            DP=selected_breaker["dp"],
                        )
                        break
            selected_contactor = component_selections.get("CONTACTOR")
            if selected_contactor:
                for item in self.current_bom:
                    if item["Category"] == "CONTACTOR":
                        item.update(
                            Item_Name=selected_contactor["item_name"],
                            **{"SUB Category": selected_contactor["sub_category"]},
                            Capacity=f"{selected_contactor['capacity']}A",
                            DP=selected_contactor["dp"],
                        )
                        break
            selected_main_incomer = component_selections.get("MAIN INCOMER")
            if selected_main_incomer:
                for item in self.current_bom:
                    if item["Category"].lower() == "switch":
                        item.update(
                            Item_Name=selected_main_incomer["item_name"],
                            **{"SUB Category": selected_main_incomer["sub_category"]},
                            Capacity=f"{selected_main_incomer['capacity']}A",
                            DP=selected_main_incomer["dp"],
                        )
                        break
            selected_separate_breaker = component_selections.get(
                "SEPARATE CIRCUIT BREAKER (MCB/MCCB/RCCB/RCB/MPCB)"
            )
            if selected_separate_breaker:
                for item in self.current_bom:
                    if item["Category"] in {"MCB", "MCCB"}:
                        item.update(
                            Item_Name=selected_separate_breaker["item_name"],
                            Category=selected_separate_breaker["category"],
                            **{"SUB Category": selected_separate_breaker["sub_category"]},
                            Capacity=f"{selected_separate_breaker['capacity']}A",
                            DP=selected_separate_breaker["dp"],
                        )
                        break

        total_material_dp = sum(item["Qty"] * item["DP"] for item in self.current_bom)
        self.current_labor_cost = 500.0 + (num_pumps * 250.0) + (panel_size_lvl * 150.0)
        self.current_grand_total = total_material_dp + self.current_labor_cost

        self.output_text.delete("1.0", tk.END)
        out = []
        
        # Display input parameters
        out.append("=" * 80)
        out.append("INPUT PARAMETERS:")
        out.append(f"  Pump Type: {pump_type}")
        out.append(f"  Pump Current: {pump_current} A")
        out.append(f"  Switch Gear Make: {switch_gear_make or 'Not specified'}")
        out.append(f"  Number of Pumps: {num_pumps}")
        out.append(f"  Number of VFD: {num_vfd}")
        out.append(f"  VFD Make: {vfd_make or 'Not selected'}")
        out.append(f"  Bypass: {bypass}")
        out.append(f"  Panel Type: {panel_type}")
        out.append(f"  Panel Size Level: {panel_size_lvl}")
        out.append(f"  Panel Class: {panel_class}")
        out.append(f"  Main Incomer: {'Yes' if incomer_req else 'No'}")
        out.append(f"  OLR Requirement: {'Yes' if olr_req else 'No'}")
        out.append(f"  Indicator Light: {'Yes' if light_req else 'No'}")
        out.append("")
        out.append("AUTO BREAKER SELECTION:")
        out.append(f"  Total Current = {pump_current:g} A x {num_pumps} = {total_panel_current:g} A")
        selected_breaker = next(
            (item for item in self.current_bom if item["Category"] in {"MCB", "MCCB"}),
            None,
        )
        if total_panel_current <= 63:
            out.append("  Result: NO MCCB")
            out.append("  Selected: Next standard 3P MCB rating >= total current")
            if selected_breaker:
                out.append(f"  Selected MCB Rating: {selected_breaker['Capacity']}")
        else:
            is_industrial = panel_class.strip().lower() == "industrial"
            out.append("  Result: MCCB REQUIRED")
            out.append(f"  Panel Application: {panel_class}")
            out.append(f"  MCCB Poles: {'4P' if is_industrial else '3P'}")
            out.append(f"  Phase/Neutral: {'R+Y+B+N' if is_industrial else 'R+Y+B + Neutral'}")
            if is_industrial:
                out.append("  Busbars: 4 (R+Y+B+N)")
            else:
                out.append("  Neutral Protection: 6A SP MCB")
            out.append(f"  MCCB Make: {mccb_make or 'Not specified'}")
            out.append(f"  Breaking Capacity: {breaking_capacity}")
            if selected_breaker:
                out.append(f"  Selected MCCB Rating: {selected_breaker['Capacity']}")
        if isinstance(component_selections, list):
            out.append("  Selected Materials:")
            for material in component_selections:
                out.append(
                    f"    - {material['category']}: {material['item']} (Qty {material['quantity']})"
                )
        elif component_selections is None:
            out.append("  [✓ Auto-Calculated] BOM generated automatically from panel inputs")
        out.append("=" * 80)
        out.append("")
        
        out.append(f"{'ITEM NAME':<42} | {'QTY':<4} | {'UNIT DP (Rs.)':<13} | {'TOTAL DP (Rs.)':<13}")
        out.append("-" * 80)

        for item in self.current_bom:
            tot_dp = item["Qty"] * item["DP"]
            out.append(f"{item['Item_Name'][:40]:<42} | {item['Qty']:<4} | {item['DP']:<13,.2f} | {tot_dp:<13,.2f}")

        out.append("-" * 80)
        out.append(f"{'TOTAL MATERIAL COST (DP):':<62} Rs.{total_material_dp:,.2f}")
        out.append(f"{'LABOR & ASSEMBLY COST:':<62} Rs.{self.current_labor_cost:,.2f}")
        out.append(f"{'GRAND TOTAL ESTIMATED COST:':<62} Rs.{self.current_grand_total:,.2f}")

        self.output_text.insert(tk.END, "\n".join(out))

    def export_bom_to_csv(self):
        if not self.current_bom:
            messagebox.showwarning("Export Warning", "No BOM generated to export. Please calculate first.")
            return

        export_file = config.BOM_EXPORT_CSV
        try:
            with open(export_file, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(
                    f,
                    fieldnames=[
                        "Item_Name",
                        "Category",
                        "SUB Category",
                        "Capacity",
                        "Qty",
                        "DP",
                        "Make",
                        "Breaking Capacity",
                        "Application",
                        "Phase/Neutral",
                    ],
                )
                writer.writeheader()
                for item in self.current_bom:
                    writer.writerow(item)
            messagebox.showinfo("Export Success", f"BOM export saved successfully to:\n{export_file}")
        except Exception as e:
            messagebox.showerror("Export Error", f"Could not export BOM CSV: {e}")

    def open_auto_calculator2(self):
        """Open the Auto Calculator2 window for VFD BOM calculations."""
        from vfd_bom_calculator_window import VFDBOMCalculatorWindow
        calculator_window = VFDBOMCalculatorWindow(self)
        calculator_window.grab_set()

    def load_vfd_makes(self):
        makes = []
        try:
            with open(self.product_price_csv, newline="", encoding="utf-8-sig") as file:
                for row in csv.DictReader(file):
                    if str(row.get("Category", "")).strip().lower() != "vfd":
                        continue
                    make = str(row.get("SUB Category Type 1", "")).strip()
                    if make and make.lower() not in {value.lower() for value in makes}:
                        makes.append(make)
        except (OSError, csv.Error):
            pass
        return makes

    def load_component_options(self, categories):
        if isinstance(categories, str):
            categories = [categories]
        category_names = {category.lower() for category in categories}
        options = []
        try:
            with open(self.product_price_csv, newline="", encoding="utf-8-sig") as file:
                for row in csv.DictReader(file):
                    row_category = str(row.get("Category", "")).strip()
                    if row_category.lower() not in category_names:
                        continue
                    item_name = str(row.get("Item Name", "")).strip()
                    sub_category = str(row.get("SUB Category Type 1", "")).strip()
                    capacity = str(row.get("CAPACITY", "")).strip()
                    if not capacity:
                        match = re.search(r"(\d+(?:\.\d+)?)\s*A\b", item_name, re.IGNORECASE)
                        capacity = match.group(1) if match else ""
                    try:
                        dp = float(row.get("DP", 0) or 0)
                    except (TypeError, ValueError):
                        dp = 0.0
                    if not item_name:
                        continue
                    options.append(
                        {
                            "label": f"{sub_category or 'Standard'} | {item_name} | {capacity or 'N/A'}A | DP Rs.{dp:,.2f}",
                            "category": row_category,
                            "item_name": item_name,
                            "sub_category": sub_category,
                            "capacity": capacity or "0",
                            "dp": dp,
                        }
                    )
        except (OSError, csv.Error):
            pass
        return options

    def load_mccb_makes(self):
        make_values = []
        try:
            with open(self.product_price_csv, newline="", encoding="utf-8-sig") as file:
                for row in csv.DictReader(file):
                    if str(row.get("Category", "")).strip().lower() != "mccb":
                        continue
                    make = str(row.get("make", "")).strip()
                    if make and make.lower() not in {value.lower() for value in make_values}:
                        make_values.append(make)
        except (OSError, csv.Error):
            pass
        return make_values or list(self.switch_gear_make_values) or ["Not specified"]

    def load_switch_gear_makes(self):
        make_values = []
        try:
            with open(self.product_price_csv, newline="", encoding="utf-8-sig") as file:
                for row in csv.DictReader(file):
                    if str(row.get("Category", "")).strip().lower() != "mcb":
                        continue
                    make = str(row.get("make", "")).strip()
                    if make and make.lower() not in {value.lower() for value in make_values}:
                        make_values.append(make)
        except (OSError, csv.Error):
            pass
        return make_values or ["Not specified"]
