"""Full-screen Human Resources workspace with local CSV registers."""

import csv
import os
import re
import shutil
import tempfile
import tkinter as tk
from datetime import date, datetime
from tkinter import filedialog, messagebox, ttk

import config
from tkcalendar import DateEntry


DEPARTMENTS = (
    "Executive / Leadership",
    "Human Resources (HR)",
    "Finance & Accounting",
    "Sales & Business Development",
    "Marketing & Communications",
    "Operations & Logistics",
    "Information Technology (IT) / Engineering",
    "Customer Support / Success",
    "Legal & Compliance",
    "Research & Development (R&D)",
)
EMPLOYMENT_TYPES = ("Full-time", "Part-time", "Contractor", "Intern")
EMPLOYEE_ID_PREFIX = "SE-01"
BLOOD_GROUPS = ("A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-")
TIME_OPTIONS = tuple(
    f"{hour:02d}:{minute:02d}" for hour in range(24) for minute in range(0, 60, 15)
)
COUNTRIES = tuple("""
Afghanistan|Albania|Algeria|Andorra|Angola|Antigua and Barbuda|Argentina|Armenia|Australia|Austria|Azerbaijan|Bahamas|Bahrain|Bangladesh|Barbados|Belarus|Belgium|Belize|Benin|Bhutan|Bolivia|Bosnia and Herzegovina|Botswana|Brazil|Brunei|Bulgaria|Burkina Faso|Burundi|Cabo Verde|Cambodia|Cameroon|Canada|Central African Republic|Chad|Chile|China|Colombia|Comoros|Congo (Republic of the)|Congo (Democratic Republic of the)|Costa Rica|Cote d'Ivoire|Croatia|Cuba|Cyprus|Czechia|Denmark|Djibouti|Dominica|Dominican Republic|Ecuador|Egypt|El Salvador|Equatorial Guinea|Eritrea|Estonia|Eswatini|Ethiopia|Fiji|Finland|France|Gabon|Gambia|Georgia|Germany|Ghana|Greece|Grenada|Guatemala|Guinea|Guinea-Bissau|Guyana|Haiti|Honduras|Hungary|Iceland|India|Indonesia|Iran|Iraq|Ireland|Israel|Italy|Jamaica|Japan|Jordan|Kazakhstan|Kenya|Kiribati|Korea (North)|Korea (South)|Kuwait|Kyrgyzstan|Laos|Latvia|Lebanon|Lesotho|Liberia|Libya|Liechtenstein|Lithuania|Luxembourg|Madagascar|Malawi|Malaysia|Maldives|Mali|Malta|Marshall Islands|Mauritania|Mauritius|Mexico|Micronesia|Moldova|Monaco|Mongolia|Montenegro|Morocco|Mozambique|Myanmar|Namibia|Nauru|Nepal|Netherlands|New Zealand|Nicaragua|Niger|Nigeria|North Macedonia|Norway|Oman|Pakistan|Palau|Palestine|Panama|Papua New Guinea|Paraguay|Peru|Philippines|Poland|Portugal|Qatar|Romania|Russia|Rwanda|Saint Kitts and Nevis|Saint Lucia|Saint Vincent and the Grenadines|Samoa|San Marino|Sao Tome and Principe|Saudi Arabia|Senegal|Serbia|Seychelles|Sierra Leone|Singapore|Slovakia|Slovenia|Solomon Islands|Somalia|South Africa|South Sudan|Spain|Sri Lanka|Sudan|Suriname|Sweden|Switzerland|Syria|Tajikistan|Tanzania|Thailand|Timor-Leste|Togo|Tonga|Trinidad and Tobago|Tunisia|Turkey|Turkmenistan|Tuvalu|Uganda|Ukraine|United Arab Emirates|United Kingdom|United States|Uruguay|Uzbekistan|Vanuatu|Vatican City|Venezuela|Vietnam|Yemen|Zambia|Zimbabwe
""".strip().split("|"))

HR_TABS = (
    ("Employee Profile", "employees", (
        ("Employee ID", "employee_id"), ("First Name", "first_name"), ("Middle Name", "middle_name"), ("Full Name", "full_name"), ("Work Email", "work_email"), ("Date of Birth", "date_of_birth"),
        ("Gender", "gender"), ("Marital Status", "marital_status"), ("Nationality", "nationality"),
        ("Blood Group", "blood_group"), ("Emergency Contact", "emergency_contact"),
        ("Personal Email", "personal_email"), ("Phone Number", "phone_number"), ("Current Address", "current_address"),
        ("Permanent Address", "permanent_address"), ("Department", "department"), ("Designation / Title", "designation"),
        ("Reporting Manager ID", "reporting_manager"), ("Reporting Manager Name", "reporting_manager_name"), ("Date of Joining", "date_of_joining"), ("Employment Type", "employment_type"), ("Employment Status", "employment_status"), ("Created At", "created_at"),
        ("Work Location / Branch", "work_location"), ("Resume / CV", "resume_cv"), ("Signed Offer Letter", "offer_letter"),
        ("Employment Contract", "employment_contract"), ("NDA", "nda"), ("Government ID Proofs", "government_id"), ("Photograph", "photograph"),
    )),
    ("Employee Documents", "employee_documents", (
        ("Document ID", "document_id"), ("Employee ID", "employee_id"), ("Document Type", "document_type"),
        ("Document File", "file_path"), ("Uploaded At", "uploaded_at"),
    )),
    ("Recruitment (ATS)", "recruitment", (
        ("Job ID", "job_id"), ("Job Title", "job_title"), ("Department", "department"), ("Open Positions", "open_positions"),
        ("Minimum Salary", "min_salary"), ("Maximum Salary", "max_salary"), ("Job Status", "job_status"), ("Job Description", "job_description"), ("Required Skills", "required_skills"),
        ("Hiring Manager", "hiring_manager"), ("Candidate ID", "candidate_id"), ("Candidate First Name", "candidate_first_name"), ("Candidate Last Name", "candidate_last_name"), ("Candidate Name", "candidate_name"), ("Candidate Email", "candidate_email"),
        ("Candidate Phone", "candidate_phone"), ("Resume File", "resume_file"), ("Cover Letter", "cover_letter"),
        ("Source", "source"), ("Current Stage", "current_stage"), ("Interview Rounds", "interview_rounds"), ("Interviewer Feedback / Score", "feedback_score"),
        ("Interview Date", "interview_date"), ("Offer Status", "offer_status"),
    )),
    ("Interviews", "interviews", (
        ("Interview ID", "interview_id"), ("Candidate ID", "candidate_id"), ("Interviewer Employee ID", "interviewer_id"),
        ("Scheduled Date / Time", "scheduled_time"), ("Score", "score"), ("Feedback", "feedback"), ("Interview Status", "interview_status"),
    )),
    ("Shift Setup", "shifts", (
        ("Shift ID", "shift_id"), ("Shift Name", "shift_name"), ("Start Time", "start_time"),
        ("End Time", "end_time"), ("Grace Period (Minutes)", "grace_period_minutes"),
    )),
    ("Attendance & Leave", "attendance_leave", (
        ("Employee ID", "employee_id"), ("Employee Name", "employee_name"), ("Attendance Date", "attendance_date"),
        ("Clock In", "clock_in"), ("Clock Out", "clock_out"), ("Total Work Hours", "total_work_hours"),
        ("IP / Geo-location", "location"), ("Attendance Status", "attendance_status"), ("Attendance Source", "attendance_source"),
        ("Over / Under Logged Hours", "logged_hours"), ("Shift ID / Schedule", "shift_schedule"),
    )),
    ("Leave Types", "leave_types", (
        ("Leave Type ID", "leave_type_id"), ("Leave Name", "leave_name"), ("Default Quota (Days)", "default_quota"), ("Carry Forward", "carry_forward"),
    )),
    ("Leave Balances", "leave_balances", (
        ("Balance ID", "balance_id"), ("Employee ID", "employee_id"), ("Leave Type ID", "leave_type_id"),
        ("Year", "year"), ("Allocated Days", "allocated_days"), ("Used Days", "used_days"), ("Remaining Days", "remaining_days"),
    )),
    ("Leave Requests", "leave_requests", (
        ("Request ID", "request_id"), ("Employee ID", "employee_id"), ("Leave Type ID", "leave_type_id"),
        ("Start Date", "start_date"), ("End Date", "end_date"), ("Total Days", "total_days"), ("Reason", "reason"),
        ("Status", "approval_status"), ("Approved By (Employee ID)", "approved_by"), ("Created At", "created_at"),
        ("Medical Certificate", "medical_certificate"), ("Approved Leave Form", "leave_form"),
    )),
    ("Payroll & Compensation", "payroll", (
        ("Employee ID", "employee_id"), ("Employee Name", "employee_name"), ("Effective Month", "effective_month"),
        ("Base Salary", "base_salary"), ("HRA", "hra"), ("Allowances", "allowances"), ("Bonus", "bonus"), ("Incentive", "incentive"), ("Effective From", "effective_from"),
        ("Overpay / Deductions", "deductions"), ("Bank Account Number", "bank_account"), ("Bank Name", "bank_name"),
        ("SWIFT / IFSC Code", "ifsc_code"), ("Tax / PF / Social Security ID", "tax_pf_id"), ("Tax Declaration", "tax_declaration"),
        ("Monthly Payslip", "payslip"), ("Tax Statement", "tax_statement"), ("Salary Revision Letter", "revision_letter"),
        ("Loan / Advance Request", "loan_advance"), ("Repayment Log", "repayment_log"),
    )),
    ("Payroll Runs", "payroll_runs", (
        ("Payroll ID", "payroll_id"), ("Employee ID", "employee_id"), ("Pay Period Month", "pay_period_month"),
        ("Pay Period Year", "pay_period_year"), ("Gross Salary", "gross_salary"), ("Total Deductions", "total_deductions"),
        ("Net Salary", "net_salary"), ("Payment Status", "payment_status"), ("Payment Date", "payment_date"),
    )),
    ("Performance & Goals", "performance", (
        ("Goal ID", "goal_id"), ("Employee ID", "employee_id"), ("Employee Name", "employee_name"), ("Goal Title / Objective", "objective"),
        ("Goal Description", "goal_description"), ("Key Result", "key_result"), ("Target Completion Date", "target_date"), ("Progress %", "progress_percentage"), ("Weightage", "weightage"),
        ("Goal Status", "goal_status"), ("Self Evaluation", "self_evaluation"), ("Manager Review", "manager_review"),
        ("Peer / 360 Feedback", "peer_feedback"), ("Feedback Score", "feedback_score"), ("PIP Document", "pip_document"),
        ("Appraisal Letter", "appraisal_letter"), ("Promotion / Demotion Letter", "promotion_letter"), ("Training Certificate", "training_certificate"),
    )),
    ("Appraisals", "appraisals", (
        ("Appraisal ID", "appraisal_id"), ("Employee ID", "employee_id"), ("Reviewer Employee ID", "reviewer_id"),
        ("Review Cycle", "review_cycle"), ("Self Rating", "self_rating"), ("Manager Rating", "manager_rating"),
        ("Final Comments", "final_comments"), ("Appraisal Status", "appraisal_status"),
    )),
    ("Learning & Development", "learning", (
        ("Course Title", "course_title"), ("Description", "description"), ("Assigned Participants", "participants"),
        ("Completion Status", "completion_status"), ("Due Date", "due_date"), ("Completion Certificate", "certificate"),
        ("Training Feedback Form", "feedback_form"), ("External Certification Receipt", "certification_receipt"),
    )),
    ("Offboarding & Separation", "offboarding", (
        ("Employee ID", "employee_id"), ("Employee Name", "employee_name"), ("Resignation Date", "resignation_date"),
        ("Last Working Day", "last_working_day"), ("Reason for Leaving", "leaving_reason"), ("IT Clearance", "it_clearance"),
        ("Finance Clearance", "finance_clearance"), ("HR Clearance / Exit Notes", "hr_clearance"), ("Resignation Letter", "resignation_letter"),
        ("Resignation Acceptance", "acceptance_letter"), ("Exit Interview Form", "exit_interview"),
        ("F&F Settlement Statement", "final_settlement"), ("Experience / Relieving Letter", "relieving_letter"),
    )),
)


def csv_path(register):
    os.makedirs(config.HR_DIR, exist_ok=True)
    legacy_dir = config.CSV_DIR

    if register == "employees":
        target = os.path.join(config.HR_DIR, "employee record.csv")
        legacy = os.path.join(legacy_dir, "employee record.csv")
    else:
        target = os.path.join(config.HR_DIR, f"hr_{register}.csv")
        legacy = os.path.join(legacy_dir, f"hr_{register}.csv")

    if not os.path.exists(target) and os.path.exists(legacy):
        try:
            shutil.copy2(legacy, target)
        except Exception:
            pass
    return target


class SearchableCombobox(ttk.Combobox):
    """A dropdown that narrows its values while the user types."""

    def __init__(self, parent, *, values, **kwargs):
        self.all_values = tuple(values)
        super().__init__(parent, values=self.all_values, **kwargs)
        self.bind("<KeyRelease>", self._filter_values, add=True)

    def _filter_values(self, event=None):
        if event and event.keysym in {"Up", "Down", "Left", "Right", "Return", "Escape", "Tab"}:
            return
        search_text = self.get().strip().casefold()
        self["values"] = tuple(
            value for value in self.all_values if search_text in value.casefold()
        )


class HRRegisterTab(ttk.Frame):
    def __init__(self, parent, title, register, fields):
        super().__init__(parent, padding=12)
        self.title, self.register, self.fields = title, register, fields
        self.headers = [key for _label, key in fields]
        self.vars = {key: tk.StringVar() for key in self.headers}
        self._build()
        self.load_records()
        self._set_next_employee_id()

    def _build(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)

        # ── Top header bar ─────────────────────────────────────────────────
        header = ttk.Frame(self)
        header.grid(row=0, column=0, sticky="ew", pady=(0, 6))

        ttk.Label(header, text=self.title,
                  font=("Helvetica", 14, "bold")).pack(side="left")

        ttk.Button(header, text="⬇ Download Excel",
                   command=self.export_excel).pack(side="right", padx=(4, 0))
        self._records_visible = True
        self._toggle_btn = ttk.Button(
            header, text="📋 Hide Saved Records",
            command=self._toggle_records,
        )
        self._toggle_btn.pack(side="right", padx=(4, 0))

        # ── PanedWindow: form (top) + records (bottom) ────────────────────
        self._outer = ttk.PanedWindow(self, orient="vertical")
        self._outer.grid(row=1, column=0, sticky="nsew")

        form_outer  = ttk.LabelFrame(self._outer, text=" Add / Update Record ", padding=(6, 6))
        self._list_box = ttk.LabelFrame(self._outer, text=" Saved Records ", padding=8)
        self._outer.add(form_outer,       weight=3)
        self._outer.add(self._list_box,   weight=2)

        # ── Scrollable form canvas ─────────────────────────────────────────
        form_canvas = tk.Canvas(form_outer, highlightthickness=0)
        form_vsb    = ttk.Scrollbar(form_outer, orient="vertical",   command=form_canvas.yview)
        form_hsb    = ttk.Scrollbar(form_outer, orient="horizontal", command=form_canvas.xview)
        form_canvas.configure(yscrollcommand=form_vsb.set, xscrollcommand=form_hsb.set)

        form_vsb.pack(side="right",  fill="y")
        form_hsb.pack(side="bottom", fill="x")
        form_canvas.pack(side="left", fill="both", expand=True)

        form_box = ttk.Frame(form_canvas, padding=8)
        form_win = form_canvas.create_window((0, 0), window=form_box, anchor="nw")

        form_box.bind("<Configure>",
                      lambda e: form_canvas.configure(scrollregion=form_canvas.bbox("all")))
        form_canvas.bind("<Configure>",
                         lambda e: form_canvas.itemconfig(form_win, width=e.width))

        def _on_form_scroll(e):
            form_canvas.yview_scroll(int(-1 * (e.delta / 120)), "units")
        form_canvas.bind("<Enter>",
                         lambda e: form_canvas.bind_all("<MouseWheel>", _on_form_scroll))
        form_canvas.bind("<Leave>",
                         lambda e: form_canvas.unbind_all("<MouseWheel>"))

        # ── Field grid ─────────────────────────────────────────────────────
        for column in (1, 3):
            form_box.columnconfigure(column, weight=1)

        for index, (label, key) in enumerate(self.fields):
            row, group = divmod(index, 2)
            col = group * 2
            ttk.Label(form_box, text=f"{label}:").grid(
                row=row, column=col, sticky="w", padx=(0, 6), pady=4)
            widget = self.make_widget(form_box, key)
            widget.grid(row=row, column=col + 1, sticky="ew",
                        padx=(0, 14) if group == 0 else (0, 4), pady=4)

        actions = ttk.Frame(form_box)
        actions.grid(row=(len(self.fields) + 1) // 2, column=0,
                     columnspan=4, sticky="e", pady=(12, 4))
        ttk.Button(actions, text="✖ Clear",
                   command=self.clear_form).pack(side="right")
        ttk.Button(actions, text="💾 Save Record",
                   command=self.save).pack(side="right", padx=(0, 8))

        # ── Records treeview ───────────────────────────────────────────────
        visible = self.headers[:6]
        self.tree = ttk.Treeview(self._list_box, columns=visible,
                                 show="headings", height=7)
        for key in visible:
            lbl = next(l for l, fk in self.fields if fk == key)
            self.tree.heading(key, text=lbl)
            self.tree.column(key, width=180, anchor="w")

        rec_vsb = ttk.Scrollbar(self._list_box, orient="vertical",   command=self.tree.yview)
        rec_hsb = ttk.Scrollbar(self._list_box, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=rec_vsb.set, xscrollcommand=rec_hsb.set)
        rec_vsb.pack(side="right",  fill="y")
        rec_hsb.pack(side="bottom", fill="x")
        self.tree.pack(side="left", fill="both", expand=True)
        self.tree.bind("<<TreeviewSelect>>", self._load_selected_record)

    def _toggle_records(self):
        """Show or hide the Saved Records pane."""
        if self._records_visible:
            self._outer.forget(self._list_box)
            self._toggle_btn.configure(text="📋 Show Saved Records")
            self._records_visible = False
        else:
            self._outer.add(self._list_box, weight=2)
            self._toggle_btn.configure(text="📋 Hide Saved Records")
            self._records_visible = True

    def make_widget(self, parent, key):
        if self.register == "employees" and key == "employee_id":
            return ttk.Entry(parent, textvariable=self.vars[key], state="readonly")
        if self.register == "employees" and key == "reporting_manager":
            widget = ttk.Combobox(
                parent, textvariable=self.vars[key], values=tuple(self._employee_directory()), state="readonly"
            )
            widget.bind("<<ComboboxSelected>>", self._fill_reporting_manager_name)
            widget.bind("<FocusIn>", self._refresh_reporting_manager_ids, add=True)
            self.reporting_manager_widget = widget
            return widget
        if self.register == "employees" and key == "reporting_manager_name":
            return ttk.Entry(parent, textvariable=self.vars[key], state="readonly")
        if self.register == "attendance_leave" and key == "employee_id":
            widget = ttk.Combobox(
                parent, textvariable=self.vars[key], values=tuple(self._employee_directory()), state="readonly"
            )
            widget.bind("<<ComboboxSelected>>", self._fill_attendance_employee_name)
            widget.bind("<FocusIn>", self._refresh_attendance_employee_ids, add=True)
            self.attendance_employee_id_widget = widget
            return widget
        if self.register == "attendance_leave" and key == "employee_name":
            return ttk.Entry(parent, textvariable=self.vars[key], state="readonly")
        if self.register == "attendance_leave" and key in {"clock_in", "clock_out"}:
            return ttk.Combobox(parent, textvariable=self.vars[key], values=TIME_OPTIONS, state="readonly")
        if key == "marital_status":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Married", "Unmarried"), state="readonly")
        if key == "nationality":
            return SearchableCombobox(parent, textvariable=self.vars[key], values=COUNTRIES)
        if key == "blood_group":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=BLOOD_GROUPS, state="readonly")
        if key in {"department"}:
            return ttk.Combobox(parent, textvariable=self.vars[key], values=DEPARTMENTS, state="readonly")
        if key == "employment_type":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=EMPLOYMENT_TYPES, state="readonly")
        if key == "employment_status":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Active", "On-Leave", "Terminated", "Resigned"), state="readonly")
        if key in {"gender"}:
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Female", "Male", "Non-binary", "Prefer not to say"), state="readonly")
        if key in {"offer_status"}:
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Pending", "Accepted", "Rejected"), state="readonly")
        if key == "job_status":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Draft", "Open", "On-Hold", "Closed"), state="readonly")
        if key == "current_stage":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Applied", "Screening", "Interview", "Offered", "Hired", "Rejected"), state="readonly")
        if key == "interview_status":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Scheduled", "Completed", "Cancelled"), state="readonly")
        if key in {"approval_status"}:
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Pending", "Approved", "Rejected"), state="readonly")
        if key in {"goal_status"}:
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Not Started", "In Progress", "On Track", "Completed", "Delayed"), state="readonly")
        if key == "attendance_status":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Present", "Absent", "Half-Day", "Late", "On-Leave"), state="readonly")
        if key == "attendance_source":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Web", "Mobile", "Biometric Device", "Manual Override"), state="readonly")
        if key == "payment_status":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Pending", "Processed", "Failed"), state="readonly")
        if key == "appraisal_status":
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Draft", "Submitted", "Approved", "Locked"), state="readonly")
        if key in {"completion_status"}:
            return ttk.Combobox(parent, textvariable=self.vars[key], values=("Assigned", "In Progress", "Completed", "Overdue"), state="readonly")
        if "date" in key or key in {"target_date", "due_date", "last_working_day"}:
            self.vars[key].set(date.today().isoformat())
            return DateEntry(parent, textvariable=self.vars[key], date_pattern="yyyy-mm-dd")
        if key in {"created_at", "uploaded_at"}:
            self.vars[key].set(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            return ttk.Entry(parent, textvariable=self.vars[key], state="readonly")
        if any(word in key for word in ("file", "document", "letter", "certificate", "form", "photograph", "resume", "contract", "nda", "government_id", "payslip", "statement", "receipt")):
            box = ttk.Frame(parent)
            box.columnconfigure(0, weight=1)
            ttk.Entry(box, textvariable=self.vars[key]).grid(row=0, column=0, sticky="ew")
            ttk.Button(box, text="Attach", command=lambda field=key: self.attach(field)).grid(row=0, column=1, padx=(5, 0))
            return box
        return ttk.Entry(parent, textvariable=self.vars[key])

    def attach(self, key):
        path = filedialog.askopenfilename(parent=self, title="Select document", filetypes=(("All files", "*.*"),))
        if path:
            self.vars[key].set(path)

    def ensure_file(self):
        path = csv_path(self.register)
        if self.register == "employees" and not os.path.exists(path):
            legacy_path = os.path.join(config.CSV_DIR, "hr_employees.csv")
            if os.path.exists(legacy_path):
                shutil.copyfile(legacy_path, path)
        if not os.path.exists(path):
            with open(path, "w", newline="", encoding="utf-8") as file:
                csv.DictWriter(file, fieldnames=self.headers).writeheader()
        elif self.register == "employees":
            self._upgrade_employee_file_if_needed(path)
        return path

    def _upgrade_employee_file_if_needed(self, path):
        """Keep the employee CSV header aligned with the current Employee Profile fields."""
        with open(path, newline="", encoding="utf-8") as file:
            reader = csv.DictReader(file)
            if reader.fieldnames == self.headers:
                return
            rows = list(reader)

        # Older files used a different Employee Profile header. Preserve every
        # shared field, and retain the old last-name value as the middle name.
        for row in rows:
            if not row.get("middle_name") and row.get("last_name"):
                row["middle_name"] = row["last_name"]

        with tempfile.NamedTemporaryFile(
            "w", newline="", encoding="utf-8", delete=False, dir=os.path.dirname(path)
        ) as file:
            temporary_path = file.name
            writer = csv.DictWriter(file, fieldnames=self.headers, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        os.replace(temporary_path, path)

    def _employee_directory(self):
        """Return saved employee IDs mapped to the names shown in Attendance."""
        path = csv_path("employees")
        if not os.path.exists(path):
            return {}

        directory = {}
        with open(path, newline="", encoding="utf-8") as file:
            for row in csv.DictReader(file):
                employee_id = (row.get("employee_id") or "").strip()
                if not employee_id:
                    continue
                full_name = (row.get("full_name") or "").strip()
                if not full_name:
                    full_name = " ".join(
                        part.strip()
                        for part in (row.get("first_name") or "", row.get("middle_name") or "")
                        if part.strip()
                    )
                directory[employee_id] = full_name
        return directory

    def _refresh_attendance_employee_ids(self, _event=None):
        if hasattr(self, "attendance_employee_id_widget"):
            self.attendance_employee_id_widget["values"] = tuple(self._employee_directory())

    def _fill_attendance_employee_name(self, _event=None):
        if self.register == "attendance_leave":
            self.vars["employee_name"].set(
                self._employee_directory().get(self.vars["employee_id"].get(), "")
            )

    def _refresh_reporting_manager_ids(self, _event=None):
        if hasattr(self, "reporting_manager_widget"):
            self.reporting_manager_widget["values"] = tuple(self._employee_directory())

    def _fill_reporting_manager_name(self, _event=None):
        if self.register == "employees":
            self.vars["reporting_manager_name"].set(
                self._employee_directory().get(self.vars["reporting_manager"].get(), "")
            )

    def save(self):
        identity_keys = ("employee_id", "candidate_id", "job_id", "document_id", "interview_id", "shift_id", "leave_type_id", "balance_id", "request_id", "payroll_id", "goal_id", "appraisal_id", "course_title")
        identity = next((self.vars[key] for key in identity_keys if key in self.vars), None)
        if not identity or not identity.get().strip():
            messagebox.showwarning("Required Details", "Enter the primary record name or employee ID before saving.", parent=self)
            return
        for key in ("created_at", "uploaded_at"):
            if key in self.vars:
                self.vars[key].set(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        path = self.ensure_file()
        record = {key: value.get().strip() for key, value in self.vars.items()}
        if self.register == "employees":
            with open(path, newline="", encoding="utf-8") as file:
                rows = list(csv.DictReader(file))
            employee_id = record["employee_id"]
            updated = False
            for row in rows:
                if (row.get("employee_id") or "").strip() == employee_id:
                    row.update(record)
                    updated = True
                    break
            if not updated:
                rows.append(record)
            with open(path, "w", newline="", encoding="utf-8") as file:
                writer = csv.DictWriter(file, fieldnames=self.headers)
                writer.writeheader()
                writer.writerows(rows)
        else:
            with open(path, "a", newline="", encoding="utf-8") as file:
                csv.DictWriter(file, fieldnames=self.headers).writerow(record)
        self.load_records()
        self.clear_form()
        messagebox.showinfo("Saved", f"{self.title} record saved successfully.", parent=self)

    def load_records(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        with open(self.ensure_file(), newline="", encoding="utf-8") as file:
            for row in csv.DictReader(file):
                self.tree.insert("", "end", values=[row.get(key, "") for key in self.headers[:6]])

    def _load_selected_record(self, _event=None):
        if self.register != "employees":
            return
        selection = self.tree.selection()
        if not selection:
            return
        employee_id = str(self.tree.item(selection[0], "values")[0]).strip()
        with open(self.ensure_file(), newline="", encoding="utf-8") as file:
            for row in csv.DictReader(file):
                if (row.get("employee_id") or "").strip() == employee_id:
                    for key in self.headers:
                        self.vars[key].set(row.get(key, ""))
                    return

    def _set_next_employee_id(self):
        """Populate the Employee Profile form with the next sequential ID."""
        if self.register != "employees":
            return

        highest_number = 0
        path = self.ensure_file()
        id_pattern = re.compile(rf"^{re.escape(EMPLOYEE_ID_PREFIX)}(\d+)$", re.IGNORECASE)
        with open(path, newline="", encoding="utf-8") as file:
            for row in csv.DictReader(file):
                match = id_pattern.match((row.get("employee_id") or "").strip())
                if match:
                    highest_number = max(highest_number, int(match.group(1)))

        self.vars["employee_id"].set(f"{EMPLOYEE_ID_PREFIX}{highest_number + 1:02d}")

    def clear_form(self):
        for key, value in self.vars.items():
            if "date" not in key and key not in {"target_date", "due_date", "last_working_day"}:
                value.set("")
        for key in ("created_at", "uploaded_at"):
            if key in self.vars:
                self.vars[key].set(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        self._set_next_employee_id()

    def export_excel(self):
        """Export all saved records for this register to an Excel file."""
        src = csv_path(self.register)
        if not os.path.exists(src):
            messagebox.showinfo("No Data",
                                f"No records found for {self.title}.", parent=self)
            return

        import csv as _csv
        with open(src, newline="", encoding="utf-8") as f:
            rows = list(_csv.DictReader(f))

        if not rows:
            messagebox.showinfo("No Data",
                                f"No records saved in {self.title} yet.", parent=self)
            return

        path = filedialog.asksaveasfilename(
            parent=self,
            title=f"Save {self.title} — Excel",
            defaultextension=".xlsx",
            initialfile=f"HR_{self.register}_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")],
        )
        if not path:
            return

        try:
            from openpyxl import Workbook as _WB
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = _WB()
            ws = wb.active
            ws.title = self.title[:31]   # sheet name max 31 chars

            # Header row — friendly labels
            label_map = {key: lbl for lbl, key in self.fields}
            headings  = [label_map.get(h, h) for h in self.headers]
            ws.append(headings)
            for cell in ws[1]:
                cell.font      = Font(bold=True, color="FFFFFF")
                cell.fill      = PatternFill("solid", fgColor="2E4057")
                cell.alignment = Alignment(horizontal="center")

            # Data rows
            for row in rows:
                ws.append([row.get(h, "") for h in self.headers])

            # Auto column width
            for col_cells in ws.columns:
                max_len = max(
                    (len(str(c.value)) for c in col_cells if c.value), default=10
                )
                ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)

            wb.save(path)
            messagebox.showinfo("Exported",
                                f"{self.title} records saved to:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Export Failed", str(e), parent=self)


class HRView(ttk.Frame):
    def __init__(self, parent, user_data=None):
        super().__init__(parent, padding=14)
        ttk.Label(self, text="Human Resources",
                  font=("Helvetica", 16, "bold")).pack(anchor="w")
        ttk.Label(self, text="Employee records, recruitment, attendance, payroll, performance, training, and separation.").pack(anchor="w", pady=(2, 10))

        # ── Main layout: left sidebar | right content ─────────────────────
        body = ttk.Frame(self)
        body.pack(fill="both", expand=True)
        body.columnconfigure(0, weight=0)   # sidebar fixed width
        body.columnconfigure(1, weight=0)   # scrollbar
        body.columnconfigure(2, weight=1)   # content expands
        body.rowconfigure(0, weight=1)

        # ── Scrollable sidebar ────────────────────────────────────────────
        sidebar_canvas = tk.Canvas(body, width=220, highlightthickness=0,
                                   bd=0, bg="#f0f0f0")
        sidebar_canvas.grid(row=0, column=0, sticky="nsew")

        sidebar_scroll = ttk.Scrollbar(body, orient="vertical",
                                       command=sidebar_canvas.yview)
        sidebar_scroll.grid(row=0, column=1, sticky="ns")
        sidebar_canvas.configure(yscrollcommand=sidebar_scroll.set)

        # Inner frame inside canvas
        sidebar_inner = ttk.Frame(sidebar_canvas)
        sidebar_window = sidebar_canvas.create_window(
            (0, 0), window=sidebar_inner, anchor="nw"
        )

        def _on_sidebar_configure(e):
            sidebar_canvas.configure(scrollregion=sidebar_canvas.bbox("all"))

        def _on_canvas_resize(e):
            sidebar_canvas.itemconfig(sidebar_window, width=e.width)

        sidebar_inner.bind("<Configure>", _on_sidebar_configure)
        sidebar_canvas.bind("<Configure>", _on_canvas_resize)

        # Mouse-wheel scroll on sidebar
        def _on_mousewheel(e):
            sidebar_canvas.yview_scroll(int(-1 * (e.delta / 120)), "units")
        sidebar_canvas.bind_all("<MouseWheel>", _on_mousewheel)

        # Separator
        ttk.Separator(body, orient="vertical").grid(
            row=0, column=1, sticky="ns", padx=(2, 2)
        )

        # ── Right content area ────────────────────────────────────────────
        content_area = ttk.Frame(body)
        content_area.grid(row=0, column=2, sticky="nsew", padx=(6, 0))
        content_area.rowconfigure(0, weight=1)
        content_area.columnconfigure(0, weight=1)

        # ── Styles ────────────────────────────────────────────────────────
        style = ttk.Style()
        style.configure("HR.Sidebar.TButton",
                        anchor="w", padding=(8, 5), font=("Helvetica", 9))
        style.configure("HR.SidebarActive.TButton",
                        anchor="w", padding=(8, 5), font=("Helvetica", 9, "bold"))

        ICONS = {
            "Employee Profile":         "👤",
            "Employee Documents":       "📄",
            "Recruitment (ATS)":        "🔍",
            "Interviews":               "🎙️",
            "Shift Setup":              "⏰",
            "Attendance & Leave":       "📅",
            "Leave Types":              "📋",
            "Leave Balances":           "⚖️",
            "Leave Requests":           "📩",
            "Payroll & Compensation":   "💰",
            "Payroll Runs":             "💳",
            "Performance & Goals":      "🎯",
            "Appraisals":               "⭐",
            "Learning & Development":   "📚",
            "Offboarding & Separation": "🚪",
        }

        btn_refs   = []
        tab_frames = {}

        def show_tab(frame, btn):
            for f in tab_frames.values():
                f.grid_remove()
            frame.grid(row=0, column=0, sticky="nsew")
            for b in btn_refs:
                b.configure(style="HR.Sidebar.TButton")
            btn.configure(style="HR.SidebarActive.TButton")

        first_frame = first_btn = None

        for title, register, fields in HR_TABS:
            # Content frame
            frame = HRRegisterTab(content_area, title, register, fields)
            frame.grid(row=0, column=0, sticky="nsew")
            frame.grid_remove()
            tab_frames[title] = frame

            icon = ICONS.get(title, "•")
            btn = ttk.Button(
                sidebar_inner,
                text=f" {icon}  {title}",
                style="HR.Sidebar.TButton",
                width=24,
            )
            btn.configure(command=lambda f=frame, b=btn: show_tab(f, b))
            btn.pack(fill="x", padx=4, pady=2)
            btn_refs.append(btn)

            if first_frame is None:
                first_frame, first_btn = frame, btn

        # Show first tab
        if first_frame:
            show_tab(first_frame, first_btn)
